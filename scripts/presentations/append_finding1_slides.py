#!/usr/bin/env python3
"""Insert a seven-slide Finding 1 section into the dated sex/APOE KDA deck."""

from __future__ import annotations

import argparse
import copy
import csv
import hashlib
import os
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Any, Sequence

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from scipy.stats import hypergeom

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
import update_phase11_seaad_simple_aggr_part2 as ui  # noqa: E402


DECK = (
    ROOT
    / "docs"
    / "presentations"
    / "09162026"
    / "09162026_sex_apoe_kda_fine_broad.pptx"
)
AUDIT_ROOT = (
    ROOT / "results" / "presentations" / "09162026_finding1_slides"
)
ROS_DIR = ROOT / "results" / "minerva_production" / "20_sex_apoe_kda_combo"
SEA_DIR = ROOT / "results" / "validation_human" / "12_sex_apoe_kda_combo"

EXPECTED_INPUT_SLIDES = 28
INSERT_INDEX = EXPECTED_INPUT_SLIDES
SECTION_MARKER = "FINDING 1"
SECTION_TITLE = (
    "Female APOE ε3/ε3 excitatory neurons show a reproducible "
    "mtDNA-OXPHOS increase"
)
NEW_SLIDE_COUNT = 7


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DECK)
    parser.add_argument("--output", type=Path, default=DECK)
    parser.add_argument("--audit-root", type=Path, default=AUDIT_ROOT)
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_tsv(path: Path) -> pd.DataFrame:
    if not path.is_file():
        raise FileNotFoundError(path)
    return pd.read_csv(path, sep="\t", low_memory=False)


def true_series(series: pd.Series) -> pd.Series:
    return series.astype(str).str.upper().isin({"TRUE", "T", "1", "YES"})


def bh_adjust(values: Sequence[float]) -> np.ndarray:
    p = np.asarray(values, dtype=float)
    order = np.argsort(p)
    adjusted = np.empty(len(p), dtype=float)
    ranked = p[order] * len(p) / np.arange(1, len(p) + 1)
    adjusted[order] = np.minimum.accumulate(ranked[::-1])[::-1]
    return np.minimum(adjusted, 1.0)


def semantic_slide_fingerprint(slide) -> tuple[Any, ...]:
    shapes = []
    for shape in slide.shapes:
        text = shape.text if getattr(shape, "has_text_frame", False) else ""
        shapes.append(
            (
                int(shape.shape_type),
                shape.name,
                int(shape.left),
                int(shape.top),
                int(shape.width),
                int(shape.height),
                text,
            )
        )
    notes = slide.notes_slide.notes_text_frame
    return tuple(shapes), notes.text if notes is not None else ""


def slide_text(slide) -> str:
    return "\n".join(
        shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)
    )


def prepare_facts() -> dict[str, Any]:
    modules = read_tsv(ROOT / "config" / "phase13_respiratory_modules.tsv")
    module_sets = {
        module_id: set(rows["current_approved_symbol"])
        for module_id, rows in modules.groupby("module_id")
    }
    mt_genes = module_sets["mtdna_oxphos_13"]

    ros_query = read_tsv(ROS_DIR / "combo_query_members.tsv.gz")
    ros_query = ros_query.loc[
        true_series(ros_query["included_at_thresholds"])
        & true_series(ros_query["effective_member"])
    ].copy()
    ros_manifest = read_tsv(ROS_DIR / "combo_run_manifest.tsv")
    ros_included = ros_manifest.loc[
        true_series(ros_manifest["included_at_thresholds"])
        & ros_manifest["signature_group"].eq("F_e33")
    ].copy()
    ros_exc_runs = ros_included.loc[
        ros_included["broad_cell_type"].eq("Excitatory_neurons")
    ].copy()

    ros_group_mt = ros_query.loc[
        ros_query["signature_group"].eq("F_e33")
        & ros_query["gene"].isin(mt_genes)
    ]
    ros_exc_mt = ros_group_mt.loc[
        ros_group_mt["broad_cell_type"].eq("Excitatory_neurons")
    ]

    background = read_tsv(
        ROOT / "results" / "minerva_production" / "12_rosmap_sex_apoe_mito_kda_runs"
        / "kda_background_members.tsv.gz"
    )
    background = background.loc[
        background["kda_run_id"].isin(ros_included["kda_run_id"])
    ]
    annotation = read_tsv(
        ROOT / "results" / "minerva_production" / "09_annotate_genes"
        / "gene_annotation_master.tsv.gz"
    )
    mito_annotation = annotation.loc[
        true_series(annotation["is_mitocarta3"])
        & annotation["symbol_hgnc_current"].notna()
    ]
    annotation_sets = {
        rds_id: set(rows["symbol_hgnc_current"])
        for rds_id, rows in mito_annotation.groupby("rds_id")
    }
    background_sets = {
        run_id: set(rows["gene"])
        for run_id, rows in background.groupby("kda_run_id")
    }
    query_sets = {
        run_id: set(rows["gene"])
        for run_id, rows in ros_query.groupby("kda_run_id")
    }

    significant_mt_runs: set[str] = set()
    for row in ros_included.itertuples(index=False):
        source_rds = row.source_contrast_id.split("::", maxsplit=1)[0]
        bg = background_sets[row.kda_run_id] & annotation_sets[source_rds]
        query = query_sets[row.kda_run_id]
        p_values: list[float] = []
        module_ids: list[str] = []
        for _, module_rows in modules.groupby("module_order"):
            module_id = str(module_rows["module_id"].iloc[0])
            genes = set(module_rows["current_approved_symbol"])
            p_values.append(
                hypergeom.sf(
                    len(query & genes) - 1,
                    len(bg),
                    len(bg & genes),
                    len(query),
                )
            )
            module_ids.append(module_id)
        adjusted = dict(zip(module_ids, bh_adjust(p_values), strict=True))
        if adjusted["mtdna_oxphos_13"] < 0.05:
            significant_mt_runs.add(row.kda_run_id)

    sea_manifest = read_tsv(
        SEA_DIR / "10a_inputs" / "seaad_kda_run_manifest.tsv"
    )
    sea_match = sea_manifest.loc[
        sea_manifest["signature_group"].eq("F_e33")
        & sea_manifest["broad_network"].eq("Excitatory_neurons")
        & sea_manifest["kda_run_id"].notna()
    ].copy()
    sea_members = read_tsv(
        SEA_DIR / "10a_inputs" / "seaad_kda_signature_members.tsv.gz"
    )
    sea_effective = sea_members.loc[
        sea_members["kda_run_id"].isin(sea_match["kda_run_id"])
        & true_series(sea_members["effective_member"])
    ]
    sea_mt_genes = sorted(set(sea_effective["gene"]) & mt_genes)
    sea_result = read_tsv(ROOT / str(sea_match.iloc[0]["result_path"]))
    sea_mt_result = sea_result.loc[
        sea_result["current_symbol_for_kda"].isin(sea_mt_genes)
        & (pd.to_numeric(sea_result["FDR"], errors="coerce") < 0.05)
    ]

    overlap_raw_p = float(hypergeom.sf(9 - 1, 392, 150, 10))
    overlap_bh = overlap_raw_p * 5
    facts = {
        "ros_group_occurrences": len(ros_group_mt),
        "ros_group_up": int(true_series(ros_group_mt["in_upregulated_query"]).sum()),
        "ros_group_down": int(
            true_series(ros_group_mt["in_downregulated_query"]).sum()
        ),
        "ros_group_significant_calls": sum(
            run_id in significant_mt_runs for run_id in ros_included["kda_run_id"]
        ),
        "ros_exc_eligible_calls": len(ros_exc_runs),
        "ros_exc_occurrences": len(ros_exc_mt),
        "ros_exc_up": int(true_series(ros_exc_mt["in_upregulated_query"]).sum()),
        "ros_exc_down": int(true_series(ros_exc_mt["in_downregulated_query"]).sum()),
        "ros_exc_significant_calls": sum(
            run_id in significant_mt_runs for run_id in ros_exc_runs["kda_run_id"]
        ),
        "sea_evaluable_calls": len(sea_match),
        "sea_query_genes": len(sea_effective),
        "sea_shared_mt_genes": sea_mt_genes,
        "sea_shared_up": int(sea_mt_result["effect_direction"].eq("Dementia_up").sum()),
        "cross_cohort_raw_p": overlap_raw_p,
        "cross_cohort_bh": overlap_bh,
        "identity_resolved_bh": 0.00441,
    }
    expected = {
        "ros_group_occurrences": 217,
        "ros_group_up": 217,
        "ros_group_down": 0,
        "ros_group_significant_calls": 29,
        "ros_exc_eligible_calls": 14,
        "ros_exc_occurrences": 123,
        "ros_exc_up": 123,
        "ros_exc_down": 0,
        "ros_exc_significant_calls": 12,
        "sea_evaluable_calls": 1,
        "sea_query_genes": 10,
        "sea_shared_up": 9,
    }
    for key, value in expected.items():
        if facts[key] != value:
            raise RuntimeError(f"Source drift for {key}: {facts[key]} != {value}")
    expected_genes = [
        "MT-ATP6",
        "MT-CO1",
        "MT-CO2",
        "MT-CO3",
        "MT-CYB",
        "MT-ND1",
        "MT-ND2",
        "MT-ND4",
        "MT-ND5",
    ]
    if sea_mt_genes != expected_genes:
        raise RuntimeError(f"Unexpected SEA-AD shared genes: {sea_mt_genes}")
    if not np.isclose(overlap_raw_p, 0.0010101273546736398):
        raise RuntimeError(f"Unexpected overlap P value: {overlap_raw_p}")
    if not np.isclose(overlap_bh, 0.005050636773368199):
        raise RuntimeError(f"Unexpected overlap BH value: {overlap_bh}")
    return facts


def add_notes(
    slide,
    goal: str,
    walkthrough: str,
    boundary: str,
    transition: str,
) -> None:
    ui.add_notes(
        slide,
        goal=goal,
        walkthrough=walkthrough,
        boundary=boundary,
        transition=transition,
    )


def add_source(slide, text: str) -> None:
    ui.add_text(
        slide,
        text,
        0.72,
        7.08,
        11.90,
        0.16,
        size=7.3,
        color=ui.GRAY,
        align=PP_ALIGN.RIGHT,
    )


def add_content_title(slide, title: str, subtitle: str) -> None:
    """Add a title block with enough height for a wrapped long title."""
    title_size = 20.0 if len(title) >= 78 else 22.0
    ui.add_text(
        slide,
        title,
        0.58,
        0.32,
        12.10,
        0.70,
        size=title_size,
        color=ui.NAVY,
        bold=True,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        subtitle,
        0.60,
        1.06,
        11.95,
        0.28,
        size=10.8,
        color=ui.GRAY,
    )


def add_chip(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    *,
    bg,
    accent,
    size: float = 11.2,
) -> None:
    ui.add_rect(slide, x, y, w, 0.48, color=bg, outline=accent)
    ui.add_text(
        slide,
        text,
        x + 0.10,
        y + 0.11,
        w - 0.20,
        0.23,
        size=size,
        color=ui.readable_accent(accent),
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_stat_card(
    slide,
    value: str,
    label: str,
    x: float,
    y: float,
    w: float,
    *,
    bg,
    accent,
    detail: str | None = None,
) -> None:
    height = 1.50 if detail else 1.25
    ui.add_rect(slide, x, y, w, height, color=bg, outline=accent)
    ui.add_text(
        slide,
        value,
        x + 0.18,
        y + 0.14,
        w - 0.36,
        0.53,
        size=28 if len(value) <= 8 else 23,
        color=ui.readable_accent(accent),
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        label,
        x + 0.18,
        y + 0.72,
        w - 0.36,
        0.34,
        size=10.2,
        color=ui.NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    if detail:
        ui.add_text(
            slide,
            detail,
            x + 0.18,
            y + 1.08,
            w - 0.36,
            0.24,
            size=8.8,
            color=ui.GRAY,
            align=PP_ALIGN.CENTER,
        )


def set_shape_text(slide, shape_name: str, value: str) -> None:
    matches = [shape for shape in slide.shapes if shape.name == shape_name]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one shape named {shape_name!r}, found {len(matches)}"
        )
    frame = matches[0].text_frame
    first = frame.paragraphs[0]
    if first.runs:
        first.runs[0].text = value
        for run in first.runs[1:]:
            run.text = ""
    else:
        first.add_run().text = value
    for paragraph in frame.paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)


def set_reviewed_text_runs(
    slide,
    shape_name: str,
    segments: Sequence[str],
    language_segment: int | None = None,
) -> None:
    """Set one reviewed text line while preserving its manual run boundaries."""
    shape = next(shape for shape in slide.shapes if shape.name == shape_name)
    paragraph = shape.text_frame.paragraphs[0]
    if not paragraph.runs:
        raise RuntimeError(f"Expected styled text in shape {shape_name!r}")
    template_properties = copy.deepcopy(
        paragraph.runs[0]._r.get_or_add_rPr()
    )
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    for line_break in list(paragraph._p.br_lst):
        paragraph._p.remove(line_break)
    if paragraph._p.endParaRPr is not None:
        paragraph._p.remove(paragraph._p.endParaRPr)

    for index, segment in enumerate(segments):
        run = paragraph.add_run()
        run.text = segment
        properties = run._r.get_or_add_rPr()
        properties.attrib.clear()
        for child in list(properties):
            properties.remove(child)
        properties.attrib.update(template_properties.attrib)
        properties.set("dirty", "0")
        if index == language_segment:
            attributes = list(properties.attrib.items())
            properties.attrib.clear()
            properties.set("lang", "en-US")
            for name, value in attributes:
                properties.set(name, value)
        else:
            properties.attrib.pop("lang", None)
        for child in template_properties:
            properties.append(copy.deepcopy(child))


def sync_reviewed_base_slides(prs: Presentation) -> None:
    """Reproduce the reviewed edits to the overview and section slides."""
    title_slide = prs.slides[0]
    title_shape = next(
        shape for shape in title_slide.shapes if shape.name == "TextBox 2"
    )
    title_shape.left = 713232
    title_shape.top = 1131130
    title_shape.width = 10287000
    title_shape.height = 1194173
    title_frame = title_shape.text_frame
    title_frame.clear()
    title_frame.word_wrap = True
    title_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    overview_paragraph = title_frame.paragraphs[0]
    overview_paragraph.alignment = None
    overview_paragraph.line_spacing = None
    overview_paragraph.space_before = None
    overview_paragraph.space_after = None
    overview_run = overview_paragraph.add_run()
    overview_run.text = "DEG and key-driver analysis overview"
    overview_run.font.name = ui.FONT_HEAD
    overview_run.font.size = Pt(34.0)
    overview_run.font.color.rgb = ui.WHITE
    overview_run.font.bold = True
    overview_paragraph.add_line_break()

    cohort_paragraph = title_frame.add_paragraph()
    cohort_run = cohort_paragraph.add_run()
    cohort_run.text = "ROSMAP "
    cohort_run.font.name = ui.FONT_HEAD
    cohort_run.font.size = Pt(20.0)
    cohort_run.font.color.rgb = ui.WHITE
    cohort_run.font.bold = True
    cohort_tail = cohort_paragraph.add_run()
    cohort_tail.text = "and SEA-AD sex/APOE"
    cohort_tail.font.name = ui.FONT_HEAD
    cohort_tail.font.size = Pt(20.0)
    cohort_tail.font.color.rgb = ui.WHITE
    cohort_tail.font.bold = True
    cohort_tail.font.italic = False

    for run in (overview_run, cohort_run, cohort_tail):
        run_properties = run._r.get_or_add_rPr()
        run_properties.set("lang", "en-US")
        run_properties.set("dirty", "0")

    break_properties = overview_paragraph._p.br_lst[-1].get_or_add_rPr()
    overview_properties = overview_run._r.get_or_add_rPr()
    break_properties.attrib.update(overview_properties.attrib)
    for child in overview_properties:
        break_properties.append(copy.deepcopy(child))

    end_properties = overview_paragraph._p.get_or_add_endParaRPr()
    cohort_properties = cohort_run._r.get_or_add_rPr()
    end_properties.attrib.update(cohort_properties.attrib)
    for child in cohort_properties:
        end_properties.append(copy.deepcopy(child))
    if overview_paragraph._p.pPr is not None:
        overview_paragraph._p.remove(overview_paragraph._p.pPr)

    part_one_matches = [
        shape for shape in title_slide.shapes if shape.name == "TextBox 1"
    ]
    if part_one_matches:
        part_one = part_one_matches[0]
        set_shape_text(title_slide, "TextBox 1", "PART 1")
    else:
        part_one = ui.add_text(
            title_slide,
            "PART 1",
            0.78,
            0.67,
            2.40,
            0.28,
            size=10.5,
            color=ui.BLUE,
            bold=True,
            font=ui.FONT_HEAD,
        )
        part_one.name = "TextBox 1"
    part_one.left, part_one.top = 713232, 612648
    part_one.width, part_one.height = 2194560, 256032

    contrast_slide = prs.slides[2]
    set_reviewed_text_runs(
        contrast_slide,
        "TextBox 21",
        ("One contrast → ", "run DEG and ", "at most one KDA call"),
        language_segment=1,
    )
    contrast_heading = next(
        shape for shape in contrast_slide.shapes if shape.name == "TextBox 21"
    )
    contrast_heading.left = 1024126
    contrast_heading.top = 2889504
    contrast_heading.width = 4855465
    contrast_heading.height = 289310

    set_reviewed_text_runs(
        prs.slides[3],
        "TextBox 2",
        ("Four steps: ", "run DEG per contrast, ", "one KDA query per contrast"),
        language_segment=1,
    )

    for slide_index, part_label in ((5, "PART 2"), (16, "PART 3"), (22, "PART 4")):
        slide = prs.slides[slide_index]
        set_shape_text(slide, "TextBox 1", part_label)
        label = next(shape for shape in slide.shapes if shape.name == "TextBox 1")
        label.height = 216982


def build_divider(prs: Presentation):
    template = prs.slides[22]
    slide = ui.new_slide(prs, bg=ui.NAVY)
    for shape in template.shapes:
        slide.shapes._spTree.insert_element_before(
            copy.deepcopy(shape._element), "p:extLst"
        )
    replacements = {
        "TextBox 1": SECTION_MARKER,
        "TextBox 3": SECTION_TITLE,
        "TextBox 4": (
            "A mitochondrial DEG-program finding, not an identified "
            "upstream key driver."
        ),
        "TextBox 6": "IN THIS FINDING",
        "TextBox 9": "Plain-language statement",
        "TextBox 12": "ROSMAP evidence",
        "TextBox 15": "SEA-AD support",
        "TextBox 18": "Meaning, limits, literature",
        "TextBox 20": (
            "SEA-AD can support the same biological pattern without "
            "matching the exact key driver."
        ),
    }
    for shape_name, value in replacements.items():
        set_shape_text(slide, shape_name, value)
    title_shape = next(shape for shape in slide.shapes if shape.name == "TextBox 3")
    title_shape.text_frame.paragraphs[0].runs[0].font.size = Pt(23.0)
    add_notes(
        slide,
        "Introduce the first biological finding.",
        "Female APOE epsilon-3 homozygous cells repeatedly show higher mitochondrial-DNA-encoded oxidative-phosphorylation RNA in Alzheimer’s disease. The matched excitatory-neuron program also appears in SEA-AD.",
        "This section describes a mitochondrial DEG program. It does not identify a causal key driver, prove better energy production, or establish a sex-by-APOE interaction.",
        "State the finding in plain language.",
    )
    return slide


def build_statement_slide(prs: Presentation):
    slide = ui.new_slide(prs)
    add_content_title(
        slide,
        "Female ε3/ε3 cells (especially excitatory neurons) raise mtDNA energy-gene RNA in AD",
        "The result is repeated across ROSMAP calls and receives program-level support from SEA-AD.",
    )
    add_chip(slide, "Female", 1.12, 1.48, 2.10, bg=ui.PALE_SKY, accent=ui.BLUE)
    add_chip(
        slide,
        "APOE ε3/ε3",
        3.52,
        1.48,
        2.42,
        bg=ui.PALE_GREEN,
        accent=ui.TEAL,
    )
    add_chip(
        slide,
        "Excitatory neurons",
        6.24,
        1.48,
        2.66,
        bg=ui.PALE_GOLD,
        accent=ui.GOLD,
    )
    add_chip(
        slide,
        "Alzheimer’s disease",
        9.20,
        1.48,
        2.82,
        bg=ui.PALE_RED,
        accent=ui.VERMILION,
    )
    ui.add_rect(slide, 0.82, 2.26, 11.70, 3.10, color=ui.PALE_SKY, outline=None)
    ui.add_text(
        slide,
        "Mitochondrial DNA",
        1.22,
        2.77,
        3.30,
        0.42,
        size=18.0,
        color=ui.NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        "makes RNA instructions\nfor OXPHOS parts",
        1.22,
        3.40,
        3.30,
        0.82,
        size=14.0,
        color=ui.DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    ui.add_text(
        slide,
        "→",
        5.02,
        3.15,
        1.00,
        0.74,
        size=34.0,
        color=ui.BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        "RNA abundance",
        6.38,
        2.77,
        3.20,
        0.42,
        size=18.0,
        color=ui.NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        "↑",
        9.72,
        2.83,
        1.24,
        1.26,
        size=52.0,
        color=ui.TEAL_TEXT,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        "more mtDNA-OXPHOS RNA\nin AD than in the comparison group",
        6.12,
        3.43,
        3.72,
        0.78,
        size=13.8,
        color=ui.DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    ui.add_rect(slide, 0.82, 5.72, 11.70, 0.76, color=ui.WHITE, outline=ui.LIGHT)
    ui.add_text(
        slide,
        "Plain meaning: the cells make more RNA instructions for mitochondrial energy-system parts.",
        1.08,
        5.94,
        11.18,
        0.30,
        size=14.0,
        color=ui.PURPLE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_source(
        slide,
        "Finding 1 synthesis; ROSMAP discovery with supplemental SEA-AD program support.",
    )
    add_notes(
        slide,
        "Explain the finding without statistical jargon.",
        "In Alzheimer’s disease, female APOE epsilon-3 homozygous cells repeatedly contain more RNA from mitochondrial OXPHOS genes. Excitatory neurons provide the cleanest comparison across ROSMAP and SEA-AD.",
        "More RNA is not the same as more protein, more ATP, or healthier mitochondria. The increase could be compensatory or stress related.",
        "Show the ROSMAP numbers behind the statement.",
    )
    return slide


def build_rosmap_slide(prs: Presentation, facts: dict[str, Any]):
    slide = ui.new_slide(prs)
    ui.add_text(
        slide,
        "ROSMAP: the female ε3/ε3 mtDNA-OXPHOS signal is consistently AD-up",
        0.58,
        0.32,
        12.10,
        0.70,
        size=22.0,
        color=ui.NAVY,
        bold=True,
        font=ui.FONT_HEAD,
    )
    top_panel = ui.add_rect(
        slide,
        0.72,
        2.628,
        11.90,
        1.24,
        color=ui.PALE_SKY,
        outline=None,
    )
    top_panel.name = "Rounded Rectangle 3"
    top_metric = ui.add_text(
        slide,
        f"{facts['ros_group_occurrences']} / {facts['ros_group_occurrences']} mtDNA-OXPHOS occurrences were AD-up",
        4.08,
        2.94,
        5.26,
        0.52,
        size=20.0,
        color=ui.BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        font=ui.FONT_HEAD,
    )
    top_metric.name = "TextBox 5"
    enriched_count = ui.add_text(
        slide,
        f"{facts['ros_group_significant_calls']} F_e33 input queries enriched",
        9.42,
        3.09,
        2.80,
        0.263,
        size=12.0,
        color=ui.GRAY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    enriched_count.name = "TextBox 6"

    excitatory_accent = ui.add_rect(
        slide,
        0.90,
        4.278,
        0.07,
        0.31,
        color=ui.TEAL,
        outline=None,
        radius=False,
    )
    excitatory_accent.name = "Rectangle 7"
    excitatory_title = ui.add_text(
        slide,
        "For excitatory-neuron only",
        1.06,
        4.248,
        11.52,
        0.316,
        size=15.2,
        color=ui.NAVY,
        bold=True,
        font=ui.FONT_HEAD,
    )
    excitatory_title.name = "TextBox 8"

    shape_start = len(slide.shapes)
    add_stat_card(
        slide,
        str(facts["ros_exc_eligible_calls"]),
        "eligible fine-cell DEG queries",
        0.90,
        4.848,
        3.56,
        bg=ui.PALE_GREEN,
        accent=ui.TEAL,
    )
    card_shapes = list(slide.shapes)[shape_start:]
    for shape, name in zip(
        card_shapes,
        ["Rounded Rectangle 9", "TextBox 10", "TextBox 11"],
        strict=True,
    ):
        shape.name = name
    card_shapes[-1].height = Inches(0.232)

    shape_start = len(slide.shapes)
    add_stat_card(
        slide,
        str(facts["ros_exc_occurrences"]),
        "mtDNA-OXPHOS DEG occurrences",
        4.88,
        4.848,
        3.56,
        bg=ui.PALE_SKY,
        accent=ui.BLUE,
    )
    for shape, name in zip(
        list(slide.shapes)[shape_start:],
        ["Rounded Rectangle 12", "TextBox 13", "TextBox 14"],
        strict=True,
    ):
        shape.name = name

    shape_start = len(slide.shapes)
    add_stat_card(
        slide,
        f"{facts['ros_exc_significant_calls']}/{facts['ros_exc_eligible_calls']}",
        "KDA input queries enriched",
        8.86,
        4.848,
        3.56,
        bg=ui.PALE_GOLD,
        accent=ui.GOLD,
    )
    for shape, name in zip(
        list(slide.shapes)[shape_start:],
        ["Rounded Rectangle 15", "TextBox 16", "TextBox 17"],
        strict=True,
    ):
        shape.name = name

    direction_panel = ui.add_rect(
        slide,
        0.90,
        6.431,
        11.52,
        0.66,
        color=ui.WHITE,
        outline=ui.LIGHT,
    )
    direction_panel.name = "Rounded Rectangle 18"
    direction_bar = ui.add_rect(
        slide,
        1.14,
        6.718,
        9.76,
        0.20,
        color=ui.TEAL,
        outline=None,
        radius=False,
    )
    direction_bar.name = "Rectangle 19"
    up_label = ui.add_text(
        slide,
        f"{facts['ros_exc_up']} AD-up",
        1.14,
        6.448,
        2.40,
        0.24,
        size=10.5,
        color=ui.TEAL_TEXT,
        bold=True,
    )
    up_label.name = "TextBox 20"
    down_label = ui.add_text(
        slide,
        f"{facts['ros_exc_down']} AD-down",
        10.88,
        6.448,
        1.24,
        0.24,
        size=10.5,
        color=ui.VERMILION_TEXT,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )
    down_label.name = "TextBox 21"

    explanation_panel = ui.add_rect(
        slide,
        0.72,
        1.324,
        11.90,
        0.641,
        color=ui.PALE_GRAY,
        outline=None,
    )
    explanation_panel.name = "Rounded Rectangle 22"
    explanation = ui.add_text(
        slide,
        "Enriched",
        1.16,
        1.455,
        11.00,
        0.464,
        size=12.0,
        color=RGBColor(255, 0, 0),
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    explanation.name = "TextBox 23"
    explanation_paragraph = explanation.text_frame.paragraphs[0]
    explanation_tail = explanation_paragraph.add_run()
    explanation_tail.text = (
        " = more mtDNA-OXPHOS genes than expected by chance in the exact "
        "background (BH-adjusted P < 0.05)"
    )
    ui.set_run(explanation_tail, size=12.0, color=ui.GRAY)
    for text in (
        ", test genes in input query to ",
        "call_key_drivers",
        "().",
    ):
        run = explanation_paragraph.add_run()
        run.text = text
        run.font.name = ui.FONT_BODY
        run.font.size = Pt(12.0)
        run.font.color.rgb = ui.GRAY

    all_calls_accent = ui.add_rect(
        slide,
        0.886,
        2.197,
        0.07,
        0.31,
        color=ui.TEAL,
        outline=None,
        radius=False,
    )
    all_calls_accent.name = "Rectangle 25"
    all_calls_title = ui.add_text(
        slide,
        "For all eligible F_e33 calls",
        1.046,
        2.167,
        11.36,
        0.316,
        size=15.2,
        color=ui.NAVY,
        bold=True,
        font=ui.FONT_HEAD,
    )
    all_calls_title.name = "TextBox 26"

    # Preserve the reviewed PowerPoint geometry on slide 31 exactly. These
    # values are EMUs copied from the deck-side layout.
    reviewed_geometry = {
        "Rounded Rectangle 3": (658368, 2402699, 10881360, 1133856),
        "TextBox 5": (3730752, 2688342, 4809744, 475488),
        "TextBox 6": (8613648, 2825940, 2560320, 240066),
        "Rectangle 7": (822960, 3911459, 64008, 283464),
        "TextBox 8": (969264, 3884027, 10387584, 289310),
        "Rounded Rectangle 9": (822960, 4432667, 3255264, 1143000),
        "TextBox 10": (987552, 4560683, 2926080, 484632),
        "TextBox 11": (987552, 5091035, 2926080, 212238),
        "Rounded Rectangle 12": (4462272, 4432667, 3255264, 1143000),
        "TextBox 13": (4626864, 4560683, 2926080, 484632),
        "TextBox 14": (4626864, 5091035, 2926080, 310896),
        "Rounded Rectangle 15": (8101583, 4432667, 3255264, 1143000),
        "TextBox 16": (8266175, 4560683, 2926080, 484632),
        "TextBox 17": (8266175, 5091035, 2926080, 310896),
        "Rounded Rectangle 18": (822960, 5880464, 10533888, 603504),
        "Rectangle 19": (1042415, 6142595, 8924544, 182880),
        "TextBox 20": (1042415, 5895707, 2194560, 219456),
        "TextBox 21": (9948672, 5895707, 1133856, 219456),
        "Rounded Rectangle 22": (658368, 1210479, 10881360, 586527),
        "TextBox 23": (1060704, 1330657, 10058400, 424732),
        "Rectangle 25": (809895, 2008624, 64008, 283464),
        "TextBox 26": (956199, 1981192, 10387584, 289310),
    }
    for shape_name, (left, top, width, height) in reviewed_geometry.items():
        shape = next(item for item in slide.shapes if item.name == shape_name)
        shape.left, shape.top = left, top
        shape.width, shape.height = width, height

    add_notes(
        slide,
        "Present the ROSMAP evidence at both stratum and excitatory-neuron levels.",
        f"Before key-driver testing, each eligible mitochondrial-DEG input query was tested for over-representation of the 13 mtDNA-OXPHOS genes against its exact mitochondrial background, with BH correction across four prespecified programs. Across all eligible female epsilon-3 homozygous calls, all {facts['ros_group_occurrences']} mtDNA-OXPHOS occurrences are AD-up, and {facts['ros_group_significant_calls']} input queries are enriched. In excitatory neurons, {facts['ros_exc_eligible_calls']} eligible calls contain {facts['ros_exc_occurrences']} occurrences, all up, and {facts['ros_exc_significant_calls']} of {facts['ros_exc_eligible_calls']} input queries are enriched.",
        "Enrichment describes the composition of a KDA input query; it is not the significance or number of key drivers returned by call_key_drivers. Occurrences repeat genes across calls and are not unique genes, donors, or independent replications.",
        "Ask whether the same program appears in SEA-AD.",
    )
    return slide


def build_seaad_slide(prs: Presentation, facts: dict[str, Any]):
    slide = ui.new_slide(prs)
    add_content_title(
        slide,
        "SEA-AD adds narrow but strong support for the excitatory-neuron program",
        "Only one matched female ε3/ε3 excitatory call was evaluable, so support is strong in direction but limited in coverage.",
    )
    add_stat_card(
        slide,
        "9/10",
        "SEA query genes are shared mtDNA-OXPHOS genes",
        0.78,
        1.50,
        3.58,
        bg=ui.PALE_SKY,
        accent=ui.BLUE,
    )
    add_stat_card(
        slide,
        "9/9",
        "shared genes are AD-up in both cohorts",
        0.78,
        3.07,
        3.58,
        bg=ui.PALE_GREEN,
        accent=ui.TEAL,
    )
    add_stat_card(
        slide,
        "BH = 0.005",
        "cross-cohort overlap after correction",
        0.78,
        4.64,
        3.58,
        bg=ui.PALE_GOLD,
        accent=ui.GOLD,
        detail="9 shared versus 3.83 expected",
    )
    ui.add_rect(slide, 4.72, 1.50, 7.82, 4.74, color=ui.WHITE, outline=ui.LIGHT)
    ui.add_panel_title(
        slide,
        "Nine shared genes",
        5.04,
        1.82,
        7.12,
        accent=ui.BLUE,
    )
    for index, gene in enumerate(facts["sea_shared_mt_genes"]):
        row, column = divmod(index, 3)
        add_chip(
            slide,
            gene,
            5.06 + column * 2.30,
            2.47 + row * 0.88,
            2.02,
            bg=ui.PALE_SKY,
            accent=ui.BLUE,
            size=10.6,
        )
    ui.add_rect(slide, 5.06, 5.28, 6.86, 0.60, color=ui.PALE_GREEN, outline=None)
    ui.add_text(
        slide,
        "Every shared gene points upward in both datasets.",
        5.28,
        5.46,
        6.42,
        0.26,
        size=12.0,
        color=ui.TEAL_TEXT,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    ui.add_text(
        slide,
        "Program-level support is useful even though SEA-AD does not select the same upstream driver.",
        0.94,
        6.53,
        11.46,
        0.28,
        size=11.2,
        color=ui.PURPLE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_notes(
        slide,
        "Show the matched SEA-AD program evidence.",
        "The only evaluable female epsilon-3 homozygous excitatory-neuron call has ten effective query genes. Nine are the same mtDNA-OXPHOS genes found in ROSMAP, and all nine are dementia-up. The overlap is larger than expected under the category-specific background model and remains significant after multiple-testing correction.",
        "This is one SEA-AD call, not broad replication across excitatory subtypes. It validates the direction of the program, not an exact key driver. Exact driver matching is not required in the analysis framework.",
        "Explain why a repeated mitochondrial program matters biologically.",
    )
    return slide


def build_why_slide(prs: Presentation):
    slide = ui.new_slide(prs)
    add_content_title(
        slide,
        "Why this matters: excitatory neurons may be mounting a coordinated mitochondrial response",
        "The repeated direction is biologically informative even though its functional meaning remains open.",
    )
    cards = [
        (
            "1",
            "High energy demand",
            "Excitatory neurons spend substantial energy sending and restoring electrical signals.",
            ui.PALE_SKY,
            ui.BLUE,
        ),
        (
            "2",
            "Essential mtDNA parts",
            "The mitochondrial genome encodes 13 core OXPHOS proteins needed for the respiratory system.",
            ui.PALE_GREEN,
            ui.TEAL,
        ),
        (
            "3",
            "Repeated AD response",
            "The same upward RNA program appears across ROSMAP calls and in the matched SEA-AD comparison.",
            ui.PALE_GOLD,
            ui.GOLD,
        ),
    ]
    for index, (number, title, detail, bg, accent) in enumerate(cards):
        x = 0.72 + index * 4.15
        ui.add_rect(slide, x, 1.58, 3.78, 3.20, color=bg, outline=None)
        ui.add_circle(slide, x + 0.28, 1.92, 0.56, accent)
        ui.add_text(
            slide,
            number,
            x + 0.28,
            2.05,
            0.56,
            0.24,
            size=13.0,
            color=ui.WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        ui.add_text(
            slide,
            title,
            x + 1.00,
            1.94,
            2.46,
            0.38,
            size=16.0,
            color=ui.NAVY,
            bold=True,
            font=ui.FONT_HEAD,
        )
        ui.add_text(
            slide,
            detail,
            x + 0.32,
            2.78,
            3.14,
            1.14,
            size=12.0,
            color=ui.DARK,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
    ui.add_rect(slide, 0.72, 5.18, 11.90, 1.14, color=ui.PALE_RED, outline=None)
    ui.add_text(
        slide,
        "Most cautious interpretation",
        0.98,
        5.46,
        2.40,
        0.30,
        size=12.0,
        color=ui.VERMILION_TEXT,
        bold=True,
    )
    ui.add_text(
        slide,
        "The increase may be compensation or a stress response—not evidence that mitochondria make more ATP.",
        3.55,
        5.38,
        8.57,
        0.46,
        size=14.0,
        color=ui.DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_notes(
        slide,
        "Explain the biological importance without claiming improved function.",
        "Excitatory neurons have high energy requirements, and mtDNA supplies essential respiratory-complex subunits. A repeated upward transcriptional response may reflect an attempt to compensate for inefficient mitochondria or a mitochondrial stress response.",
        "RNA direction alone cannot distinguish compensation, injury, altered RNA processing, or surviving-cell composition.",
        "Separate the observation from what still needs to be tested.",
    )
    return slide


def build_boundaries_slide(prs: Presentation):
    slide = ui.new_slide(prs)
    add_content_title(
        slide,
        "What the finding does not prove—and the next evidence needed",
        "RNA-program evidence generates a testable mechanism; it does not finish the mechanism.",
    )
    ui.add_rect(slide, 0.72, 1.42, 5.78, 4.92, color=ui.PALE_RED, outline=None)
    ui.add_panel_title(
        slide,
        "Not proven",
        1.04,
        1.76,
        5.10,
        accent=ui.VERMILION,
    )
    ui.add_bullets(
        slide,
        [
            "More OXPHOS protein, respiration, or ATP production",
            "A female- or ε3/ε3-specific disease effect",
            "Which upstream key driver caused the RNA increase",
        ],
        1.04,
        2.38,
        5.00,
        size=13.0,
        accent=ui.VERMILION,
        line_h=1.00,
    )
    ui.add_rect(slide, 6.82, 1.42, 5.80, 4.92, color=ui.PALE_GREEN, outline=None)
    ui.add_panel_title(
        slide,
        "Best next tests",
        7.14,
        1.76,
        5.12,
        accent=ui.TEAL,
    )
    ui.add_bullets(
        slide,
        [
            "Donor-aware disease × sex × APOE interaction model",
            "OXPHOS protein abundance and complex assembly",
            "Oxygen consumption, membrane potential, and ATP",
        ],
        7.14,
        2.38,
        5.02,
        size=13.0,
        accent=ui.TEAL,
        line_h=1.00,
    )
    ui.add_text(
        slide,
        "A causal claim requires perturbation and functional rescue—not only another expression dataset.",
        0.92,
        6.62,
        11.50,
        0.30,
        size=11.4,
        color=ui.PURPLE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_notes(
        slide,
        "State the scientific boundaries and next experiments.",
        "The current result measures RNA. Donor-aware interaction models are needed to test sex and APOE specificity. Protein assembly and respiratory assays are needed to determine whether the transcriptional increase improves or impairs mitochondrial function.",
        "SEA-AD support is valuable but optional. It cannot substitute for donor-aware statistics or functional experiments.",
        "Place the result in prior literature.",
    )
    return slide


def build_literature_slide(prs: Presentation):
    slide = ui.new_slide(prs)
    add_content_title(
        slide,
        "Prior research makes the result plausible—but does not duplicate it",
        "Each study supports one part of the interpretation; none establishes this exact female ε3/ε3 excitatory-neuron pattern.",
    )
    cards = [
        (
            "Guo et al. 2023",
            "Sex-specific AD networks",
            "Supports sex-aware key-driver analysis.",
            "Does not test this exact mtDNA program.",
            ui.PALE_SKY,
            ui.BLUE,
        ),
        (
            "Mathys et al. 2024",
            "Cell-resolved AD responses",
            "Supports strong cell-type heterogeneity.",
            "Uses ROSMAP, so it is not independent.",
            ui.PALE_GREEN,
            ui.TEAL,
        ),
        (
            "Lee et al. 2023",
            "APOE and mitochondria",
            "Links APOE4 astrocytes to impaired mitochondrial homeostasis.",
            "Different APOE group and cell type.",
            ui.PALE_GOLD,
            ui.GOLD,
        ),
    ]
    for index, (citation, title, support, limit, bg, accent) in enumerate(cards):
        x = 0.72 + index * 4.15
        ui.add_rect(slide, x, 1.48, 3.78, 4.56, color=bg, outline=None)
        ui.add_text(
            slide,
            citation,
            x + 0.30,
            1.82,
            3.18,
            0.36,
            size=16.0,
            color=ui.readable_accent(accent),
            bold=True,
            align=PP_ALIGN.CENTER,
            font=ui.FONT_HEAD,
        )
        ui.add_text(
            slide,
            title,
            x + 0.30,
            2.44,
            3.18,
            0.46,
            size=14.0,
            color=ui.NAVY,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        ui.add_text(
            slide,
            "SUPPORTS",
            x + 0.32,
            3.20,
            1.10,
            0.24,
            size=9.2,
            color=ui.readable_accent(accent),
            bold=True,
        )
        ui.add_text(
            slide,
            support,
            x + 0.32,
            3.52,
            3.14,
            0.72,
            size=11.2,
            color=ui.DARK,
            align=PP_ALIGN.CENTER,
        )
        ui.add_text(
            slide,
            "BOUNDARY",
            x + 0.32,
            4.48,
            1.10,
            0.24,
            size=9.2,
            color=ui.GRAY,
            bold=True,
        )
        ui.add_text(
            slide,
            limit,
            x + 0.32,
            4.80,
            3.14,
            0.72,
            size=10.8,
            color=ui.GRAY,
            align=PP_ALIGN.CENTER,
        )
    ui.add_rect(slide, 0.72, 6.28, 11.90, 0.58, color=ui.WHITE, outline=ui.LIGHT)
    ui.add_text(
        slide,
        "Novelty assessment: moderate—known mitochondrial biology in a less-established sex/APOE/cell context.",
        0.98,
        6.45,
        11.38,
        0.26,
        size=11.0,
        color=ui.PURPLE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_notes(
        slide,
        "Connect the finding to prior literature while preserving novelty boundaries.",
        "Guo and colleagues support sex-aware Alzheimer’s key-driver analysis. Mathys and colleagues show cell-resolved, multiregion disease responses but use ROSMAP, so that paper is supporting discovery-resource context. Lee and colleagues show an APOE–mitochondria link in human APOE4 astrocytes, a different genotype and cell type.",
        "None of these studies independently establishes the exact female epsilon-3 homozygous excitatory-neuron mtDNA-up pattern.",
        "Return to the broader finding list or proceed to sensitivity analyses.",
    )
    return slide


def build_slides(prs: Presentation, facts: dict[str, Any]) -> list[Any]:
    return [
        build_divider(prs),
        build_statement_slide(prs),
        build_rosmap_slide(prs, facts),
        build_seaad_slide(prs, facts),
        build_why_slide(prs),
        build_boundaries_slide(prs),
        build_literature_slide(prs),
    ]


def write_audit(
    audit_root: Path,
    checks: list[dict[str, Any]],
    facts: dict[str, Any],
    output_path: Path,
    backup_path: Path,
) -> None:
    audit_root.mkdir(parents=True, exist_ok=True)
    with (audit_root / "finding1_slide_checks.tsv").open(
        "w", encoding="utf-8", newline=""
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["check_id", "observed", "expected", "passed"],
            delimiter="\t",
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(checks)
    with (audit_root / "finding1_slide_facts.tsv").open(
        "w", encoding="utf-8", newline=""
    ) as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(["field", "value"])
        for key, value in facts.items():
            if isinstance(value, list):
                value = ";".join(map(str, value))
            writer.writerow([key, value])
        writer.writerow(["output_deck", output_path.resolve()])
        writer.writerow(["output_sha256", sha256_file(output_path)])
        writer.writerow(["backup_deck", backup_path.resolve()])


def main() -> int:
    args = parse_args()
    input_path = args.input.resolve()
    output_path = args.output.resolve()
    audit_root = args.audit_root.resolve()
    if not input_path.is_file():
        raise FileNotFoundError(input_path)

    facts = prepare_facts()
    prs = Presentation(input_path)
    if len(prs.slides) != EXPECTED_INPUT_SLIDES:
        raise RuntimeError(
            f"Expected {EXPECTED_INPUT_SLIDES} slides, found {len(prs.slides)}"
        )
    if SECTION_TITLE in "\n".join(slide_text(slide) for slide in prs.slides):
        raise RuntimeError("Finding 1 slides already exist; refusing to insert twice")

    sync_reviewed_base_slides(prs)
    before = [semantic_slide_fingerprint(slide) for slide in prs.slides]
    input_hash = sha256_file(input_path)
    backup_dir = audit_root / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / f"input_deck_{input_hash[:12]}.pptx"
    if not backup_path.exists():
        shutil.copy2(input_path, backup_path)

    ui.set_notes_body_template(
        prs.slides[0].notes_slide.notes_placeholder._element
    )
    new_slides = build_slides(prs, facts)
    slide_ids = prs.slides._sldIdLst
    new_ids = list(slide_ids[-NEW_SLIDE_COUNT:])
    for slide_id in new_ids:
        slide_ids.remove(slide_id)
    for offset, slide_id in enumerate(new_ids):
        slide_ids.insert(INSERT_INDEX + offset, slide_id)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}.",
        suffix=".pptx",
        dir=output_path.parent,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        prs.save(temporary)
        reloaded = Presentation(temporary)
        after = [semantic_slide_fingerprint(slide) for slide in reloaded.slides]
        inserted = [
            reloaded.slides[index]
            for index in range(INSERT_INDEX, INSERT_INDEX + NEW_SLIDE_COUNT)
        ]
        inserted_text = [slide_text(slide) for slide in inserted]
        previous_slides_unchanged = (
            after[:INSERT_INDEX] == before[:INSERT_INDEX]
            and after[INSERT_INDEX + NEW_SLIDE_COUNT :] == before[INSERT_INDEX:]
        )
        all_new_have_notes = all(
            slide.notes_slide.notes_text_frame
            and slide.notes_slide.notes_text_frame.text.strip()
            for slide in inserted
        )
        all_shapes_in_bounds = all(
            0 <= shape.left
            and 0 <= shape.top
            and shape.left + shape.width <= reloaded.slide_width
            and shape.top + shape.height <= reloaded.slide_height
            for slide in inserted
            for shape in slide.shapes
        )
        required_text = [
            SECTION_MARKER,
            "SEA-AD can support the same biological pattern",
            "Plain meaning",
            "(especially excitatory neurons)",
            "123",
            "KDA input queries enriched",
            "Enriched = more mtDNA-OXPHOS genes than expected by chance",
            "9/10",
            "Why this matters",
            "Not proven",
            "Guo et al. 2023",
        ]
        combined = "\n".join(inserted_text)
        checks = [
            {
                "check_id": "output_slide_count",
                "observed": len(reloaded.slides),
                "expected": EXPECTED_INPUT_SLIDES + NEW_SLIDE_COUNT,
                "passed": len(reloaded.slides)
                == EXPECTED_INPUT_SLIDES + NEW_SLIDE_COUNT,
            },
            {
                "check_id": "finding1_block_appended_after_sensitivity",
                "observed": SECTION_MARKER in inserted_text[0]
                and "Query-size sensitivity — validation interpretation"
                in slide_text(reloaded.slides[INSERT_INDEX - 1]),
                "expected": True,
                "passed": SECTION_MARKER in inserted_text[0]
                and "Query-size sensitivity — validation interpretation"
                in slide_text(reloaded.slides[INSERT_INDEX - 1]),
            },
            {
                "check_id": "all_required_content_present",
                "observed": all(text in combined for text in required_text),
                "expected": True,
                "passed": all(text in combined for text in required_text),
            },
            {
                "check_id": "all_new_slides_have_notes",
                "observed": all_new_have_notes,
                "expected": True,
                "passed": all_new_have_notes,
            },
            {
                "check_id": "all_new_shapes_in_slide_bounds",
                "observed": all_shapes_in_bounds,
                "expected": True,
                "passed": all_shapes_in_bounds,
            },
            {
                "check_id": "all_previous_slides_semantically_unchanged",
                "observed": previous_slides_unchanged,
                "expected": True,
                "passed": previous_slides_unchanged,
            },
        ]
        failed = [row["check_id"] for row in checks if not row["passed"]]
        if failed:
            raise RuntimeError("Finding 1 slide checks failed: " + ", ".join(failed))
        os.replace(temporary, output_path)
    finally:
        if temporary.exists():
            temporary.unlink()

    write_audit(audit_root, checks, facts, output_path, backup_path)
    print(f"Updated deck: {output_path}")
    print(
        f"Inserted slides: {INSERT_INDEX + 1}–"
        f"{INSERT_INDEX + NEW_SLIDE_COUNT}"
    )
    print(f"Backup: {backup_path}")
    print(f"Audit: {audit_root}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
