#!/usr/bin/env python3
"""Insert an all-gene ROSMAP fine-cell DEG summary after slide 3.

The slide summarizes Phase 08 DEG rows across the six sex/APOE
strata. Phase 11's validated contrast manifest supplies the planned and
estimable comparison counts. A count is a gene-by-fine-cell comparison
occurrence, so the same gene can contribute in more than one comparison.
"""

from __future__ import annotations

import argparse
import glob
import hashlib
import os
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Any

os.environ.setdefault(
    "MPLCONFIGDIR", str(Path(tempfile.gettempdir()) / "deg_summary_mpl")
)
os.environ.setdefault("MPLBACKEND", "Agg")

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from matplotlib.patches import Patch  # noqa: E402
from matplotlib.ticker import FuncFormatter  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402


ROOT = Path(__file__).resolve().parents[2]
PRESENTATION_SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(PRESENTATION_SCRIPTS))
import update_phase11_seaad_simple_aggr_part2 as ui  # noqa: E402

SKILL_SCRIPTS = ROOT / ".agents/skills/scientific-visualization/scripts"
sys.path.insert(0, str(SKILL_SCRIPTS))
from figure_export import save_publication_figure  # noqa: E402
from style_presets import apply_publication_style  # noqa: E402


DEFAULT_DECK = (
    ROOT
    / "docs"
    / "presentations"
    / "09162026"
    / "09162026_sex_apoe_kda_fine_broad.pptx"
)
AUDIT_ROOT = ROOT / "results" / "presentations" / "09162026_deg_summary"
PHASE08_DIR = ROOT / "results" / "minerva_production" / "08_deg_fine"
CONTRAST_MANIFEST = (
    ROOT
    / "results"
    / "minerva_production"
    / "11_pathway_deg_fine"
    / "fine_deg_pathway_contrast_manifest.tsv"
)

INPUT_SLIDES = 49
OUTPUT_SLIDES = 50
UPDATED_OUTPUT_SLIDES = {50, 52, 157, 158, 159}
INSERT_INDEX = 3
NEW_TITLE = "Overall DEG counts vary across sex/APOE groups"
NEXT_TITLE = "Pathway analysis"
GROUP_ORDER = ["F_e2", "F_e33", "F_e4", "M_e2", "M_e33", "M_e4"]
GROUP_LABELS = {
    "F_e2": "Female ε2",
    "F_e33": "Female ε3/ε3",
    "F_e4": "Female ε4",
    "M_e2": "Male ε2",
    "M_e33": "Male ε3/ε3",
    "M_e4": "Male ε4",
}
EXPECTED_DIRECTION_COUNTS = {
    "F_e2": (7136, 3816),
    "F_e33": (5240, 5996),
    "F_e4": (5038, 15904),
    "M_e2": (25545, 20236),
    "M_e33": (6768, 6352),
    "M_e4": (8385, 7881),
}
EXPECTED_SUMMARY = {
    "planned": 324,
    "estimable": 321,
    "contrasts_with_degs": 277,
    "deg_occurrences": 118297,
    "distinct_genes": 14840,
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_DECK)
    parser.add_argument("--output", type=Path, default=DEFAULT_DECK)
    parser.add_argument("--audit-root", type=Path, default=AUDIT_ROOT)
    parser.add_argument(
        "--figure-only",
        action="store_true",
        help="Render and validate figure artifacts without modifying the deck.",
    )
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def truth(series: pd.Series) -> pd.Series:
    return series.astype(str).str.upper().isin({"TRUE", "T", "1", "YES"})


def slide_title(slide) -> str:
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text.strip():
            return " ".join(shape.text.split())
    return ""


def slide_fingerprint(slide) -> tuple[Any, ...]:
    notes = slide.notes_slide.notes_text_frame
    return (
        tuple(
            (
                int(shape.shape_type),
                shape.name,
                int(shape.left),
                int(shape.top),
                int(shape.width),
                int(shape.height),
                shape.text if getattr(shape, "has_text_frame", False) else "",
            )
            for shape in slide.shapes
        ),
        notes.text if notes is not None else "",
    )


def prepare_statistics() -> tuple[pd.DataFrame, dict[str, int]]:
    files = sorted(glob.glob(str(PHASE08_DIR / "*.yu_mast_de.tsv.gz")))
    if len(files) != 9:
        raise RuntimeError(f"Expected 9 Phase 08 DEG files, found {len(files)}")
    if not CONTRAST_MANIFEST.is_file():
        raise FileNotFoundError(CONTRAST_MANIFEST)

    columns = ["contrast_id", "sex", "apoe_group", "gene", "logFC", "paper_deg"]
    significant_frames: list[pd.DataFrame] = []
    for path in files:
        frame = pd.read_csv(path, sep="\t", usecols=columns, low_memory=False)
        significant_frames.append(frame.loc[truth(frame["paper_deg"])].copy())
    degs = pd.concat(significant_frames, ignore_index=True)
    degs["signature_group"] = degs["sex"].str[0] + "_" + degs["apoe_group"]
    degs["direction"] = np.where(degs["logFC"] > 0, "higher_in_ad", "lower_in_ad")

    counts = (
        degs.groupby(["signature_group", "direction"])
        .size()
        .unstack(fill_value=0)
        .reindex(GROUP_ORDER)
    )
    counts = counts.rename(
        columns={"higher_in_ad": "up_count", "lower_in_ad": "down_count"}
    ).reset_index()
    counts["group_label"] = counts["signature_group"].map(GROUP_LABELS)
    counts["total_count"] = counts["up_count"] + counts["down_count"]
    counts["up_percent"] = 100.0 * counts["up_count"] / counts["total_count"]
    counts["down_percent"] = 100.0 - counts["up_percent"]
    counts["distinct_genes"] = counts["signature_group"].map(
        degs.groupby("signature_group")["gene"].nunique()
    )
    counts["contrasts_with_degs"] = counts["signature_group"].map(
        degs.groupby("signature_group")["contrast_id"].nunique()
    )

    manifest = pd.read_csv(CONTRAST_MANIFEST, sep="\t", low_memory=False)
    summary = {
        "planned": int(len(manifest)),
        "estimable": int(truth(manifest["structurally_available"]).sum()),
        "contrasts_with_degs": int(degs["contrast_id"].nunique()),
        "deg_occurrences": int(len(degs)),
        "distinct_genes": int(degs["gene"].nunique()),
    }
    observed_counts = {
        row.signature_group: (int(row.up_count), int(row.down_count))
        for row in counts.itertuples()
    }
    if observed_counts != EXPECTED_DIRECTION_COUNTS:
        raise RuntimeError(f"Directional DEG counts changed: {observed_counts!r}")
    if summary != EXPECTED_SUMMARY:
        raise RuntimeError(f"Overall DEG statistics changed: {summary!r}")

    manifest_totals = (
        manifest.groupby("signature_group")["paper_degs_phase08"]
        .sum()
        .reindex(GROUP_ORDER)
        .astype(int)
    )
    if manifest_totals.to_dict() != counts.set_index("signature_group")[
        "total_count"
    ].astype(int).to_dict():
        raise RuntimeError("Phase 08 DEG rows do not match the Phase 11 manifest")
    return counts, summary


def render_figure(frame: pd.DataFrame, output_dir: Path) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    apply_publication_style("presentation")
    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"],
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "axes.linewidth": 0.8,
        }
    )

    fig, ax = plt.subplots(figsize=(12.0, 3.65), constrained_layout=False)
    fig.subplots_adjust(left=0.135, right=0.955, top=0.80, bottom=0.18)
    positions = np.arange(len(frame))
    up = frame["up_count"].to_numpy(dtype=float)
    down = frame["down_count"].to_numpy(dtype=float)
    totals = frame["total_count"].to_numpy(dtype=float)
    up_color = "#0072B2"
    down_color = "#D55E00"

    up_bars = ax.barh(
        positions,
        up,
        height=0.64,
        color=up_color,
        edgecolor="white",
        linewidth=0.9,
        zorder=3,
    )
    down_bars = ax.barh(
        positions,
        down,
        left=up,
        height=0.64,
        color=down_color,
        edgecolor="#6B3418",
        linewidth=0.8,
        hatch="///",
        zorder=3,
    )
    for index, (up_bar, down_bar, total) in enumerate(
        zip(up_bars, down_bars, totals, strict=True)
    ):
        ax.text(
            up_bar.get_x() + up_bar.get_width() / 2,
            up_bar.get_y() + up_bar.get_height() / 2,
            f"{int(up[index]):,}",
            ha="center",
            va="center",
            color="white",
            fontsize=11.0,
            fontweight="bold",
            zorder=4,
        )
        ax.text(
            down_bar.get_x() + down_bar.get_width() / 2,
            down_bar.get_y() + down_bar.get_height() / 2,
            f"{int(down[index]):,}",
            ha="center",
            va="center",
            color="white",
            fontsize=11.0,
            fontweight="bold",
            zorder=4,
        )
        ax.text(
            total + 900,
            positions[index],
            f"{int(total):,} total",
            ha="left",
            va="center",
            color="#1F2937",
            fontsize=10.8,
            fontweight="bold",
        )

    ax.set_yticks(positions)
    ax.set_yticklabels(frame["group_label"], fontsize=12.0, fontweight="bold")
    for index, tick in enumerate(ax.get_yticklabels()):
        if index in {1, 4}:
            tick.set_color("#5B3A87")
    ax.invert_yaxis()
    ax.axhline(2.5, color="#D1D5DB", linewidth=1.1, zorder=1)
    ax.set_xlim(0, 52500)
    ax.set_xlabel("DEG occurrences across fine-cell comparisons", fontsize=12.0)
    ax.xaxis.set_major_formatter(
        FuncFormatter(lambda value, _position: "0" if value == 0 else f"{int(value / 1000)}k")
    )
    ax.tick_params(axis="x", labelsize=10.5, length=0, pad=5)
    ax.tick_params(axis="y", length=0, pad=8)
    ax.grid(axis="x", color="#E5E7EB", linewidth=0.8, zorder=0)
    for spine in ax.spines.values():
        spine.set_visible(False)
    legend = ax.legend(
        handles=[
            Patch(
                facecolor=up_color,
                edgecolor="white",
                label="Up-regulated in AD",
            ),
            Patch(
                facecolor=down_color,
                edgecolor="#6B3418",
                hatch="///",
                label="Down-regulated in AD",
            ),
        ],
        loc="lower center",
        bbox_to_anchor=(0.5, 1.01),
        ncol=2,
        frameon=False,
        fontsize=11.5,
        handlelength=1.8,
        columnspacing=2.8,
    )
    for item in legend.get_texts():
        item.set_color("#1F2937")

    base = output_dir / "rosmap_all_deg_counts_by_sex_apoe"
    saved = save_publication_figure(
        fig,
        base,
        formats=["png", "svg", "pdf"],
        dpi=320,
        bbox_inches="tight",
        pad_inches=0.08,
    )
    plt.close(fig)
    artifacts = {path.suffix.lstrip("."): path for path in saved}
    if set(artifacts) != {"png", "svg", "pdf"}:
        raise RuntimeError(f"Missing figure outputs: {artifacts}")
    frame.to_csv(base.with_name(base.name + "_statistics.tsv"), sep="\t", index=False)
    return artifacts


def add_metric_card(
    slide,
    *,
    x: float,
    value: str,
    label: str,
    detail: str = "",
) -> None:
    ui.add_rect(slide, x, 1.18, 3.72, 0.82, color=ui.WHITE, outline=ui.LIGHT)
    ui.add_text(
        slide,
        value,
        x + 0.18,
        1.35,
        1.28,
        0.38,
        size=23.0,
        color=ui.BLUE,
        bold=True,
        font=ui.FONT_HEAD,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    ui.add_text(
        slide,
        label,
        x + 1.52,
        1.31,
        2.00,
        0.26,
        size=11.0,
        color=ui.NAVY,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    if detail:
        ui.add_text(
            slide,
            detail,
            x + 1.52,
            1.60,
            2.00,
            0.22,
            size=9.2,
            color=ui.GRAY,
            valign=MSO_ANCHOR.MIDDLE,
        )


def clear_slide(slide) -> None:
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = ui.OFF_WHITE


def build_slide(slide, figure_path: Path) -> None:
    ui.add_title_block(
        slide,
        NEW_TITLE,
        "All genes; each count is one significant gene in one fine-cell AD-versus-NCI contrast.",
    )
    subtitle = next(shape for shape in slide.shapes if shape.name == "TextBox 2")
    subtitle.top = 861567
    subtitle.height = 227755
    add_metric_card(
        slide,
        x=0.72,
        value="321",
        label="of 324 planned contrasts have enough cells.",
    )
    add_metric_card(
        slide,
        x=4.81,
        value="277",
        label="of 321 estimable comparisons had at least one DEG",
        detail=" ",
    )
    add_metric_card(
        slide,
        x=8.90,
        value="118,297",
        label="DEG occurrences",
        detail="14,840 distinct genes represented",
    )
    # Match the reviewed slide: retain the original shape-number sequence, then
    # remove the no-longer-displayed percentage line and apply the manual card
    # positions recorded in the deck.
    for name in ("Rounded Rectangle 3", "Rounded Rectangle 6", "Rounded Rectangle 10"):
        next(shape for shape in slide.shapes if shape.name == name).top = 1155193
    for name in ("TextBox 4", "TextBox 7", "TextBox 11"):
        next(shape for shape in slide.shapes if shape.name == name).top = 1310641
    first_label = next(shape for shape in slide.shapes if shape.name == "TextBox 5")
    first_label.top = 1314491
    first_label.height = 393954
    next(shape for shape in slide.shapes if shape.name == "TextBox 8").top = 1358734
    next(shape for shape in slide.shapes if shape.name == "TextBox 12").top = 1274065
    next(shape for shape in slide.shapes if shape.name == "TextBox 13").top = 1539241
    removed_detail = next(
        shape for shape in slide.shapes if shape.name == "TextBox 9"
    )
    removed_detail._element.getparent().remove(removed_detail._element)
    ui.add_picture_contain(
        slide,
        figure_path,
        0.72,
        2.10,
        11.90,
        3.76,
        alt=(
            "Stacked horizontal bar chart of all significant ROSMAP fine-cell DEG "
            "occurrences across six sex/APOE groups. Blue segments show genes "
            "up-regulated in Alzheimer’s disease and hatched orange segments show "
            "genes down-regulated in Alzheimer’s disease."
        ),
    )
    ui.add_rect(slide, 0.72, 6.12, 11.90, 0.72, color=ui.PALE_GRAY, outline=None)
    ui.add_text(
        slide,
        "Key patterns",
        1.01,
        6.34,
        1.30,
        0.24,
        size=12.0,
        color=ui.PURPLE,
        bold=True,
    )
    key_pattern = ui.add_text(
        slide,
        "Male ε2 had the largest DEG count; female ε4 was strongly dominated by down-regulated genes (76%).",
        2.42,
        6.29,
        9.76,
        0.34,
        size=13.4,
        color=ui.DARK,
        valign=MSO_ANCHOR.MIDDLE,
    )
    key_pattern.top = 5776219
    key_pattern.height = 261610
    ui.add_notes(
        slide,
        goal="Summarize the complete fine-cell DEG results before pathway analysis.",
        walkthrough=(
            "ROSMAP contained 324 planned AD-versus-NCI comparisons: 54 fine cell "
            "types crossed with six sex/APOE groups. Three male epsilon-2 comparisons "
            "were not estimable, leaving 321 analyzed comparisons; 277 of those had "
            "at least one DEG. Across all genes, there were 118,297 significant "
            "gene-by-fine-cell comparison occurrences representing 14,840 distinct "
            "genes. The bars divide these occurrences into higher expression in AD "
            "and lower expression in AD. Female epsilon-2 had 10,952 occurrences, "
            "65 percent higher in "
            "AD. Female epsilon-3 homozygous had 11,236, 47 percent higher. Female "
            "epsilon-4 had 20,942, with 76 percent lower in AD. Male epsilon-2 had the "
            "largest count, 45,781, with 56 percent higher. Male epsilon-3 homozygous "
            "had 13,120, and male epsilon-4 had 16,266; both were close to evenly split."
        ),
        boundary=(
            "A DEG met the prespecified within-contrast criteria: BH FDR below 0.05, "
            "absolute fold change above 1.3, and detection in at least 10 percent of "
            "cells in either group. Counts repeat a gene when it is significant in "
            "another fine-cell comparison. They are descriptive and can reflect "
            "sample size, power, and the number of genes tested; they are not donor "
            "counts, independent replications, or a formal disease-by-sex/APOE "
            "interaction test."
        ),
        transition=(
            "Now move from individual DEG counts to pathway analysis, which asks "
            "whether those changes concentrate in predefined biological programs."
        ),
    )


def validate_deck(prs: Presentation, preserved: tuple[Any, ...]) -> None:
    expected_slides = len(preserved) + 1
    if len(prs.slides) != expected_slides:
        raise RuntimeError(
            f"Expected {expected_slides} slides, found {len(prs.slides)}"
        )
    if slide_title(prs.slides[INSERT_INDEX]) != NEW_TITLE:
        raise RuntimeError("The all-DEG summary is not slide 4")
    if slide_title(prs.slides[INSERT_INDEX + 1]) != NEXT_TITLE:
        raise RuntimeError("Pathway analysis does not immediately follow the DEG summary")
    retained = tuple(
        slide_fingerprint(slide)
        for index, slide in enumerate(prs.slides)
        if index != INSERT_INDEX
    )
    if retained != preserved:
        raise RuntimeError("A retained slide changed during DEG-summary insertion")
    slide = prs.slides[INSERT_INDEX]
    visible = "\n".join(
        shape.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    for required in (
        "of 324 planned contrasts have enough cells",
        "of 321",
        "118,297",
        "14,840",
    ):
        if required not in visible:
            raise RuntimeError(f"Missing DEG statistic on slide 4: {required}")
    if "86% of estimable comparisons" in visible:
        raise RuntimeError("Removed percentage line is still visible on slide 4")
    if "down-regulated genes (76%)" not in visible:
        raise RuntimeError("Reviewed DEG direction summary is missing from slide 4")
    notes = slide.notes_slide.notes_text_frame
    if notes is None or "BH FDR below 0.05" not in notes.text:
        raise RuntimeError("DEG criteria are missing from the slide 4 notes")
    for shape in slide.shapes:
        if (
            shape.left < 0
            or shape.top < 0
            or shape.left + shape.width > prs.slide_width + 10
            or shape.top + shape.height > prs.slide_height + 10
        ):
            raise RuntimeError(f"Out-of-bounds shape on slide 4: {shape.name}")


def main() -> int:
    args = parse_args()
    input_path = args.input.resolve()
    output_path = args.output.resolve()
    audit_root = args.audit_root.resolve()
    counts, summary = prepare_statistics()
    figures = render_figure(counts, audit_root / "figures")
    pd.DataFrame([summary]).to_csv(
        audit_root / "figures" / "rosmap_all_deg_overall_statistics.tsv",
        sep="\t",
        index=False,
    )
    if args.figure_only:
        print(f"Rendered figure: {figures['png']}")
        return 0
    if not input_path.is_file():
        raise FileNotFoundError(input_path)

    input_hash = sha256(input_path)
    prs = Presentation(input_path)
    if slide_title(prs.slides[INSERT_INDEX - 1]) != "What is a contrast":
        raise RuntimeError("Slide 3 is not the expected contrast slide")

    if (
        len(prs.slides) == INPUT_SLIDES
        and slide_title(prs.slides[INSERT_INDEX]) == NEXT_TITLE
        and NEW_TITLE not in [slide_title(slide) for slide in prs.slides]
    ):
        preserved = tuple(slide_fingerprint(slide) for slide in prs.slides)
        ui.set_notes_body_template(
            prs.slides[0].notes_slide.notes_placeholder._element
        )
        new_slide = ui.new_slide(prs)
        build_slide(new_slide, figures["png"])
        slide_ids = prs.slides._sldIdLst
        new_id = slide_ids[-1]
        slide_ids.remove(new_id)
        slide_ids.insert(INSERT_INDEX, new_id)
        action = "Inserted"
    elif (
        len(prs.slides) in UPDATED_OUTPUT_SLIDES
        and slide_title(prs.slides[INSERT_INDEX]) == NEW_TITLE
        and slide_title(prs.slides[INSERT_INDEX + 1]) == NEXT_TITLE
    ):
        preserved = tuple(
            slide_fingerprint(slide)
            for index, slide in enumerate(prs.slides)
            if index != INSERT_INDEX
        )
        clear_slide(prs.slides[INSERT_INDEX])
        build_slide(prs.slides[INSERT_INDEX], figures["png"])
        action = "Updated"
    else:
        raise RuntimeError(
            f"Unexpected deck state: {len(prs.slides)} slides; "
            f"slide 4 is {slide_title(prs.slides[INSERT_INDEX])!r}"
        )
    validate_deck(prs, preserved)

    audit_root.mkdir(parents=True, exist_ok=True)
    backup_dir = audit_root / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / f"input_deck_{input_hash[:12]}.pptx"
    if not backup_path.exists():
        shutil.copy2(input_path, backup_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    handle, temp_name = tempfile.mkstemp(
        prefix=output_path.stem + ".", suffix=".pptx", dir=output_path.parent
    )
    os.close(handle)
    temp_path = Path(temp_name)
    try:
        prs.save(temp_path)
        reopened = Presentation(temp_path)
        validate_deck(reopened, preserved)
        os.replace(temp_path, output_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()

    print(f"{action} slide 4: {output_path}")
    print(f"Figure: {figures['png']}")
    print(f"Backup: {backup_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
