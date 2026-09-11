#!/usr/bin/env python3
"""Append a sensitivity-analysis section to the 2026-09-16 deck.

The target deck is edited in place by default. Existing slides, geometry, and
speaker notes are preserved; a Part 4 divider, one comparison-table slide,
and one findings slide are appended. All displayed values are recomputed from
the frozen Phase 22 and VH14 manifests and significant-return tables before
the deck is written.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
import update_phase11_seaad_simple_aggr_part2 as ui  # noqa: E402


DEFAULT_DECK = (
    ROOT
    / "docs/presentations/09162026/sex_apoe_kda_fine_broad_09162026.pptx"
)
SOURCE_ANALYSIS = (
    ROOT
    / "docs/validation_human/harmonized_broad_kda/"
    "donor_per_arm_3_vs_5_sensitivity.md"
)
ROS_DIR = ROOT / "results/minerva_production/22_harmonized_broad_kda_combo"
SEA_DIR = ROOT / "results/validation_human/14_harmonized_broad_kda_combo"
AUDIT_DIR = ROOT / "results/presentations/09162026_broad_kda_sensitivity"

EXPECTED_INPUT_SLIDES = 24
EXPECTED_OUTPUT_SLIDES = 27
DIVIDER_TITLE = "Sensitivity analyses"
TABLE_TITLE = "Donor-threshold sensitivity — numerical comparison"
FINDINGS_TITLE = "Donor-threshold sensitivity — key findings"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_DECK)
    parser.add_argument("--output", type=Path, default=DEFAULT_DECK)
    parser.add_argument("--audit-dir", type=Path, default=AUDIT_DIR)
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_tsv(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        raise FileNotFoundError(path)
    with path.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def write_tsv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=fields, delimiter="\t", lineterminator="\n"
        )
        writer.writeheader()
        writer.writerows(rows)


def truth(value: str) -> bool:
    return value.strip().upper() in {"TRUE", "T", "1", "YES"}


def slide_text(slide) -> str:
    return "\n".join(
        shape.text for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )


def slide_signature(slide) -> tuple[Any, ...]:
    shapes = tuple(
        (
            shape.name,
            str(shape.shape_type),
            int(shape.left),
            int(shape.top),
            int(shape.width),
            int(shape.height),
            shape.text if getattr(shape, "has_text_frame", False) else "",
        )
        for shape in slide.shapes
    )
    notes_frame = slide.notes_slide.notes_text_frame
    notes = notes_frame.text if notes_frame is not None else ""
    return shapes, notes


def load_cohort_facts(cohort: str, directory: Path) -> dict[int, dict[str, int]]:
    status = read_tsv(directory / "status.tsv")
    if (
        len(status) != 1
        or status[0]["validation_status"] != "validated_complete"
        or int(status[0]["failed_checks"]) != 0
    ):
        raise RuntimeError(f"{cohort} KDA release is not validated complete")

    manifest = read_tsv(directory / "run_manifest.tsv")
    returns = read_tsv(directory / "significant_returns.tsv")
    if len(manifest) != 42:
        raise RuntimeError(f"Expected 42 {cohort} structural slots")

    output: dict[int, dict[str, int]] = {}
    for threshold in (3, 5):
        eligible = {
            row["kda_run_id"]
            for row in manifest
            if row["source_terminal_status"] == "completed"
            and int(row["n_case_donors"]) >= threshold
            and int(row["n_reference_donors"]) >= threshold
        }
        executed = {
            row["kda_run_id"]
            for row in manifest
            if truth(row["execute_kda"])
            and int(row["n_case_donors"]) >= threshold
            and int(row["n_reference_donors"]) >= threshold
        }
        retained = [row for row in returns if row["kda_run_id"] in executed]
        genes = {row["key_driver"] for row in retained}
        non_mt = {gene for gene in genes if not gene.startswith("MT-")}
        global_rows = [row for row in retained if truth(row["global_key_driver"])]
        global_genes = {row["key_driver"] for row in global_rows}
        global_non_mt = {gene for gene in global_genes if not gene.startswith("MT-")}
        output[threshold] = {
            "structural_contrasts": len(manifest),
            "eligible_contrasts": len(eligible),
            "kda_runs": len(executed),
            "driver_rows": len(retained),
            "unique_raw_drivers": len(genes),
            "unique_non_mt_drivers": len(non_mt),
            "global_driver_rows": len(global_rows),
            "unique_global_drivers": len(global_genes),
            "unique_global_non_mt_drivers": len(global_non_mt),
        }
    return output


def load_facts() -> dict[str, dict[int, dict[str, int]]]:
    required = [
        DEFAULT_DECK,
        SOURCE_ANALYSIS,
        ROS_DIR / "status.tsv",
        ROS_DIR / "run_manifest.tsv",
        ROS_DIR / "significant_returns.tsv",
        SEA_DIR / "status.tsv",
        SEA_DIR / "run_manifest.tsv",
        SEA_DIR / "significant_returns.tsv",
    ]
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing slide input(s): " + ", ".join(missing))

    facts = {
        "rosmap": load_cohort_facts("ROSMAP", ROS_DIR),
        "seaad": load_cohort_facts("SEA-AD", SEA_DIR),
    }
    expected = {
        "rosmap": {
            3: (42, 3, 39, 39, 33, 14, 14, 13),
            5: (40, 2, 13, 13, 13, 9, 9, 9),
        },
        "seaad": {
            3: (28, 5, 24, 21, 18, 14, 12, 9),
            5: (20, 5, 24, 21, 18, 14, 12, 9),
        },
    }
    fields = (
        "eligible_contrasts",
        "kda_runs",
        "driver_rows",
        "unique_raw_drivers",
        "unique_non_mt_drivers",
        "global_driver_rows",
        "unique_global_drivers",
        "unique_global_non_mt_drivers",
    )
    for cohort, thresholds in expected.items():
        for threshold, expected_values in thresholds.items():
            observed = tuple(facts[cohort][threshold][field] for field in fields)
            if observed != expected_values:
                raise RuntimeError(
                    f"Sensitivity source drift for {cohort} threshold "
                    f"{threshold}: {observed} != {expected_values}"
                )
    return facts


def append_divider(prs: Presentation) -> None:
    slide = ui.new_slide(prs, bg=ui.NAVY)
    ui.add_text(
        slide, "PART 4", 0.78, 0.67, 2.4, 0.28,
        size=10.5, color=ui.GOLD, bold=True,
    )
    ui.add_rect(
        slide, 0.78, 1.26, 0.10, 2.38,
        color=ui.GOLD, outline=None, radius=False,
    )
    ui.add_text(
        slide, DIVIDER_TITLE,
        1.18, 1.36, 8.20, 1.05,
        size=32, color=ui.WHITE, bold=True,
        font=ui.FONT_HEAD, valign=MSO_ANCHOR.MIDDLE,
    )
    ui.add_text(
        slide,
        "Robustness checks across sample support, profile coverage, "
        "query definition, and KDA decision rules.",
        1.20, 2.70, 8.10, 1.00,
        size=15.0, color=ui.RGBColor(204, 219, 234),
    )
    ui.add_rect(
        slide, 9.78, 1.28, 2.70, 4.70,
        color=ui.NAVY_2, outline=None,
    )
    ui.add_text(
        slide, "IN THIS PART", 10.08, 1.67, 2.10, 0.24,
        size=9.4, color=ui.GOLD, bold=True,
    )
    topics = ["Donor support", "Profile coverage", "Query definition", "KDA filters"]
    for index, topic in enumerate(topics, start=1):
        y = 2.20 + (index - 1) * 0.82
        ui.add_circle(slide, 10.06, y + 0.03, 0.34, ui.GOLD)
        ui.add_text(
            slide, str(index), 10.06, y + 0.08, 0.34, 0.17,
            size=8.2, color=ui.NAVY, bold=True,
            align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE,
        )
        ui.add_text(
            slide, topic, 10.53, y, 1.58, 0.52,
            size=11.1, color=ui.WHITE, bold=True,
        )
    ui.add_rect(
        slide, 1.18, 4.48, 7.95, 0.72,
        color=ui.NAVY_2, outline=ui.GOLD,
    )
    ui.add_text(
        slide,
        "Each sensitivity test is reported separately with its own evidence boundary.",
        1.48, 4.67, 7.35, 0.30,
        size=11.2, color=ui.WHITE, bold=True, align=PP_ALIGN.CENTER,
    )
    ui.add_notes(
        slide,
        goal="Introduce the section reserved for all harmonized broad-cell KDA sensitivity analyses.",
        walkthrough="Sensitivity checks may address donor support, nuclei-per-profile coverage, DEG query definition, and KDA filtering. The first check raises the disease-arm donor gate from three to five in both cohorts.",
        boundary="Each test must change one frozen decision at a time and retain its own interpretation boundary. The current test concerns donors per disease arm, not nuclei.",
        transition="Begin with the cohort-level consequences of the donor-threshold test.",
    )


def append_comparison_table(
    prs: Presentation,
    facts: dict[str, dict[int, dict[str, int]]],
) -> None:
    slide = ui.new_slide(prs)
    ui.add_title_block(
        slide,
        TABLE_TITLE,
        "Raising minimum donors per disease arm from 3 to 5 in ROSMAP and SEA-AD.",
    )
    rows = []
    for cohort_key, cohort_label in (("rosmap", "ROSMAP"), ("seaad", "SEA-AD")):
        for threshold in (3, 5):
            values = facts[cohort_key][threshold]
            rows.append(
                [
                    cohort_label,
                    str(threshold),
                    str(values["eligible_contrasts"]),
                    str(values["kda_runs"]),
                    str(values["driver_rows"]),
                    str(values["unique_non_mt_drivers"]),
                ]
            )
    ui.add_table(
        slide,
        [
            "Cohort",
            "Minimum donors / arm",
            "Eligible contrasts",
            "KDA runs",
            "Significant driver rows",
            "Unique non-MT drivers",
        ],
        rows,
        0.72,
        1.75,
        [1.45, 2.05, 1.85, 1.35, 2.45, 2.75],
        row_h=0.72,
        header_h=0.76,
        font_size=10.4,
    )
    ui.add_rect(
        slide, 0.72, 5.64, 11.90, 0.72,
        color=ui.PALE_SKY, outline=ui.SKY,
    )
    ui.add_text(
        slide,
        "Driver rows count significant returned KDA rows. Unique non-MT drivers exclude MT-* genes and are deduplicated within each cohort.",
        1.00, 5.84, 11.34, 0.31,
        size=10.8, color=ui.NAVY, bold=True, align=PP_ALIGN.CENTER,
    )
    ui.add_text(
        slide,
        "Source: Phase 22 and VH14 run manifests/returns; donor-per-arm 3-vs-5 sensitivity assessment (2026-09-11).",
        0.72, 7.08, 11.90, 0.16,
        size=6.7, color=ui.MID, align=PP_ALIGN.CENTER,
    )
    ui.add_notes(
        slide,
        goal="Show the numerical impact of raising the minimum disease-arm donor count from three to five in both cohorts.",
        walkthrough="ROSMAP falls from 42 to 40 eligible contrasts, three to two KDA runs, 39 to 13 returned rows, and 33 to 13 unique non-mitochondrial drivers. SEA-AD falls from 28 to 20 eligible contrasts, while its five KDA runs, 24 returned rows, and 18 unique non-mitochondrial drivers are unchanged.",
        boundary="The threshold test changes only the donor eligibility gate. It does not change retained contrast models, DEG queries, networks, or KDA settings.",
        transition="Next, interpret why the two cohorts respond differently and what this means for validation.",
    )


def append_findings(prs: Presentation) -> None:
    slide = ui.new_slide(prs)
    ui.add_title_block(
        slide,
        FINDINGS_TITLE,
        "Raising minimum donors per disease arm from 3 to 5 in ROSMAP and SEA-AD.",
    )
    ui.add_rect(
        slide, 0.72, 1.42, 11.90, 5.22,
        color=ui.WHITE, outline=ui.LIGHT,
    )
    ui.add_text(
        slide, "KEY FINDINGS", 1.05, 1.70, 2.40, 0.26,
        size=10.2, color=ui.BLUE, bold=True,
    )
    ui.add_bullets(
        slide,
        [
            "ROSMAP loses the Vasculature × M_e2 4-versus-3-donor call, which alone generated 26 driver rows. This removes the low-support result dominating the ROSMAP release.",
            "SEA-AD loses eight eligible contrasts, but all already had empty mitochondrial queries. Therefore, all five KDA runs and all returned drivers remain unchanged.",
            "Cross-cohort validation remains zero at both thresholds.",
            "The only exact shared executed stratum remains Astrocytes × F_e4, with no shared drivers.",
            "The excluded contrasts are treated as unavailable comparisons—not negative validation results.",
        ],
        1.05, 2.10, 10.95,
        size=14.2, accent=ui.BLUE, line_h=0.84,
    )
    ui.add_text(
        slide,
        "Source: Phase 22 and VH14 run manifests/returns; donor-per-arm 3-vs-5 sensitivity assessment (2026-09-11).",
        0.72, 7.08, 11.90, 0.16,
        size=6.7, color=ui.MID, align=PP_ALIGN.CENTER,
    )
    ui.add_notes(
        slide,
        goal="Summarize the cohort-specific and cross-cohort conclusions of the donor-threshold sensitivity analysis.",
        walkthrough="ROSMAP loses the low-support Vasculature M_e2 call that generated 26 driver rows. SEA-AD loses eight eligible contrasts, but none had progressed to KDA. Cross-cohort validation remains zero under both donor thresholds; Astrocytes F_e4 is the only exact stratum executed in both cohorts, and it has no shared drivers.",
        boundary="Contrasts excluded by the stricter donor threshold are unavailable comparisons. They must not be counted as evidence against validation.",
        transition="Continue with the next sensitivity analysis when it is available.",
    )


def validate_deck(
    path: Path,
    preserved: tuple[tuple[Any, ...], ...],
) -> list[dict[str, Any]]:
    prs = Presentation(str(path))
    checks: list[dict[str, Any]] = []

    def add(check_id: str, observed: Any, expected: Any, passed: bool) -> None:
        checks.append(
            {
                "schema_version": "broad_kda_sensitivity_deck_checks_v1",
                "check_id": check_id,
                "observed": observed,
                "expected": expected,
                "passed": str(bool(passed)).upper(),
            }
        )

    add("slide_count", len(prs.slides), EXPECTED_OUTPUT_SLIDES, len(prs.slides) == EXPECTED_OUTPUT_SLIDES)
    add(
        "widescreen_ratio",
        round(prs.slide_width / prs.slide_height, 3),
        1.778,
        round(prs.slide_width / prs.slide_height, 3) == 1.778,
    )
    actual_preserved = tuple(
        slide_signature(slide) for slide in list(prs.slides)[:24]
    )
    add("existing_slides_preserved", actual_preserved == preserved, True, actual_preserved == preserved)
    add(
        "divider_title",
        DIVIDER_TITLE in slide_text(prs.slides[24]),
        True,
        DIVIDER_TITLE in slide_text(prs.slides[24]),
    )
    table_text = slide_text(prs.slides[25])
    required_table = [
        TABLE_TITLE,
        "Minimum donors / arm",
        "Eligible contrasts",
        "KDA runs",
        "Significant driver rows",
        "Unique non-MT drivers",
    ]
    missing_table = [text for text in required_table if text not in table_text]
    add("comparison_table_present", missing_table, [], not missing_table)
    findings_text = slide_text(prs.slides[26])
    required_findings = [
        FINDINGS_TITLE,
        "ROSMAP loses the Vasculature × M_e2 4-versus-3-donor call",
        "SEA-AD loses eight eligible contrasts",
        "Cross-cohort validation remains zero at both thresholds.",
        "The only exact shared executed stratum remains Astrocytes × F_e4, with no shared drivers.",
        "The excluded contrasts are treated as unavailable comparisons—not negative validation results.",
    ]
    missing_findings = [
        text for text in required_findings if text not in findings_text
    ]
    add("finding_claims_present", missing_findings, [], not missing_findings)
    add(
        "admissible_wording_removed",
        "admissible" in findings_text.lower(),
        False,
        "admissible" not in findings_text.lower(),
    )

    notes_ok = 0
    in_bounds = 0
    for slide in list(prs.slides)[24:]:
        notes_frame = slide.notes_slide.notes_text_frame
        notes = notes_frame.text if notes_frame is not None else ""
        if all(
            marker in notes
            for marker in (
                "Teaching goal:",
                "Walk through:",
                "Scientific boundary:",
                "Transition:",
            )
        ):
            notes_ok += 1
        bounded = all(
            shape.left >= 0
            and shape.top >= 0
            and shape.left + shape.width <= prs.slide_width + 10
            and shape.top + shape.height <= prs.slide_height + 10
            for shape in slide.shapes
        )
        if bounded:
            in_bounds += 1
    add("structured_notes", notes_ok, 3, notes_ok == 3)
    add("new_shapes_in_bounds", in_bounds, 3, in_bounds == 3)

    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
    add("pptx_zip_integrity", bad_member, None, bad_member is None)
    return checks


def main() -> int:
    args = parse_args()
    input_path = args.input.resolve()
    output_path = args.output.resolve()
    audit_dir = args.audit_dir.resolve()
    if not input_path.is_file():
        raise FileNotFoundError(input_path)

    facts = load_facts()
    source_hash = sha256(input_path)
    prs = Presentation(str(input_path))
    if len(prs.slides) != EXPECTED_INPUT_SLIDES:
        raise RuntimeError(
            f"Expected {EXPECTED_INPUT_SLIDES}-slide input deck, found {len(prs.slides)}; "
            "refusing to append duplicate or misplaced slides"
        )
    if "What the SEA-AD-network rerun" not in slide_text(prs.slides[-1]):
        raise RuntimeError("The expected Part 3 closing slide is not last")

    preserved = tuple(slide_signature(slide) for slide in prs.slides)
    ui.set_notes_body_template(
        prs.slides[0].notes_slide.notes_placeholder._element
    )
    append_divider(prs)
    append_comparison_table(prs, facts)
    append_findings(prs)
    if len(prs.slides) != EXPECTED_OUTPUT_SLIDES:
        raise RuntimeError("The three requested slides were not appended")

    old_comments = prs.core_properties.comments or ""
    addition = "Appended Part 4 sensitivity-analysis section, donor-threshold comparison table, and key findings on 2026-09-11."
    if addition not in old_comments:
        prs.core_properties.comments = f"{old_comments} {addition}".strip()

    audit_dir.mkdir(parents=True, exist_ok=True)
    backup = audit_dir / f"source_before_append_{source_hash[:12]}.pptx"
    if backup.exists():
        if sha256(backup) != source_hash:
            raise RuntimeError(f"Existing backup hash mismatch: {backup}")
    else:
        shutil.copy2(input_path, backup)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temp_name = tempfile.mkstemp(
        prefix=f".{output_path.name}.", suffix=".tmp", dir=output_path.parent
    )
    os.close(descriptor)
    temporary = Path(temp_name)
    try:
        prs.save(str(temporary))
        checks = validate_deck(temporary, preserved)
        failed = [row["check_id"] for row in checks if row["passed"] != "TRUE"]
        if failed:
            raise RuntimeError("Deck validation failed: " + ", ".join(failed))
        temporary.replace(output_path)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise

    output_hash = sha256(output_path)
    write_tsv(
        audit_dir / "checks.tsv",
        checks,
        ["schema_version", "check_id", "observed", "expected", "passed"],
    )
    sensitivity_rows = []
    for cohort in ("rosmap", "seaad"):
        for threshold in (3, 5):
            sensitivity_rows.append(
                {
                    "schema_version": "broad_kda_donor_threshold_sensitivity_v1",
                    "cohort": cohort,
                    "minimum_donors_per_arm": threshold,
                    **facts[cohort][threshold],
                }
            )
    write_tsv(
        audit_dir / "sensitivity_summary.tsv",
        sensitivity_rows,
        [
            "schema_version",
            "cohort",
            "minimum_donors_per_arm",
            "structural_contrasts",
            "eligible_contrasts",
            "kda_runs",
            "driver_rows",
            "unique_raw_drivers",
            "unique_non_mt_drivers",
            "global_driver_rows",
            "unique_global_drivers",
            "unique_global_non_mt_drivers",
        ],
    )
    write_tsv(
        audit_dir / "status.tsv",
        [
            {
                "schema_version": "broad_kda_sensitivity_deck_status_v1",
                "validation_status": "validated_complete",
                "failed_checks": 0,
                "input_sha256": source_hash,
                "output_sha256": output_hash,
                "slide_count": EXPECTED_OUTPUT_SLIDES,
                "backup_path": str(backup.relative_to(ROOT)),
                "output_path": str(output_path.relative_to(ROOT)),
            }
        ],
        [
            "schema_version",
            "validation_status",
            "failed_checks",
            "input_sha256",
            "output_sha256",
            "slide_count",
            "backup_path",
            "output_path",
        ],
    )
    print(f"output={output_path}")
    print(f"slides={EXPECTED_OUTPUT_SLIDES}")
    print(f"sha256={output_hash}")
    print(f"backup={backup}")
    print(f"audit={audit_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
