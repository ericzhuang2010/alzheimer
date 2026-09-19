#!/usr/bin/env python3
"""Build or refresh the two sex/APOE OXPHOS DEG result slides.

The replacement slides summarize the two DEG programs that lead directly to
the later biological findings. Existing slides, notes, and order are preserved
outside the two target positions.
"""

from __future__ import annotations

import argparse
import hashlib
import os
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
import update_phase11_seaad_simple_aggr_part2 as ui  # noqa: E402


DEFAULT_DECK = (
    ROOT
    / "docs"
    / "presentations"
    / "09162026"
    / "09162026_sex_apoe_kda_fine_broad.pptx"
)
AUDIT_ROOT = ROOT / "results" / "presentations" / "09162026_deg_results"

EXPECTED_SLIDE_COUNTS = {48, 49, 50, 52, 157, 158, 159}
OLD_TITLES = (
    "What is a differentially expressed gene (DEG)?",
    "Why start with fine cell types?",
)
NEW_TITLES = (
    "Female ε3/ε3 neurons repeatedly increase mtDNA-encoded OXPHOS RNA",
    "Male ε3/ε3 cells show opposite directions in two OXPHOS gene sets",
)
LEGACY_RESULT_TITLES = (
    "Female ε3/ε3 neurons repeatedly increase mitochondrial energy-gene RNA",
    "Male ε3/ε3 cells show a split between mitochondrial and nuclear energy genes",
)
def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_DECK)
    parser.add_argument("--output", type=Path, default=DEFAULT_DECK)
    parser.add_argument("--audit-root", type=Path, default=AUDIT_ROOT)
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def slide_title(slide) -> str:
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text.strip():
            return " ".join(shape.text.split())
    return ""


def slide_fingerprint(slide) -> tuple[Any, ...]:
    notes = slide.notes_slide.notes_text_frame
    return (
        tuple(
            (
                int(shape.shape_type),
                shape.name,
                int(shape.left),
                int(shape.top),
                int(shape.width),
                int(shape.height),
                shape.text if getattr(shape, "has_text_frame", False) else "",
            )
            for shape in slide.shapes
        ),
        notes.text if notes is not None else "",
    )


def resolve_replace_start(prs: Presentation) -> int:
    """Locate the adjacent result slides without imposing an obsolete order."""
    if len(prs.slides) not in EXPECTED_SLIDE_COUNTS:
        raise RuntimeError(f"Unexpected slide count: {len(prs.slides)}")
    titles = [slide_title(slide) for slide in prs.slides]
    accepted_pairs = {OLD_TITLES, LEGACY_RESULT_TITLES, NEW_TITLES}
    matches = [
        index
        for index in range(len(titles) - 1)
        if tuple(titles[index : index + 2]) in accepted_pairs
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one adjacent DEG-result pair, found {matches}")
    return matches[0]


def add_metric_card(
    slide,
    *,
    x: float,
    y: float,
    w: float,
    value: str,
    label: str,
    detail: str,
    accent,
    background,
) -> None:
    ui.add_rect(slide, x, y, w, 1.62, color=background, outline=None)
    ui.add_text(
        slide,
        value,
        x + 0.22,
        y + 0.18,
        w - 0.44,
        0.52,
        size=27.0,
        color=ui.readable_accent(accent),
        bold=True,
        font=ui.FONT_HEAD,
        align=PP_ALIGN.CENTER,
    )
    ui.add_text(
        slide,
        label,
        x + 0.22,
        y + 0.75,
        w - 0.44,
        0.32,
        size=13.0,
        color=ui.NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    ui.add_text(
        slide,
        detail,
        x + 0.22,
        y + 1.12,
        w - 0.44,
        0.28,
        size=10.8,
        color=ui.GRAY,
        align=PP_ALIGN.CENTER,
    )


def clear_slide(slide) -> None:
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = ui.OFF_WHITE


def build_female_e33_slide(slide):
    clear_slide(slide)
    ui.add_title_block(
        slide,
        NEW_TITLES[0],
        "This is the clearest sex/APOE-related mitochondrial RNA pattern.",
    )

    ui.add_rect(slide, 0.72, 1.35, 11.90, 0.72, color=ui.WHITE, outline=ui.LIGHT)
    ui.add_text(
        slide,
        "Mitochondrial DNA encodes 13 structural subunits of OXPHOS.",
        1.03,
        1.57,
        11.28,
        0.28,
        size=14.2,
        color=ui.NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    ui.add_rect(slide, 0.72, 2.33, 7.52, 3.05, color=ui.PALE_SKY, outline=None)
    ui.add_text(
        slide,
        "ROSMAP FINE-CELL RESULTS",
        1.03,
        2.65,
        6.90,
        0.27,
        size=11.0,
        color=ui.BLUE,
        bold=True,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        "More increased mtDNA-encoded OXPHOS genes than expected by chance",
        1.03,
        3.04,
        6.90,
        0.38,
        size=16.0,
        color=ui.NAVY,
        bold=True,
    )
    add_metric_card(
        slide,
        x=1.03,
        y=3.56,
        w=3.22,
        value="12 of 14",
        label="excitatory comparisons",
        detail="showed this upward pattern",
        accent=ui.BLUE,
        background=ui.WHITE,
    )
    add_metric_card(
        slide,
        x=4.65,
        y=3.56,
        w=3.22,
        value="13 of 14",
        label="inhibitory comparisons",
        detail="showed this upward pattern",
        accent=ui.BLUE,
        background=ui.WHITE,
    )

    ui.add_rect(slide, 8.52, 2.33, 4.10, 3.05, color=ui.PALE_GREEN, outline=None)
    ui.add_text(
        slide,
        "SEA-AD SUPPORT",
        8.84,
        2.65,
        3.46,
        0.27,
        size=11.0,
        color=ui.TEAL_TEXT,
        bold=True,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        "One matched excitatory comparison",
        8.84,
        3.04,
        3.46,
        0.50,
        size=16.0,
        color=ui.NAVY,
        bold=True,
    )
    ui.add_text(
        slide,
        "9 of 10",
        8.84,
        3.77,
        3.46,
        0.52,
        size=28.0,
        color=ui.TEAL_TEXT,
        bold=True,
        font=ui.FONT_HEAD,
        align=PP_ALIGN.CENTER,
    )
    ui.add_text(
        slide,
        "genes were the same mtDNA-encoded OXPHOS genes seen in ROSMAP",
        8.94,
        4.34,
        3.26,
        0.58,
        size=12.4,
        color=ui.DARK,
        align=PP_ALIGN.CENTER,
    )
    ui.add_text(
        slide,
        "All 9 increased in disease in both cohorts.",
        8.84,
        4.94,
        3.46,
        0.28,
        size=11.8,
        color=ui.TEAL_TEXT,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    ui.add_rect(slide, 0.72, 5.70, 11.90, 0.94, color=ui.PALE_GOLD, outline=None)
    ui.add_text(
        slide,
        "Most careful reading",
        1.03,
        5.98,
        1.72,
        0.28,
        size=12.4,
        color=ui.GOLD_TEXT,
        bold=True,
    )
    ui.add_text(
        slide,
        "A repeated mitochondrial response that may reflect compensation or stress—not proof of better energy production.",
        2.86,
        5.94,
        9.26,
        0.38,
        size=14.0,
        color=ui.DARK,
        valign=MSO_ANCHOR.MIDDLE,
    )

    ui.add_notes(
        slide,
        goal="Preview the strongest sex/APOE-related mitochondrial DEG program.",
        walkthrough=(
            "Mitochondrial DNA encodes 13 structural subunits of OXPHOS. In female "
            "APOE epsilon-3 homozygous "
            "ROSMAP comparisons, 12 of 14 excitatory and 13 of 14 inhibitory "
            "fine-cell comparisons contained more increased genes from this program "
            "than expected by chance. In the one matched SEA-AD excitatory "
            "comparison, nine of ten genes in the SEA-AD list were the same "
            "mtDNA-encoded OXPHOS genes, and all nine increased in disease in both "
            "cohorts."
        ),
        boundary=(
            "The SEA-AD support comes from one matched comparison, so it is strong "
            "in direction but narrow in coverage. More RNA may reflect compensation "
            "or stress and does not demonstrate better energy production."
        ),
        transition=(
            "Contrast this coordinated increase with the male epsilon-3 homozygous "
            "split between two genomes."
        ),
    )
    return slide


def add_direction_panel(
    slide,
    *,
    x: float,
    heading: str,
    value: str,
    direction: str,
    detail: str,
    accent,
    background,
) -> None:
    ui.add_rect(slide, x, 1.86, 5.54, 2.62, color=background, outline=None)
    ui.add_text(
        slide,
        heading,
        x + 0.30,
        2.17,
        4.94,
        0.30,
        size=14.0,
        color=ui.NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    ui.add_text(
        slide,
        value,
        x + 0.30,
        2.72,
        4.94,
        0.62,
        size=31.0,
        color=ui.readable_accent(accent),
        bold=True,
        font=ui.FONT_HEAD,
        align=PP_ALIGN.CENTER,
    )
    ui.add_text(
        slide,
        direction,
        x + 0.30,
        3.40,
        4.94,
        0.34,
        size=16.0,
        color=ui.readable_accent(accent),
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    ui.add_text(
        slide,
        detail,
        x + 0.48,
        3.89,
        4.58,
        0.30,
        size=11.2,
        color=ui.GRAY,
        align=PP_ALIGN.CENTER,
    )


def build_male_e33_slide(slide):
    clear_slide(slide)
    ui.add_title_block(
        slide,
        NEW_TITLES[1],
        "Mitochondrial and nuclear DNA encode structural subunits of the same OXPHOS system.",
    )

    add_direction_panel(
        slide,
        x=0.72,
        heading="13 mtDNA-encoded OXPHOS genes",
        value="70 ↑  |  11 ↓",
        direction="86% increased in AD",
        detail="81 gene appearances across fine-cell comparisons",
        accent=ui.BLUE,
        background=ui.PALE_SKY,
    )
    add_direction_panel(
        slide,
        x=7.08,
        heading="86 nuclear-encoded structural OXPHOS genes",
        value="8 ↑  |  60 ↓",
        direction="88% decreased in AD",
        detail="68 gene appearances across fine-cell comparisons",
        accent=ui.VERMILION,
        background=ui.PALE_RED,
    )
    ui.add_text(
        slide,
        "↔",
        6.24,
        2.74,
        0.84,
        0.58,
        size=27.0,
        color=ui.MID,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    ui.add_rect(slide, 0.72, 4.78, 7.52, 1.48, color=ui.PALE_GREEN, outline=None)
    ui.add_text(
        slide,
        "SEA-AD MATCHED INHIBITORY-NEURON SUPPORT",
        1.03,
        5.05,
        6.90,
        0.26,
        size=10.7,
        color=ui.TEAL_TEXT,
        bold=True,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        "19 of 22 shared genes moved in the same direction",
        1.03,
        5.43,
        6.90,
        0.34,
        size=17.0,
        color=ui.NAVY,
        bold=True,
    )
    ui.add_text(
        slide,
        "Nine mtDNA-encoded OXPHOS genes rose while several nuclear-encoded OXPHOS genes fell.",
        1.03,
        5.84,
        6.90,
        0.28,
        size=11.5,
        color=ui.DARK,
    )

    ui.add_rect(slide, 8.52, 4.78, 4.10, 1.48, color=ui.PALE_GOLD, outline=None)
    ui.add_text(
        slide,
        "IMPORTANT LIMIT",
        8.84,
        5.05,
        3.46,
        0.26,
        size=10.7,
        color=ui.GOLD_TEXT,
        bold=True,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        "The broad-cell ROSMAP result was no longer statistically clear after accounting for subtype mixture.",
        8.84,
        5.43,
        3.46,
        0.60,
        size=12.2,
        color=ui.DARK,
    )

    ui.add_text(
        slide,
        "Each percentage uses gene-by-fine-cell DEG occurrences; a gene can recur across comparisons.",
        0.76,
        6.55,
        11.80,
        0.22,
        size=9.8,
        color=ui.GRAY,
        italic=True,
        align=PP_ALIGN.CENTER,
    )

    ui.add_notes(
        slide,
        goal=(
            "Preview the male epsilon-3 homozygous split between mitochondrial and "
            "nuclear-encoded OXPHOS genes."
        ),
        walkthrough=(
            "Mitochondrial and nuclear DNA encode structural subunits of the same "
            "OXPHOS system. Across eligible male APOE epsilon-3 homozygous ROSMAP "
            "fine-cell comparisons, mtDNA-encoded OXPHOS genes appeared 81 times: "
            "70 increased and 11 decreased. Nuclear-encoded structural OXPHOS genes "
            "appeared 68 times: 8 increased and 60 decreased. In matched SEA-AD "
            "inhibitory-neuron evidence, 19 of 22 shared genes moved in the same "
            "direction, including nine mtDNA-encoded OXPHOS genes that increased and "
            "several nuclear-encoded OXPHOS genes that decreased."
        ),
        boundary=(
            "Gene appearances repeat genes across fine-cell comparisons and are not "
            "independent donors. In addition, the direct broad-cell ROSMAP pattern "
            "was no longer statistically clear after accounting for subtype mixture, "
            "so this remains a testable RNA-level hypothesis rather than proof of a "
            "sex/APOE interaction or impaired energy production."
        ),
        transition="Move to Part 3, where mitochondrial DEGs become network-analysis queries.",
    )
    return slide


def validate_deck(
    prs: Presentation, preserved: tuple[Any, ...], replace_start: int
) -> None:
    if len(prs.slides) not in EXPECTED_SLIDE_COUNTS:
        raise RuntimeError(f"Unexpected slide count: {len(prs.slides)}")
    if tuple(slide_title(prs.slides[replace_start + i]) for i in range(2)) != NEW_TITLES:
        raise RuntimeError("DEG result slides are not in the expected position")
    after_existing = tuple(
        slide_fingerprint(slide)
        for index, slide in enumerate(prs.slides)
        if index not in {replace_start, replace_start + 1}
    )
    if after_existing != preserved:
        raise RuntimeError("A retained slide changed during DEG-slide replacement")
    for slide in list(prs.slides)[replace_start : replace_start + 2]:
        for shape in slide.shapes:
            if (
                shape.left < 0
                or shape.top < 0
                or shape.left + shape.width > prs.slide_width + 10
                or shape.top + shape.height > prs.slide_height + 10
            ):
                raise RuntimeError(f"Out-of-bounds shape on {slide_title(slide)!r}")
        notes = slide.notes_slide.notes_text_frame
        if notes is None or not notes.text.strip():
            raise RuntimeError(f"Missing speaker notes on {slide_title(slide)!r}")


def main() -> int:
    args = parse_args()
    input_path = args.input.resolve()
    output_path = args.output.resolve()
    audit_root = args.audit_root.resolve()
    if not input_path.is_file():
        raise FileNotFoundError(input_path)

    input_hash = sha256(input_path)
    prs = Presentation(input_path)
    replace_start = resolve_replace_start(prs)
    observed_titles = tuple(
        slide_title(prs.slides[replace_start + i]) for i in range(2)
    )
    if observed_titles not in {OLD_TITLES, LEGACY_RESULT_TITLES, NEW_TITLES}:
        raise RuntimeError(
            f"Unexpected DEG result slides: {observed_titles!r}"
        )

    preserved = tuple(
        slide_fingerprint(slide)
        for index, slide in enumerate(prs.slides)
        if index not in {replace_start, replace_start + 1}
    )
    ui.set_notes_body_template(
        prs.slides[0].notes_slide.notes_placeholder._element
    )
    build_female_e33_slide(prs.slides[replace_start])
    build_male_e33_slide(prs.slides[replace_start + 1])
    validate_deck(prs, preserved, replace_start)

    audit_root.mkdir(parents=True, exist_ok=True)
    backup_dir = audit_root / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / f"input_deck_{input_hash[:12]}.pptx"
    if not backup_path.exists():
        shutil.copy2(input_path, backup_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    handle, temp_name = tempfile.mkstemp(
        prefix=output_path.stem + ".", suffix=".pptx", dir=output_path.parent
    )
    os.close(handle)
    temp_path = Path(temp_name)
    try:
        prs.save(temp_path)
        reopened = Presentation(temp_path)
        validate_deck(reopened, preserved, replace_start)
        os.replace(temp_path, output_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()

    print(f"Updated deck: {output_path}")
    print(f"Refreshed slides: {replace_start + 1}-{replace_start + 2}")
    print(f"Backup: {backup_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
