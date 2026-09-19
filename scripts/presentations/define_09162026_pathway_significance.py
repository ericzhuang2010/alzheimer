#!/usr/bin/env python3
"""Define "significant result" on slide 6 without changing its layout."""

from __future__ import annotations

import argparse
import hashlib
import os
import shutil
import tempfile
from pathlib import Path

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DECK = (
    ROOT
    / "docs"
    / "presentations"
    / "09162026"
    / "09162026_sex_apoe_kda_fine_broad.pptx"
)
AUDIT_ROOT = ROOT / "results" / "presentations" / "09162026_pathway_significance"
SLIDE_NUMBER = 6
EXPECTED_SLIDES = 159
EXPECTED_TITLE = "How the pathway analysis was done"
OLD_TEXT = (
    "Significant result = one DEG list showed greater-than-expected overlap with one "
    "pathway after correction across that pathway collection (false-discovery rate < 5%)."
)
NEW_TEXT = (
    "Significant result = greater-than-expected pathway overlap after multiple-test "
    "correction (false-discovery rate < 5%)."
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_DECK)
    parser.add_argument("--output", type=Path, default=DEFAULT_DECK)
    parser.add_argument("--audit-root", type=Path, default=AUDIT_ROOT)
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def slide_title(slide) -> str:
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text.strip():
            return " ".join(shape.text.split())
    return ""


def shape_geometry(prs: Presentation) -> tuple[tuple[object, ...], ...]:
    return tuple(
        (
            slide_number,
            int(shape.shape_type),
            shape.name,
            int(shape.left),
            int(shape.top),
            int(shape.width),
            int(shape.height),
        )
        for slide_number, slide in enumerate(prs.slides, start=1)
        for shape in slide.shapes
    )


def replace_text_preserving_format(slide) -> None:
    matches = [
        shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text == OLD_TEXT
    ]
    if len(matches) != 1:
        if any(
            getattr(shape, "has_text_frame", False) and shape.text == NEW_TEXT
            for shape in slide.shapes
        ):
            return
        raise RuntimeError(f"Expected one significance-definition text box, found {len(matches)}")

    frame = matches[0].text_frame
    if len(frame.paragraphs) != 1 or len(frame.paragraphs[0].runs) != 1:
        raise RuntimeError("The significance-definition text box format changed")
    frame.paragraphs[0].runs[0].text = NEW_TEXT


def validate(prs: Presentation, geometry: tuple[tuple[object, ...], ...]) -> None:
    if len(prs.slides) != EXPECTED_SLIDES:
        raise RuntimeError(f"Expected {EXPECTED_SLIDES} slides, found {len(prs.slides)}")
    slide = prs.slides[SLIDE_NUMBER - 1]
    if slide_title(slide) != EXPECTED_TITLE:
        raise RuntimeError(f"Unexpected slide {SLIDE_NUMBER} title: {slide_title(slide)!r}")
    if shape_geometry(prs) != geometry:
        raise RuntimeError("A slide object's geometry changed")
    visible = "\n".join(
        shape.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    for required in (
        "Significant result =",
        "greater-than-expected pathway overlap",
        "false-discovery rate < 5%",
    ):
        if required not in visible:
            raise RuntimeError(f"Missing {required!r} from slide {SLIDE_NUMBER}")


def main() -> int:
    args = parse_args()
    input_path = args.input.resolve()
    output_path = args.output.resolve()
    audit_root = args.audit_root.resolve()
    if not input_path.is_file():
        raise FileNotFoundError(input_path)

    input_hash = sha256(input_path)
    prs = Presentation(input_path)
    if len(prs.slides) != EXPECTED_SLIDES:
        raise RuntimeError(f"Expected {EXPECTED_SLIDES} slides, found {len(prs.slides)}")
    geometry = shape_geometry(prs)
    replace_text_preserving_format(prs.slides[SLIDE_NUMBER - 1])
    validate(prs, geometry)

    audit_root.mkdir(parents=True, exist_ok=True)
    backup_dir = audit_root / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / f"input_deck_{input_hash[:12]}.pptx"
    if not backup_path.exists():
        shutil.copy2(input_path, backup_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    handle, temp_name = tempfile.mkstemp(
        prefix=output_path.stem + ".", suffix=".pptx", dir=output_path.parent
    )
    os.close(handle)
    temp_path = Path(temp_name)
    try:
        prs.save(temp_path)
        reopened = Presentation(temp_path)
        validate(reopened, geometry)
        os.replace(temp_path, output_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()

    print(f"Defined significant result on slide {SLIDE_NUMBER}: {output_path}")
    print(f"Backup: {backup_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
