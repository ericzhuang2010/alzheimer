#!/usr/bin/env python3
"""Append an eight-slide Finding 2 section to the dated sex/APOE KDA deck."""

from __future__ import annotations

import argparse
import copy
import csv
import os
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Any

import pandas as pd
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from scipy.stats import hypergeom

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
import append_finding1_slides as f1  # noqa: E402

ui = f1.ui

DECK = (
    ROOT
    / "docs"
    / "presentations"
    / "09162026"
    / "09162026_sex_apoe_kda_fine_broad.pptx"
)
AUDIT_ROOT = ROOT / "results" / "presentations" / "09162026_finding2_slides"
ROS_DIR = ROOT / "results" / "minerva_production" / "20_sex_apoe_kda_combo"
SEA_DIR = ROOT / "results" / "validation_human" / "12_sex_apoe_kda_combo"

EXPECTED_INPUT_SLIDES = 35
NEW_SLIDE_COUNT = 8
INSERT_INDEX = EXPECTED_INPUT_SLIDES
SECTION_MARKER = "FINDING 2"
SECTION_TITLE = (
    "Male APOE ε3/ε3 cells show a mtDNA-up, "
    "nuclear-OXPHOS-down mismatch"
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DECK)
    parser.add_argument("--output", type=Path, default=DECK)
    parser.add_argument("--audit-root", type=Path, default=AUDIT_ROOT)
    return parser.parse_args()


def prepare_facts() -> dict[str, Any]:
    modules = f1.read_tsv(ROOT / "config" / "phase13_respiratory_modules.tsv")
    module_sets = {
        module_id: set(rows["current_approved_symbol"])
        for module_id, rows in modules.groupby("module_id")
    }
    mt_genes = module_sets["mtdna_oxphos_13"]
    nuclear_genes = module_sets["nuclear_oxphos_structural_86"]

    ros_query = f1.read_tsv(ROS_DIR / "combo_query_members.tsv.gz")
    ros_query = ros_query.loc[
        f1.true_series(ros_query["included_at_thresholds"])
        & f1.true_series(ros_query["effective_member"])
    ].copy()
    ros_manifest = f1.read_tsv(ROS_DIR / "combo_run_manifest.tsv")
    ros_included = ros_manifest.loc[
        f1.true_series(ros_manifest["included_at_thresholds"])
        & ros_manifest["signature_group"].eq("M_e33")
    ].copy()
    ros_group_query = ros_query.loc[
        ros_query["signature_group"].eq("M_e33")
    ].copy()
    ros_mt = ros_group_query.loc[ros_group_query["gene"].isin(mt_genes)]
    ros_nuclear = ros_group_query.loc[
        ros_group_query["gene"].isin(nuclear_genes)
    ]

    background = f1.read_tsv(
        ROOT
        / "results"
        / "minerva_production"
        / "12_rosmap_sex_apoe_mito_kda_runs"
        / "kda_background_members.tsv.gz"
    )
    background = background.loc[
        background["kda_run_id"].isin(ros_included["kda_run_id"])
    ]
    annotation = f1.read_tsv(
        ROOT
        / "results"
        / "minerva_production"
        / "09_annotate_genes"
        / "gene_annotation_master.tsv.gz"
    )
    mito_annotation = annotation.loc[
        f1.true_series(annotation["is_mitocarta3"])
        & annotation["symbol_hgnc_current"].notna()
    ]
    annotation_sets = {
        rds_id: set(rows["symbol_hgnc_current"])
        for rds_id, rows in mito_annotation.groupby("rds_id")
    }
    background_sets = {
        run_id: set(rows["gene"])
        for run_id, rows in background.groupby("kda_run_id")
    }
    query_sets = {
        run_id: set(rows["gene"])
        for run_id, rows in ros_query.groupby("kda_run_id")
    }

    significant_runs = {
        "mtdna_oxphos_13": set(),
        "nuclear_oxphos_structural_86": set(),
    }
    for row in ros_included.itertuples(index=False):
        source_rds = row.source_contrast_id.split("::", maxsplit=1)[0]
        exact_background = (
            background_sets[row.kda_run_id] & annotation_sets[source_rds]
        )
        query = query_sets[row.kda_run_id]
        module_ids: list[str] = []
        p_values: list[float] = []
        for _, module_rows in modules.groupby("module_order"):
            module_id = str(module_rows["module_id"].iloc[0])
            genes = set(module_rows["current_approved_symbol"])
            module_ids.append(module_id)
            p_values.append(
                hypergeom.sf(
                    len(query & genes) - 1,
                    len(exact_background),
                    len(exact_background & genes),
                    len(query),
                )
            )
        for module_id, adjusted_p in zip(
            module_ids,
            f1.bh_adjust(p_values),
            strict=True,
        ):
            if module_id in significant_runs and adjusted_p < 0.05:
                significant_runs[module_id].add(row.kda_run_id)

    ros_inhibitory = ros_group_query.loc[
        ros_group_query["broad_cell_type"].eq("Inhibitory_neurons")
    ]
    ros_directions: dict[str, str] = {}
    for gene, rows in ros_inhibitory.groupby("gene"):
        directions = set(
            "up" if value else "down"
            for value in f1.true_series(rows["in_upregulated_query"])
        )
        ros_directions[str(gene)] = (
            next(iter(directions)) if len(directions) == 1 else "mixed"
        )

    sea_manifest = f1.read_tsv(
        SEA_DIR / "10a_inputs" / "seaad_kda_run_manifest.tsv"
    )
    sea_match = sea_manifest.loc[
        sea_manifest["signature_group"].eq("M_e33")
        & sea_manifest["broad_network"].eq("Inhibitory_neurons")
        & sea_manifest["kda_run_id"].notna()
    ].copy()
    sea_members = f1.read_tsv(
        SEA_DIR / "10a_inputs" / "seaad_kda_signature_members.tsv.gz"
    )
    sea_effective = sea_members.loc[
        sea_members["kda_run_id"].isin(sea_match["kda_run_id"])
        & f1.true_series(sea_members["effective_member"])
    ].copy()
    sea_effective_sets = {
        run_id: set(rows["gene"])
        for run_id, rows in sea_effective.groupby("kda_run_id")
    }

    sea_direction_rows: list[dict[str, str]] = []
    for row in sea_match.itertuples(index=False):
        result = f1.read_tsv(ROOT / str(row.result_path))
        result = result.loc[
            result["current_symbol_for_kda"].isin(
                sea_effective_sets[row.kda_run_id]
            )
        ]
        for result_row in result.itertuples(index=False):
            sea_direction_rows.append(
                {
                    "gene": str(result_row.current_symbol_for_kda),
                    "direction": (
                        "up"
                        if result_row.effect_direction == "Dementia_up"
                        else "down"
                    ),
                }
            )
    sea_direction_frame = pd.DataFrame(sea_direction_rows)
    sea_directions: dict[str, str] = {}
    for gene, rows in sea_direction_frame.groupby("gene"):
        directions = set(rows["direction"])
        sea_directions[str(gene)] = (
            next(iter(directions)) if len(directions) == 1 else "mixed"
        )

    shared_genes = sorted(set(ros_directions) & set(sea_directions))
    concordant = sorted(
        gene
        for gene in shared_genes
        if ros_directions[gene] != "mixed"
        and ros_directions[gene] == sea_directions[gene]
    )
    discordant = sorted(
        gene
        for gene in shared_genes
        if ros_directions[gene] != "mixed"
        and sea_directions[gene] != "mixed"
        and ros_directions[gene] != sea_directions[gene]
    )
    concordant_mt_up = sorted(
        gene
        for gene in concordant
        if gene in mt_genes and ros_directions[gene] == "up"
    )
    highlighted_down = [
        "UQCRFS1",
        "SLC25A5",
        "TIMM13",
        "TIMM17A",
        "MRPL16",
        "MRPS34",
        "ENDOG",
        "MTCH1",
    ]
    for gene in highlighted_down:
        if gene not in concordant or ros_directions[gene] != "down":
            raise RuntimeError(f"Unexpected cross-cohort direction for {gene}")
    other_concordant = sorted(
        set(concordant) - set(concordant_mt_up) - set(highlighted_down)
    )

    facts = {
        "ros_eligible_calls": len(ros_included),
        "ros_mt_occurrences": len(ros_mt),
        "ros_mt_up": int(f1.true_series(ros_mt["in_upregulated_query"]).sum()),
        "ros_mt_down": int(
            f1.true_series(ros_mt["in_downregulated_query"]).sum()
        ),
        "ros_mt_significant_calls": len(
            significant_runs["mtdna_oxphos_13"]
        ),
        "ros_nuclear_occurrences": len(ros_nuclear),
        "ros_nuclear_up": int(
            f1.true_series(ros_nuclear["in_upregulated_query"]).sum()
        ),
        "ros_nuclear_down": int(
            f1.true_series(ros_nuclear["in_downregulated_query"]).sum()
        ),
        "ros_nuclear_significant_calls": len(
            significant_runs["nuclear_oxphos_structural_86"]
        ),
        "sea_evaluable_calls": len(sea_match),
        "sea_shared_genes": len(shared_genes),
        "sea_concordant_genes": len(concordant),
        "sea_discordant_genes": len(discordant),
        "sea_mt_up_genes": concordant_mt_up,
        "sea_highlighted_down_genes": highlighted_down,
        "sea_other_concordant": other_concordant,
        "sea_discordant": discordant,
        "cross_cohort_expected": 15.17,
        "cross_cohort_raw_p": 0.0167,
        "cross_cohort_bh": 0.0417,
        "identity_resolved_shared": 21,
        "identity_resolved_bh": 0.0486,
    }
    expected = {
        "ros_eligible_calls": 31,
        "ros_mt_occurrences": 81,
        "ros_mt_up": 70,
        "ros_mt_down": 11,
        "ros_mt_significant_calls": 12,
        "ros_nuclear_occurrences": 68,
        "ros_nuclear_up": 8,
        "ros_nuclear_down": 60,
        "ros_nuclear_significant_calls": 1,
        "sea_evaluable_calls": 8,
        "sea_shared_genes": 22,
        "sea_concordant_genes": 19,
        "sea_discordant_genes": 3,
    }
    for key, value in expected.items():
        if facts[key] != value:
            raise RuntimeError(f"Source drift for {key}: {facts[key]} != {value}")
    expected_mt = [
        "MT-ATP6",
        "MT-CO2",
        "MT-CO3",
        "MT-CYB",
        "MT-ND1",
        "MT-ND2",
        "MT-ND3",
        "MT-ND4",
        "MT-ND4L",
    ]
    if concordant_mt_up != expected_mt:
        raise RuntimeError(f"Unexpected concordant mtDNA genes: {concordant_mt_up}")
    if other_concordant != ["DBP", "PRELID2"]:
        raise RuntimeError(
            f"Unexpected other concordant genes: {other_concordant}"
        )
    if discordant != ["HIBCH", "ISCU", "TMEM126B"]:
        raise RuntimeError(f"Unexpected discordant genes: {discordant}")
    return facts


def build_divider(prs: Presentation):
    template = prs.slides[28]
    slide = ui.new_slide(prs, bg=ui.NAVY)
    for shape in template.shapes:
        slide.shapes._spTree.insert_element_before(
            copy.deepcopy(shape._element), "p:extLst"
        )
    replacements = {
        "TextBox 1": SECTION_MARKER,
        "TextBox 3": SECTION_TITLE,
        "TextBox 4": (
            "A mitochondrial DEG-program finding, not an identified "
            "upstream key driver."
        ),
        "TextBox 6": "IN THIS FINDING",
        "TextBox 9": "Plain-language statement",
        "TextBox 12": "ROSMAP evidence",
        "TextBox 15": "SEA-AD support",
        "TextBox 18": "Meaning, limits, literature",
        "TextBox 20": (
            "SEA-AD supports the inhibitory-neuron DEG pattern without "
            "matching the exact ROSMAP drivers."
        ),
    }
    for shape_name, value in replacements.items():
        f1.set_shape_text(slide, shape_name, value)
    f1.add_notes(
        slide,
        "Introduce the second biological finding.",
        "Male APOE epsilon-3 homozygous cells show opposite RNA directions across the two genomes that encode oxidative-phosphorylation machinery. The matched inhibitory-neuron program also appears in SEA-AD.",
        "This is a transcript-level mitochondrial program. It does not identify a causal upstream driver, prove protein imbalance, or establish a male-by-APOE interaction.",
        "State the mismatch in plain language.",
    )
    return slide


def build_statement_slide(prs: Presentation):
    slide = ui.new_slide(prs)
    f1.add_content_title(
        slide,
        "Male ε3/ε3 cells raise mtDNA energy-gene RNA while lowering nuclear energy-gene RNA",
        "The two genomes often move in opposite directions; matched inhibitory neurons show the pattern in SEA-AD.",
    )
    f1.add_chip(
        slide,
        "Male",
        1.12,
        1.48,
        2.10,
        bg=ui.PALE_SKY,
        accent=ui.BLUE,
    )
    f1.add_chip(
        slide,
        "APOE ε3/ε3",
        3.52,
        1.48,
        2.42,
        bg=ui.PALE_GREEN,
        accent=ui.TEAL,
    )
    f1.add_chip(
        slide,
        "Inhibitory support",
        6.24,
        1.48,
        2.66,
        bg=ui.PALE_GOLD,
        accent=ui.GOLD,
    )
    f1.add_chip(
        slide,
        "Alzheimer’s disease",
        9.20,
        1.48,
        2.82,
        bg=ui.PALE_RED,
        accent=ui.VERMILION,
    )

    ui.add_rect(slide, 0.82, 2.28, 5.52, 3.02, color=ui.PALE_SKY, outline=None)
    ui.add_text(
        slide,
        "Mitochondrial DNA",
        1.12,
        2.70,
        4.92,
        0.40,
        size=18.0,
        color=ui.NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        "↑",
        1.30,
        3.23,
        1.30,
        1.05,
        size=48.0,
        color=ui.TEAL_TEXT,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        "mtDNA-OXPHOS RNA\nmostly increases in AD",
        2.48,
        3.39,
        3.22,
        0.72,
        size=15.0,
        color=ui.DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    ui.add_rect(slide, 6.98, 2.28, 5.54, 3.02, color=ui.PALE_RED, outline=None)
    ui.add_text(
        slide,
        "Nuclear DNA",
        7.28,
        2.70,
        4.94,
        0.40,
        size=18.0,
        color=ui.NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        "↓",
        7.46,
        3.23,
        1.30,
        1.05,
        size=48.0,
        color=ui.VERMILION_TEXT,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        "nuclear OXPHOS RNA\nmostly decreases in AD",
        8.64,
        3.39,
        3.24,
        0.72,
        size=15.0,
        color=ui.DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    ui.add_rect(slide, 0.82, 5.72, 11.70, 0.76, color=ui.WHITE, outline=ui.LIGHT)
    ui.add_text(
        slide,
        "Plain meaning: RNA instructions for two matching sets of respiratory parts move in opposite directions.",
        1.08,
        5.93,
        11.18,
        0.32,
        size=13.4,
        color=ui.PURPLE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    f1.add_source(
        slide,
        "Finding 2 synthesis; ROSMAP discovery with supplemental SEA-AD program support.",
    )
    f1.add_notes(
        slide,
        "Explain the mitonuclear mismatch without statistical jargon.",
        "In male APOE epsilon-3 homozygous cells, mitochondrial-DNA OXPHOS RNA usually rises in Alzheimer’s disease while nuclear-encoded OXPHOS RNA usually falls. Inhibitory neurons provide the clearest cross-cohort support.",
        "Opposite RNA directions do not demonstrate mismatched proteins, reduced respiration, or unhealthy mitochondria.",
        "Show the ROSMAP numbers behind the statement.",
    )
    return slide


def build_rosmap_slide(prs: Presentation, facts: dict[str, Any]):
    slide = ui.new_slide(prs)
    f1.add_content_title(
        slide,
        "ROSMAP: male ε3/ε3 cells show opposite OXPHOS RNA directions",
        f"Across {facts['ros_eligible_calls']} eligible fine-cell DEG queries; occurrences repeat genes across calls.",
    )

    ui.add_rect(slide, 0.72, 1.52, 5.78, 4.78, color=ui.PALE_SKY, outline=None)
    ui.add_panel_title(
        slide,
        "Mitochondrial-DNA OXPHOS",
        1.04,
        1.86,
        5.10,
        accent=ui.BLUE,
    )
    ui.add_text(
        slide,
        str(facts["ros_mt_occurrences"]),
        1.08,
        2.35,
        1.58,
        0.66,
        size=31.0,
        color=ui.BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        "DEG occurrences",
        2.66,
        2.56,
        2.98,
        0.28,
        size=12.0,
        color=ui.DARK,
        bold=True,
    )
    ui.add_text(
        slide,
        f"{facts['ros_mt_up']} AD-up  •  {facts['ros_mt_down']} AD-down",
        1.12,
        3.30,
        5.00,
        0.30,
        size=13.0,
        color=ui.DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    ui.add_rect(slide, 1.12, 3.80, 4.72, 0.34, color=ui.WHITE, outline=None)
    ui.add_rect(slide, 1.12, 3.80, 4.08, 0.34, color=ui.TEAL, outline=None, radius=False)
    ui.add_rect(slide, 5.20, 3.80, 0.64, 0.34, color=ui.VERMILION, outline=None, radius=False)
    ui.add_text(
        slide,
        "86% up",
        1.18,
        4.23,
        2.02,
        0.24,
        size=10.4,
        color=ui.TEAL_TEXT,
        bold=True,
    )
    ui.add_text(
        slide,
        f"{facts['ros_mt_significant_calls']} input queries enriched",
        1.12,
        5.08,
        5.00,
        0.38,
        size=15.0,
        color=ui.BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    ui.add_text(
        slide,
        "All 12 enriched queries were in the AD-up direction.",
        1.12,
        5.54,
        5.00,
        0.32,
        size=10.6,
        color=ui.GRAY,
        align=PP_ALIGN.CENTER,
    )

    ui.add_rect(slide, 6.82, 1.52, 5.80, 4.78, color=ui.PALE_RED, outline=None)
    ui.add_panel_title(
        slide,
        "Nuclear-encoded OXPHOS",
        7.14,
        1.86,
        5.12,
        accent=ui.VERMILION,
    )
    ui.add_text(
        slide,
        str(facts["ros_nuclear_occurrences"]),
        7.18,
        2.35,
        1.58,
        0.66,
        size=31.0,
        color=ui.VERMILION_TEXT,
        bold=True,
        align=PP_ALIGN.CENTER,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        "DEG occurrences",
        8.76,
        2.56,
        2.98,
        0.28,
        size=12.0,
        color=ui.DARK,
        bold=True,
    )
    ui.add_text(
        slide,
        f"{facts['ros_nuclear_up']} AD-up  •  {facts['ros_nuclear_down']} AD-down",
        7.22,
        3.30,
        5.00,
        0.30,
        size=13.0,
        color=ui.DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    ui.add_rect(slide, 7.22, 3.80, 4.72, 0.34, color=ui.WHITE, outline=None)
    ui.add_rect(slide, 7.22, 3.80, 0.56, 0.34, color=ui.TEAL, outline=None, radius=False)
    ui.add_rect(slide, 7.78, 3.80, 4.16, 0.34, color=ui.VERMILION, outline=None, radius=False)
    ui.add_text(
        slide,
        "88% down",
        9.92,
        4.23,
        2.02,
        0.24,
        size=10.4,
        color=ui.VERMILION_TEXT,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )
    ui.add_text(
        slide,
        f"{facts['ros_nuclear_significant_calls']} input query enriched",
        7.22,
        5.08,
        5.00,
        0.38,
        size=15.0,
        color=ui.VERMILION_TEXT,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    ui.add_text(
        slide,
        "The enriched query was in the AD-down direction.",
        7.22,
        5.54,
        5.00,
        0.32,
        size=10.6,
        color=ui.GRAY,
        align=PP_ALIGN.CENTER,
    )
    ui.add_rect(slide, 0.72, 6.55, 11.90, 0.52, color=ui.PALE_GRAY, outline=None)
    ui.add_text(
        slide,
        "Enriched = more program genes than expected in a KDA input query (BH < 0.05), not the number of key drivers returned.",
        0.98,
        6.70,
        11.38,
        0.24,
        size=10.4,
        color=ui.PURPLE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    f1.add_notes(
        slide,
        "Present the ROSMAP evidence for both genomes.",
        f"Across {facts['ros_eligible_calls']} eligible male epsilon-3 homozygous calls, mtDNA OXPHOS contributes {facts['ros_mt_occurrences']} occurrences: {facts['ros_mt_up']} up and {facts['ros_mt_down']} down. Nuclear OXPHOS contributes {facts['ros_nuclear_occurrences']} occurrences: {facts['ros_nuclear_up']} up and {facts['ros_nuclear_down']} down. Twelve mtDNA input queries and one nuclear input query are enriched.",
        "Occurrences repeat genes across calls and are not independent donors or replications. Enrichment describes the input query rather than returned key drivers.",
        "Ask whether the matched inhibitory-neuron program appears in SEA-AD.",
    )
    return slide


def build_seaad_summary_slide(prs: Presentation, facts: dict[str, Any]):
    slide = ui.new_slide(prs)
    f1.add_content_title(
        slide,
        "SEA-AD supports the matched male ε3/ε3 inhibitory-neuron program",
        f"{facts['sea_evaluable_calls']} eligible SEA-AD fine-cell DEG queries contributed. Support comes from shared MitoCarta MT DEGs and matching AD directions.",
    )
    f1.add_stat_card(
        slide,
        str(facts["sea_shared_genes"]),
        "shared MitoCarta MT DEGs",
        0.72,
        1.52,
        3.78,
        bg=ui.PALE_SKY,
        accent=ui.BLUE,
    )
    f1.add_stat_card(
        slide,
        f"{facts['sea_concordant_genes']}/{facts['sea_shared_genes']}",
        "shared genes have the same AD direction",
        4.78,
        1.52,
        3.78,
        bg=ui.PALE_GREEN,
        accent=ui.TEAL,
    )
    f1.add_stat_card(
        slide,
        f"BH = {facts['cross_cohort_bh']:.4f}",
        "cross-cohort overlap after correction",
        8.84,
        1.52,
        3.78,
        bg=ui.PALE_GOLD,
        accent=ui.GOLD,
        detail=f"22 shared versus {facts['cross_cohort_expected']:.2f} expected",
    )

    ui.add_rect(slide, 0.72, 3.16, 5.78, 2.82, color=ui.PALE_SKY, outline=None)
    ui.add_text(
        slide,
        "Mitochondrial DNA",
        1.06,
        3.54,
        3.58,
        0.38,
        size=17.0,
        color=ui.NAVY,
        bold=True,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        "↑",
        4.78,
        3.42,
        1.08,
        0.82,
        size=38.0,
        color=ui.TEAL_TEXT,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    ui.add_text(
        slide,
        "9 mtDNA-OXPHOS genes are AD-up in both cohorts",
        1.06,
        4.40,
        5.10,
        0.70,
        size=15.0,
        color=ui.DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    ui.add_rect(slide, 6.82, 3.16, 5.80, 2.82, color=ui.PALE_RED, outline=None)
    ui.add_text(
        slide,
        "Nuclear support systems",
        7.16,
        3.54,
        3.78,
        0.38,
        size=17.0,
        color=ui.NAVY,
        bold=True,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        "↓",
        11.04,
        3.42,
        1.08,
        0.82,
        size=38.0,
        color=ui.VERMILION_TEXT,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    ui.add_text(
        slide,
        "OXPHOS, import, ribosome, and maintenance genes are AD-down in both cohorts",
        7.16,
        4.30,
        5.10,
        0.90,
        size=14.0,
        color=ui.DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    ui.add_rect(slide, 0.72, 6.34, 11.90, 0.58, color=ui.WHITE, outline=ui.LIGHT)
    ui.add_text(
        slide,
        f"Identity-resolved sensitivity: {facts['identity_resolved_shared']} shared genes, BH = {facts['identity_resolved_bh']:.4f}.",
        0.98,
        6.50,
        11.38,
        0.26,
        size=10.8,
        color=ui.PURPLE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    f1.add_notes(
        slide,
        "Explain the matched SEA-AD DEG support and what cross-cohort directional agreement means here.",
        f"Eight eligible SEA-AD male epsilon-3 homozygous inhibitory-neuron fine-cell DEG queries contributed to this comparison. ROSMAP and SEA-AD shared {facts['sea_shared_genes']} MitoCarta MT DEGs, and {facts['sea_concordant_genes']} of the {facts['sea_shared_genes']} had the same AD direction. Nine core MT genes were upregulated in both cohorts. Nuclear-encoded OXPHOS, protein-import, mitochondrial-ribosome, and maintenance genes were downregulated in both. The overlap was {facts['sea_shared_genes']} genes compared with {facts['cross_cohort_expected']:.2f} expected and remained significant after correction, with a Benjamini-Hochberg adjusted P value of {facts['cross_cohort_bh']:.4f}. Excluding unresolved mitochondrial identity-conflict genes left {facts['identity_resolved_shared']} shared genes and a similar adjusted P value of {facts['identity_resolved_bh']:.4f}. The support therefore comes from shared DEGs and their directions across cohorts.",
        "This is focused cross-cohort DEG support within the matched inhibitory-neuron context. It does not establish protein-level imbalance, mitochondrial dysfunction, or a formal disease-by-sex-by-APOE interaction.",
        "Show which genes carry the agreement.",
    )
    return slide


def build_seaad_genes_slide(prs: Presentation, facts: dict[str, Any]):
    slide = ui.new_slide(prs)
    f1.add_content_title(
        slide,
        "Which genes carry the matched inhibitory-neuron signal?",
        "The agreement spans respiratory subunits, transport, import, mitochondrial ribosomes, and maintenance.",
    )
    ui.add_rect(slide, 0.72, 1.48, 5.78, 3.94, color=ui.PALE_SKY, outline=None)
    ui.add_panel_title(
        slide,
        "9 mtDNA-OXPHOS genes ↑ in both",
        1.04,
        1.82,
        5.10,
        accent=ui.BLUE,
    )
    for index, gene in enumerate(facts["sea_mt_up_genes"]):
        row, column = divmod(index, 3)
        f1.add_chip(
            slide,
            gene,
            1.04 + column * 1.67,
            2.48 + row * 0.82,
            1.45,
            bg=ui.WHITE,
            accent=ui.BLUE,
            size=10.0,
        )

    ui.add_rect(slide, 6.82, 1.48, 5.80, 3.94, color=ui.PALE_RED, outline=None)
    ui.add_panel_title(
        slide,
        "8 nuclear or maintenance genes ↓ in both",
        7.14,
        1.82,
        5.12,
        accent=ui.VERMILION,
    )
    for index, gene in enumerate(facts["sea_highlighted_down_genes"]):
        row, column = divmod(index, 2)
        f1.add_chip(
            slide,
            gene,
            7.18 + column * 2.48,
            2.34 + row * 0.68,
            2.18,
            bg=ui.WHITE,
            accent=ui.VERMILION,
            size=10.3,
        )

    ui.add_rect(slide, 0.72, 5.72, 5.78, 1.04, color=ui.PALE_GREEN, outline=None)
    ui.add_text(
        slide,
        "Other concordant genes",
        1.02,
        5.96,
        2.22,
        0.26,
        size=11.0,
        color=ui.TEAL_TEXT,
        bold=True,
    )
    ui.add_text(
        slide,
        "DBP ↓   •   PRELID2 ↑",
        3.08,
        5.93,
        3.04,
        0.34,
        size=12.0,
        color=ui.DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    ui.add_rect(slide, 6.82, 5.72, 5.80, 1.04, color=ui.PALE_GOLD, outline=None)
    ui.add_text(
        slide,
        "3 discordant genes",
        7.12,
        5.96,
        1.94,
        0.26,
        size=11.0,
        color=ui.GOLD_TEXT,
        bold=True,
    )
    ui.add_text(
        slide,
        "HIBCH   •   ISCU   •   TMEM126B",
        8.94,
        5.93,
        3.34,
        0.34,
        size=10.8,
        color=ui.DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    f1.add_notes(
        slide,
        "Name the genes underlying the cross-cohort directional agreement.",
        "Nine mtDNA OXPHOS genes rise in both cohorts. Eight highlighted nuclear or mitochondrial-support genes fall in both cohorts: UQCRFS1, SLC25A5, TIMM13, TIMM17A, MRPL16, MRPS34, ENDOG, and MTCH1. DBP and PRELID2 complete the nineteen concordant genes.",
        "HIBCH, ISCU, and TMEM126B point in opposite directions across cohorts. Gene agreement still does not demonstrate protein-complex imbalance.",
        "Explain why a two-genome mismatch may matter.",
    )
    return slide


def build_why_slide(prs: Presentation):
    slide = ui.new_slide(prs)
    f1.add_content_title(
        slide,
        "Why this matters: one respiratory system depends on two genomes",
        "Opposite RNA directions could challenge coordination, but the functional outcome remains unknown.",
    )
    cards = [
        (
            "1",
            "Two-genome machine",
            "OXPHOS complexes combine proteins encoded by mitochondrial DNA and nuclear DNA.",
            ui.PALE_SKY,
            ui.BLUE,
        ),
        (
            "2",
            "Support systems also fall",
            "Import, mitochondrial-ribosome, transport, and maintenance genes are down in both cohorts.",
            ui.PALE_RED,
            ui.VERMILION,
        ),
        (
            "3",
            "Possible cellular strain",
            "The pattern could reflect compensation, unassembled parts, altered cell state, or mitochondrial stress.",
            ui.PALE_GOLD,
            ui.GOLD,
        ),
    ]
    for index, (number, title, detail, bg, accent) in enumerate(cards):
        x = 0.72 + index * 4.15
        ui.add_rect(slide, x, 1.58, 3.78, 3.20, color=bg, outline=None)
        ui.add_circle(slide, x + 0.28, 1.92, 0.56, accent)
        ui.add_text(
            slide,
            number,
            x + 0.28,
            2.05,
            0.56,
            0.24,
            size=13.0,
            color=ui.WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        ui.add_text(
            slide,
            title,
            x + 1.00,
            1.94,
            2.46,
            0.48,
            size=15.0,
            color=ui.NAVY,
            bold=True,
            font=ui.FONT_HEAD,
        )
        ui.add_text(
            slide,
            detail,
            x + 0.32,
            2.78,
            3.14,
            1.28,
            size=11.6,
            color=ui.DARK,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
    ui.add_rect(slide, 0.72, 5.18, 11.90, 1.14, color=ui.PALE_RED, outline=None)
    ui.add_text(
        slide,
        "Most cautious interpretation",
        0.98,
        5.46,
        2.40,
        0.30,
        size=12.0,
        color=ui.VERMILION_TEXT,
        bold=True,
    )
    ui.add_text(
        slide,
        "This is an RNA-level mitonuclear mismatch—not proof that respiratory proteins or ATP production are unbalanced.",
        3.55,
        5.34,
        8.57,
        0.56,
        size=13.6,
        color=ui.DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    f1.add_notes(
        slide,
        "Explain the biological importance without claiming impaired function.",
        "Respiratory complexes require matching parts encoded by two genomes. The cross-cohort decrease in import, mitochondrial-ribosome, transport, and maintenance genes makes the pattern broader than a single OXPHOS gene change.",
        "RNA abundance cannot distinguish compensation, failed assembly, cell-state changes, or altered survival of cell populations.",
        "Separate the observation from what still needs to be tested.",
    )
    return slide


def build_boundaries_slide(prs: Presentation):
    slide = ui.new_slide(prs)
    f1.add_content_title(
        slide,
        "What the finding does not prove—and the next evidence needed",
        "A transcript-level mismatch is a testable hypothesis, not a completed mechanism.",
    )
    ui.add_rect(slide, 0.72, 1.42, 5.78, 4.92, color=ui.PALE_RED, outline=None)
    ui.add_panel_title(
        slide,
        "Not proven",
        1.04,
        1.76,
        5.10,
        accent=ui.VERMILION,
    )
    ui.add_bullets(
        slide,
        [
            "Physical imbalance of OXPHOS proteins or complexes",
            "Lower ATP, respiration, or higher oxidative stress",
            "A male- or ε3/ε3-specific disease interaction",
            "Which upstream key driver caused the pattern",
        ],
        1.04,
        2.32,
        5.00,
        size=12.4,
        accent=ui.VERMILION,
        line_h=0.82,
    )
    ui.add_rect(slide, 6.82, 1.42, 5.80, 4.92, color=ui.PALE_GREEN, outline=None)
    ui.add_panel_title(
        slide,
        "Best next tests",
        7.14,
        1.76,
        5.12,
        accent=ui.TEAL,
    )
    ui.add_bullets(
        slide,
        [
            "Donor-aware disease × sex × APOE interaction model",
            "Paired mtDNA- and nuclear-encoded OXPHOS proteins",
            "Respiratory-complex assembly, oxygen use, ATP, and ROS",
            "Perturb candidate drivers and test functional rescue",
        ],
        7.14,
        2.32,
        5.02,
        size=12.4,
        accent=ui.TEAL,
        line_h=0.82,
    )
    ui.add_text(
        slide,
        "SEA-AD strengthens the program observation but cannot replace interaction tests or functional measurements.",
        0.92,
        6.62,
        11.50,
        0.30,
        size=11.2,
        color=ui.PURPLE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    f1.add_notes(
        slide,
        "State the scientific boundaries and next experiments.",
        "Donor-aware interaction models are needed to test whether the mismatch differs by sex or APOE group. Paired protein, complex-assembly, respiration, ATP, and oxidative-stress assays are needed to determine functional consequences.",
        "SEA-AD supplies program-level support but not exact driver replication or causal evidence.",
        "Place the finding in prior literature.",
    )
    return slide


def build_literature_slide(prs: Presentation):
    slide = ui.new_slide(prs)
    f1.add_content_title(
        slide,
        "Prior research supports the ingredients—but not this exact mismatch",
        "Known mitonuclear biology and cell-resolved AD studies make the finding plausible without independently duplicating it.",
    )
    cards = [
        (
            "Mitonuclear biology",
            "Two-genome assembly",
            "OXPHOS requires coordinated mitochondrial- and nuclear-encoded parts.",
            "Does not show protein imbalance in these samples.",
            ui.PALE_SKY,
            ui.BLUE,
        ),
        (
            "Guo et al. 2023",
            "Sex-aware AD networks",
            "Supports analyzing Alzheimer’s biology separately by sex.",
            "Does not test this exact male ε3/ε3 mismatch.",
            ui.PALE_GREEN,
            ui.TEAL,
        ),
        (
            "Mathys et al. 2024",
            "Cell-resolved AD responses",
            "Supports strong differences among brain cell populations.",
            "Uses ROSMAP, so it is not independent replication.",
            ui.PALE_GOLD,
            ui.GOLD,
        ),
    ]
    for index, (citation, title, support, limit, bg, accent) in enumerate(cards):
        x = 0.72 + index * 4.15
        ui.add_rect(slide, x, 1.48, 3.78, 4.56, color=bg, outline=None)
        ui.add_text(
            slide,
            citation,
            x + 0.30,
            1.80,
            3.18,
            0.30,
            size=10.2,
            color=ui.readable_accent(accent),
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        ui.add_text(
            slide,
            title,
            x + 0.30,
            2.44,
            3.18,
            0.46,
            size=14.0,
            color=ui.NAVY,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        ui.add_text(
            slide,
            "SUPPORTS",
            x + 0.32,
            3.20,
            1.10,
            0.24,
            size=9.2,
            color=ui.readable_accent(accent),
            bold=True,
        )
        ui.add_text(
            slide,
            support,
            x + 0.32,
            3.52,
            3.14,
            0.72,
            size=11.0,
            color=ui.DARK,
            align=PP_ALIGN.CENTER,
        )
        ui.add_text(
            slide,
            "BOUNDARY",
            x + 0.32,
            4.48,
            1.10,
            0.24,
            size=9.2,
            color=ui.GRAY,
            bold=True,
        )
        ui.add_text(
            slide,
            limit,
            x + 0.32,
            4.80,
            3.14,
            0.72,
            size=10.7,
            color=ui.GRAY,
            align=PP_ALIGN.CENTER,
        )
    ui.add_rect(slide, 0.72, 6.28, 11.90, 0.58, color=ui.WHITE, outline=ui.LIGHT)
    ui.add_text(
        slide,
        "Novelty assessment: high—mitonuclear imbalance is known, but this cross-cohort male ε3/ε3 inhibitory-neuron pattern was not identified previously.",
        0.98,
        6.40,
        11.38,
        0.34,
        size=10.4,
        color=ui.PURPLE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    f1.add_notes(
        slide,
        "Connect the finding to prior research while preserving novelty boundaries.",
        "Established mitonuclear biology explains why opposite directions across the two genomes could matter. Guo and colleagues support sex-aware Alzheimer’s network analysis. Mathys and colleagues support cell-resolved disease heterogeneity but use ROSMAP.",
        "None of these sources independently establishes the exact male epsilon-3 homozygous inhibitory-neuron pattern. No sex-stratified human-genetic claim is used here.",
        "Return to the broader finding list or proceed to Finding 3.",
    )
    return slide


def build_slides(prs: Presentation, facts: dict[str, Any]) -> list[Any]:
    return [
        build_divider(prs),
        build_statement_slide(prs),
        build_rosmap_slide(prs, facts),
        build_seaad_summary_slide(prs, facts),
        build_seaad_genes_slide(prs, facts),
        build_why_slide(prs),
        build_boundaries_slide(prs),
        build_literature_slide(prs),
    ]


def write_audit(
    audit_root: Path,
    checks: list[dict[str, Any]],
    facts: dict[str, Any],
    output_path: Path,
    backup_path: Path,
) -> None:
    audit_root.mkdir(parents=True, exist_ok=True)
    with (audit_root / "finding2_slide_checks.tsv").open(
        "w", encoding="utf-8", newline=""
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["check_id", "observed", "expected", "passed"],
            delimiter="\t",
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(checks)
    with (audit_root / "finding2_slide_facts.tsv").open(
        "w", encoding="utf-8", newline=""
    ) as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(["field", "value"])
        for key, value in facts.items():
            if isinstance(value, list):
                value = ";".join(map(str, value))
            writer.writerow([key, value])
        writer.writerow(["output_deck", output_path.resolve()])
        writer.writerow(["output_sha256", f1.sha256_file(output_path)])
        writer.writerow(["backup_deck", backup_path.resolve()])


def main() -> int:
    args = parse_args()
    input_path = args.input.resolve()
    output_path = args.output.resolve()
    audit_root = args.audit_root.resolve()
    if not input_path.is_file():
        raise FileNotFoundError(input_path)

    facts = prepare_facts()
    prs = Presentation(input_path)
    if len(prs.slides) != EXPECTED_INPUT_SLIDES:
        raise RuntimeError(
            f"Expected {EXPECTED_INPUT_SLIDES} slides, found {len(prs.slides)}"
        )
    if SECTION_TITLE in "\n".join(f1.slide_text(slide) for slide in prs.slides):
        raise RuntimeError("Finding 2 slides already exist; refusing to append twice")

    before = [f1.semantic_slide_fingerprint(slide) for slide in prs.slides]
    input_hash = f1.sha256_file(input_path)
    backup_dir = audit_root / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / f"input_deck_{input_hash[:12]}.pptx"
    if not backup_path.exists():
        shutil.copy2(input_path, backup_path)

    ui.set_notes_body_template(
        prs.slides[0].notes_slide.notes_placeholder._element
    )
    build_slides(prs, facts)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}.",
        suffix=".pptx",
        dir=output_path.parent,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        prs.save(temporary)
        reloaded = Presentation(temporary)
        after = [
            f1.semantic_slide_fingerprint(slide) for slide in reloaded.slides
        ]
        inserted = list(reloaded.slides)[INSERT_INDEX:]
        inserted_text = [f1.slide_text(slide) for slide in inserted]
        combined = "\n".join(inserted_text)
        previous_slides_unchanged = after[:INSERT_INDEX] == before
        all_new_have_notes = all(
            slide.notes_slide.notes_text_frame
            and slide.notes_slide.notes_text_frame.text.strip()
            for slide in inserted
        )
        all_shapes_in_bounds = all(
            0 <= shape.left
            and 0 <= shape.top
            and shape.left + shape.width <= reloaded.slide_width
            and shape.top + shape.height <= reloaded.slide_height
            for slide in inserted
            for shape in slide.shapes
        )
        required_text = [
            SECTION_MARKER,
            "mtDNA-up",
            "Plain meaning",
            "81",
            "70 AD-up",
            "60 AD-down",
            "BH = 0.0417",
            "19/22",
            "MT-ND4L",
            "TIMM17A",
            "HIBCH",
            "Why this matters",
            "Not proven",
            "Guo et al. 2023",
        ]
        checks = [
            {
                "check_id": "output_slide_count",
                "observed": len(reloaded.slides),
                "expected": EXPECTED_INPUT_SLIDES + NEW_SLIDE_COUNT,
                "passed": len(reloaded.slides)
                == EXPECTED_INPUT_SLIDES + NEW_SLIDE_COUNT,
            },
            {
                "check_id": "finding2_block_appended_after_finding1",
                "observed": (
                    SECTION_MARKER in inserted_text[0]
                    and "Prior research makes the result plausible"
                    in f1.slide_text(reloaded.slides[INSERT_INDEX - 1])
                ),
                "expected": True,
                "passed": (
                    SECTION_MARKER in inserted_text[0]
                    and "Prior research makes the result plausible"
                    in f1.slide_text(reloaded.slides[INSERT_INDEX - 1])
                ),
            },
            {
                "check_id": "all_required_content_present",
                "observed": all(text in combined for text in required_text),
                "expected": True,
                "passed": all(text in combined for text in required_text),
            },
            {
                "check_id": "all_new_slides_have_notes",
                "observed": all_new_have_notes,
                "expected": True,
                "passed": all_new_have_notes,
            },
            {
                "check_id": "all_new_shapes_in_slide_bounds",
                "observed": all_shapes_in_bounds,
                "expected": True,
                "passed": all_shapes_in_bounds,
            },
            {
                "check_id": "all_previous_slides_semantically_unchanged",
                "observed": previous_slides_unchanged,
                "expected": True,
                "passed": previous_slides_unchanged,
            },
        ]
        failed = [row["check_id"] for row in checks if not row["passed"]]
        if failed:
            raise RuntimeError(
                "Finding 2 slide checks failed: " + ", ".join(failed)
            )
        os.replace(temporary, output_path)
    finally:
        if temporary.exists():
            temporary.unlink()

    write_audit(audit_root, checks, facts, output_path, backup_path)
    print(f"Updated deck: {output_path}")
    print(
        f"Inserted slides: {INSERT_INDEX + 1}–"
        f"{INSERT_INDEX + NEW_SLIDE_COUNT}"
    )
    print(f"Backup: {backup_path}")
    print(f"Audit: {audit_root}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
