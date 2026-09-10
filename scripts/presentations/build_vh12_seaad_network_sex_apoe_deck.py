#!/usr/bin/env python3
"""Build the dated sex/APOE KDA deck with VH12 SEA-AD-network results.

Slides 1-10 are retained from the 2026-08-31 deck (with the SEA-AD summary
on slides 1-2 updated). Part II and Part III are rebuilt from the validated
VH12 result tables and canonical Phase 12 figures. The source deck is never
modified.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import sys
import tempfile
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
import update_phase11_seaad_simple_aggr_part2 as ui  # noqa: E402


SOURCE_DECK = (
    ROOT
    / "docs/presentations/08312026/sex_apoe_kda_fine_broad_08312026.pptx"
)
DEFAULT_OUTPUT = (
    ROOT
    / "docs/presentations/09162026/sex_apoe_kda_fine_broad_09162026.pptx"
)
AUDIT_DIR = ROOT / "results/presentations/vh12_seaad_network_sex_apoe_deck"

SEA_RESULT = ROOT / "results/validation_human/12_sex_apoe_kda_simple_aggr_network"
SEA_MANIFEST = SEA_RESULT / "10a_inputs/seaad_kda_run_manifest.tsv"
SEA_STATUS = SEA_RESULT / "simple_status.tsv"
SEA_CATEGORIES = SEA_RESULT / "simple_category_gene_aggregates.tsv"
SEA_SUMMARY = SEA_RESULT / "simple_category_summary.tsv"
SEA_FIGURE_ROOT = (
    ROOT
    / "results/figures/validation_human/phase_12_sex_apoe_simple_aggr_network"
)
SEA_TOP5 = (
    SEA_FIGURE_ROOT
    / "top5_candidates/phase12_seaad_simple_aggr_network_top5_candidates.png"
)
SEA_TOP5_DATA = (
    SEA_FIGURE_ROOT
    / "top5_candidates/phase12_seaad_simple_aggr_network_top5_candidates_plot_data.tsv"
)
SEA_RECURRENCE = (
    SEA_FIGURE_ROOT
    / "driver_recurrence/phase12_seaad_simple_aggr_network_driver_recurrence.png"
)
SEA_RECURRENCE_DATA = (
    SEA_FIGURE_ROOT
    / "driver_recurrence/phase12_seaad_simple_aggr_network_driver_recurrence_plot_data.tsv"
)
ROS_CATEGORIES = (
    ROOT
    / "results/minerva_production/20_sex_apoe_kda_simple_aggr/simple_category_gene_aggregates.tsv"
)

GROUP_ORDER = ["F_e2", "F_e33", "F_e4", "M_e2", "M_e33", "M_e4"]
NETWORK_LABELS = {
    "Astrocytes": "Astrocytes",
    "Excitatory_neurons": "Excitatory neurons",
    "Inhibitory_neurons": "Inhibitory neurons",
    "Microglia": "Microglia",
    "OPCs": "OPCs",
    "Oligodendrocytes": "Oligodendrocytes",
    "Vasculature_cells": "Vasculature",
}
EXPECTED_TITLES = [
    "ROSMAP and SEAAD sex/APOE key-driver analysis",
    "Two cohorts, one aggregation rule",
    "PART 1",
    "Four steps",
    "54 fine cell types create 648 planned directional slots",
    "Source validity and query size decide whether KDA is called",
    "Mitochondrial query genes",
    "Top-five: 149 non-MT entries across 32 categories",
    "RPS15 recurs across 12 returned-only categories",
    "Simple output: 689 non-MT category units represent 433 genes",
    "PART 2",
    "129 SEA-AD supertypes create 1,548 planned directional slots",
    "Top-five display: 20 non-MT entries across 5 categories",
    "PAPOLA and PJVK each recur across 2 SEA-AD categories",
    "SEA-AD output: 96 non-MT category units represent 94 genes",
    "PART 3",
    "Two ways to ask whether a driver reappears",
    "10 of 94 SEA-AD drivers also return in ROSMAP",
    "One driver-category pair matches exactly",
    "MIPOL1 is the sole exact context match",
    "Shared genes are not concentrated in matching M_e33",
    "Neither cell type nor sex/APOE context is strongly preserved",
    "Shared SEA-AD male-e33 genes span both ROSMAP sexes",
    "What the SEA-AD-network rerun does—and does not—show",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=SOURCE_DECK)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--audit-dir", type=Path, default=AUDIT_DIR)
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def write_tsv(rows: Iterable[dict[str, Any]], path: Path, fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=fields, delimiter="\t", lineterminator="\n"
        )
        writer.writeheader()
        writer.writerows(rows)


def slide_text(slide) -> str:
    return "\n".join(
        shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
    )


def replace_shape_text(slide, shape_name: str, old: str, new: str) -> None:
    matches = [shape for shape in slide.shapes if shape.name == shape_name]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one shape named {shape_name!r}")
    shape = matches[0]
    if shape.text != old:
        raise RuntimeError(
            f"Text contract failed for {shape_name}: {shape.text!r} != {old!r}"
        )
    runs = [run for p in shape.text_frame.paragraphs for run in p.runs]
    if len(runs) == 1:
        runs[0].text = new
    else:
        shape.text_frame.text = new


def drop_slides_after(prs: Presentation, keep: int) -> None:
    while len(prs.slides) > keep:
        slide_id = prs.slides._sldIdLst[-1]
        prs.part.drop_rel(slide_id.rId)
        del prs.slides._sldIdLst[-1]


def load_non_mt(path: Path) -> pd.DataFrame:
    frame = pd.read_csv(
        path, sep="\t", dtype=str, keep_default_na=False, na_values=["NA"]
    )
    frame = frame[
        frame["case_id"].eq("non_mt_driver")
        & ~frame["is_core_mito"].isin(["TRUE"])
    ].copy()
    frame["score"] = frame["returned_run_q_acat_score"].astype(float)
    frame["calls"] = frame["returned_call_count"].astype(int)
    frame = frame.sort_values(
        ["signature_group", "broad_network", "score", "current_symbol"],
        kind="mergesort",
    )
    frame["display_rank"] = (
        frame.groupby(["signature_group", "broad_network"]).cumcount() + 1
    )
    return frame


def load_facts() -> dict[str, Any]:
    required = [
        SOURCE_DECK,
        SEA_MANIFEST,
        SEA_STATUS,
        SEA_CATEGORIES,
        SEA_SUMMARY,
        SEA_TOP5,
        SEA_TOP5_DATA,
        SEA_RECURRENCE,
        SEA_RECURRENCE_DATA,
        ROS_CATEGORIES,
    ]
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing deck input(s): " + ", ".join(missing))

    status_rows = read_tsv(SEA_STATUS)
    if len(status_rows) != 1:
        raise RuntimeError("VH12 simple_status.tsv must contain one row")
    status = status_rows[0]
    if (
        status["analysis_id"]
        != "seaad_network_simple_returned_only_non_core_mt_acat_v1"
        or status["execution_status"] != "complete"
        or status["failed_check_count"] != "0"
    ):
        raise RuntimeError("VH12 simple aggregation is not validated complete")

    for phase in ("10a_inputs", "10b_kda"):
        phase_status = read_tsv(SEA_RESULT / phase / "status.tsv")
        if len(phase_status) != 1 or phase_status[0]["validation_status"] != "validated_complete":
            raise RuntimeError(f"{phase} is not validated complete")
    for figure_id in ("top5_candidates", "driver_recurrence"):
        stem = f"phase12_seaad_simple_aggr_network_{figure_id}"
        fig_status = read_tsv(SEA_FIGURE_ROOT / figure_id / f"{stem}_status.tsv")
        fig_checks = read_tsv(SEA_FIGURE_ROOT / figure_id / f"{stem}_checks.tsv")
        if (
            len(fig_status) != 1
            or fig_status[0]["validation_status"] != "validated_complete"
            or fig_status[0]["failed_checks"] != "0"
            or any(row["passed"].upper() != "TRUE" for row in fig_checks)
        ):
            raise RuntimeError(f"Figure source is not validated: {figure_id}")

    manifest = read_tsv(SEA_MANIFEST)
    eligible = [row for row in manifest if row["eligibility_status"] == "eligible"]
    terminal = Counter(row["terminal_status"] for row in manifest)
    eligible_groups = Counter(row["signature_group"] for row in eligible)

    sea = load_non_mt(SEA_CATEGORIES)
    ros = load_non_mt(ROS_CATEGORIES)
    recurrence = read_tsv(SEA_RECURRENCE_DATA)
    top5 = read_tsv(SEA_TOP5_DATA)
    categories = sorted(
        set(zip(sea["signature_group"], sea["broad_network"], strict=True))
    )
    group_counts = Counter(sea["signature_group"])
    network_counts = Counter(sea["broad_network"])

    shared_genes = sorted(set(sea["current_symbol"]) & set(ros["current_symbol"]))
    matched = sea.merge(
        ros,
        on=["signature_group", "broad_network", "current_symbol"],
        suffixes=("_sea", "_ros"),
    ).sort_values(["signature_group", "broad_network", "score_sea"])
    pairs = sea.merge(ros, on="current_symbol", suffixes=("_sea", "_ros"))
    same_network = pairs["broad_network_sea"].eq(pairs["broad_network_ros"])
    same_group = pairs["signature_group_sea"].eq(pairs["signature_group_ros"])
    ros_net = ros["broad_network"].value_counts(normalize=True)
    sea_net = sea["broad_network"].value_counts(normalize=True)
    chance_network = sum(
        ros_net.get(network, 0.0) * sea_net.get(network, 0.0)
        for network in set(ros_net.index) | set(sea_net.index)
    )

    sea_m_genes = set(
        sea.loc[sea["signature_group"].eq("M_e33"), "current_symbol"]
    )
    sea_m_shared = sea_m_genes & set(ros["current_symbol"])
    group_matrix = {
        group: len(
            sea_m_shared
            & set(ros.loc[ros["signature_group"].eq(group), "current_symbol"])
        )
        for group in GROUP_ORDER
    }
    sex_split: dict[str, list[str]] = {
        "male_only": [],
        "female_only": [],
        "both": [],
    }
    for gene in sorted(sea_m_shared):
        sexes = set(
            ros.loc[ros["current_symbol"].eq(gene), "signature_group"].str[0]
        )
        key = "male_only" if sexes == {"M"} else "female_only" if sexes == {"F"} else "both"
        sex_split[key].append(gene)

    values: dict[str, Any] = {
        "status": status,
        "manifest": manifest,
        "planned_slots": len(manifest),
        "supertypes": len({row["supertype_id"] for row in manifest}),
        "not_estimable": terminal["source_contrast_not_estimable"],
        "query_empty": terminal["query_empty"],
        "query_below": terminal["query_below_minimum"],
        "small_calls": terminal["eligible_small_query"],
        "large_calls": terminal["eligible_phase18_sized"],
        "active_calls": len(eligible),
        "eligible_groups": eligible_groups,
        "significant_calls": int(status["completed_significant_call_count"]),
        "empty_calls": int(status["completed_empty_call_count"]),
        "stock_rows": int(status["all_class_stock_returned_row_count"]),
        "mt_excluded": int(status["mt_excluded_returned_row_count"]),
        "non_mt_rows": int(status["non_mt_retained_returned_row_count"]),
        "sea": sea,
        "ros": ros,
        "sea_units": len(sea),
        "sea_genes": sea["current_symbol"].nunique(),
        "ros_units": len(ros),
        "ros_genes": ros["current_symbol"].nunique(),
        "categories": categories,
        "category_count": len(categories),
        "group_counts": group_counts,
        "network_counts": network_counts,
        "recurrence": recurrence,
        "top5": top5,
        "top5_count": len(top5),
        "shared_genes": shared_genes,
        "matched": matched,
        "pairs": pairs,
        "pair_count": len(pairs),
        "same_network_count": int(same_network.sum()),
        "same_group_count": int(same_group.sum()),
        "chance_network": float(chance_network),
        "sea_m_gene_count": len(sea_m_genes),
        "sea_m_shared_count": len(sea_m_shared),
        "group_matrix": group_matrix,
        "sex_split": sex_split,
    }

    expected = {
        "planned_slots": 1548,
        "supertypes": 129,
        "not_estimable": 786,
        "query_empty": 708,
        "query_below": 20,
        "small_calls": 15,
        "large_calls": 19,
        "active_calls": 34,
        "significant_calls": 29,
        "empty_calls": 5,
        "stock_rows": 221,
        "mt_excluded": 86,
        "non_mt_rows": 135,
        "sea_units": 96,
        "sea_genes": 94,
        "ros_units": 689,
        "ros_genes": 433,
        "category_count": 5,
        "top5_count": 20,
        "pair_count": 22,
        "same_network_count": 3,
        "same_group_count": 4,
        "sea_m_gene_count": 89,
        "sea_m_shared_count": 10,
    }
    for key, expected_value in expected.items():
        if values[key] != expected_value:
            raise RuntimeError(
                f"VH12 source drift for {key}: {values[key]} != {expected_value}"
            )
    if eligible_groups != Counter({"M_e33": 32, "F_e33": 1, "F_e4": 1}):
        raise RuntimeError(f"Active-call group distribution drifted: {eligible_groups}")
    if shared_genes != [
        "CHD3", "CTNND2", "CXXC1", "DEF8", "LAGE3", "MIPOL1", "MKLN1",
        "PAPOLA", "SECISBP2", "SYN1",
    ]:
        raise RuntimeError(f"Cross-cohort shared-gene set drifted: {shared_genes}")
    if len(matched) != 1 or matched.iloc[0]["current_symbol"] != "MIPOL1":
        raise RuntimeError("Expected MIPOL1 to be the sole exact category match")
    if group_matrix != {
        "F_e2": 3, "F_e33": 2, "F_e4": 7,
        "M_e2": 1, "M_e33": 4, "M_e4": 3,
    }:
        raise RuntimeError(f"ROSMAP group overlap matrix drifted: {group_matrix}")
    if {key: len(value) for key, value in sex_split.items()} != {
        "male_only": 2, "female_only": 4, "both": 4,
    }:
        raise RuntimeError(f"Sex-context split drifted: {sex_split}")
    if [row["current_symbol"] for row in recurrence[:2]] != ["PAPOLA", "PJVK"]:
        raise RuntimeError("VH12 recurrence leaders drifted")
    return values


def update_overview_slides(prs: Presentation, facts: dict[str, Any]) -> None:
    replace_shape_text(
        prs.slides[0],
        "TextBox 11",
        "42 active KDA calls → returned-only non-MT ACAT → 96 category units",
        "34 active SEA-AD-network KDA calls → returned-only non-MT ACAT → 96 category units",
    )
    replace_shape_text(prs.slides[1], "TextBox 29", "42", "34")
    replace_shape_text(
        prs.slides[1],
        "TextBox 39",
        "40 of 42 calls are M_e33; 4 categories have returns.",
        "32 of 34 calls are M_e33; 5 categories have returns.",
    )
    ui.add_notes(
        prs.slides[0],
        goal="Introduce the two cohorts and the matched returned-only aggregation rule.",
        walkthrough="Part I summarizes the existing ROSMAP fine-cell reaggregation. Part II uses 34 newly rerun SEA-AD calls against SEA-AD-specific Bayesian networks and yields 96 non-MT gene-by-category units.",
        boundary="The SEA-AD rerun changes network topology as well as cohort, so differences from ROSMAP or the prior deck are not attributable to cohort biology alone.",
        transition="Compare the cohort and run structures before reviewing each analysis branch.",
    )
    ui.add_notes(
        prs.slides[1],
        goal="Give a concise side-by-side map of the ROSMAP and VH12 SEA-AD analyses.",
        walkthrough="ROSMAP contributes 295 fine-cell runs and 689 units. The VH12 SEA-AD rerun contributes 34 active calls and 96 units from 94 genes, using the same returned-only non-MT ACAT aggregation recipe.",
        boundary="Thirty-two of 34 SEA-AD calls are M_e33. Categories without calls are unavailable, not completed null tests.",
        transition="Review the ROSMAP branch first, then the network-specific SEA-AD rerun.",
    )


def append_part2_divider(prs: Presentation, facts: dict[str, Any]) -> None:
    slide = ui.new_slide(prs, bg=ui.NAVY)
    ui.add_text(slide, "PART 2", 0.78, 0.67, 2.4, 0.28, size=10.5, color=ui.TEAL, bold=True)
    ui.add_rect(slide, 0.78, 1.26, 0.10, 2.38, color=ui.TEAL, outline=None, radius=False)
    ui.add_text(
        slide, "SEA-AD network-specific KDA rerun",
        1.18, 1.36, 8.65, 1.34, size=32, color=ui.WHITE, bold=True,
        font=ui.FONT_HEAD, valign=MSO_ANCHOR.MIDDLE,
    )
    ui.add_text(
        slide,
        f"{facts['supertypes']} supertypes define the planned universe; "
        f"{facts['active_calls']} calls were rerun against SEA-AD-specific Bayesian networks.",
        1.20, 3.00, 8.30, 0.84, size=15.0, color=ui.RGBColor(204, 219, 234),
    )
    ui.add_rect(slide, 9.78, 1.28, 2.70, 4.70, color=ui.NAVY_2, outline=None)
    ui.add_text(slide, "IN THIS PART", 10.08, 1.67, 2.10, 0.24, size=9.4, color=ui.TEAL, bold=True)
    for index, topic in enumerate(
        ["Run universe", "Top-five results", "Driver recurrence", "Result scale"], start=1
    ):
        y = 2.20 + (index - 1) * 0.82
        ui.add_circle(slide, 10.06, y + 0.03, 0.34, ui.TEAL)
        ui.add_text(slide, str(index), 10.06, y + 0.08, 0.34, 0.17, size=8.2, color=ui.NAVY, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        ui.add_text(slide, topic, 10.53, y, 1.58, 0.52, size=11.1, color=ui.WHITE, bold=True)
    ui.add_notes(
        slide,
        goal="Introduce the VH12 SEA-AD-specific-network rerun.",
        walkthrough="The section moves from the full supertype-by-stratum universe to the 34 executable KDA calls, then shows the top-five, recurrence, and aggregate summaries.",
        boundary="The aggregation is exploratory and returned-only; the underlying KDA calls were rerun rather than borrowed from the prior network release.",
        transition="Start with how 1,548 planned directional slots become 34 active calls.",
    )


def append_run_universe(prs: Presentation, facts: dict[str, Any]) -> None:
    slide = ui.new_slide(prs)
    ui.add_title_block(
        slide,
        f"{facts['supertypes']} SEA-AD supertypes create {facts['planned_slots']:,} planned directional slots",
        "Supertype × sex/APOE contrast × AD-up/AD-down mitochondrial query; network mapping determines executable calls.",
    )
    cards = [
        (f"{facts['supertypes']} supertypes", "× 6 groups = 774 contrasts", ui.PALE_GREEN, ui.TEAL),
        ("774 contrasts", f"× up/down = {facts['planned_slots']:,} slots", ui.PALE_SKY, ui.SKY),
        (f"{facts['active_calls']} KDA calls", "effective mapped query n ≥ 3", ui.PALE_GOLD, ui.GOLD),
    ]
    for i, (value, label, bg, accent) in enumerate(cards):
        x = 0.70 + i * 4.12
        ui.add_rect(slide, x, 1.47, 3.70, 1.31, color=bg, outline=accent)
        ui.add_text(slide, value, x + 0.25, 1.74, 3.20, 0.31, size=19, color=ui.NAVY, bold=True, align=PP_ALIGN.CENTER)
        ui.add_text(slide, label, x + 0.25, 2.18, 3.20, 0.30, size=12.3, color=ui.readable_accent(accent), bold=True, align=PP_ALIGN.CENTER)
    ui.add_table(
        slide,
        ["Mutually exclusive slot outcome", "Slots", "KDA called?"],
        [
            ["Source DEG contrast not estimable", str(facts["not_estimable"]), "No"],
            ["Estimable source; FDR-only query = 0", str(facts["query_empty"]), "No"],
            ["Effective mapped query 1–2", str(facts["query_below"]), "No"],
            ["Effective mapped query 3–9", str(facts["small_calls"]), "Yes (small query)"],
            ["Effective mapped query ≥10", str(facts["large_calls"]), "Yes"],
            ["TOTAL", f"{facts['planned_slots']:,}", f"{facts['active_calls']} calls"],
        ],
        0.86, 3.10, [6.35, 1.22, 3.95], row_h=0.47, header_h=0.47,
        font_size=9.6, highlight_last=True,
    )
    ui.add_text(
        slide,
        "Availability remains unbalanced: 32 M_e33 calls, one F_e33 call, and one F_e4 call.",
        0.88, 6.52, 11.60, 0.28, size=11.0, color=ui.PURPLE, bold=True, align=PP_ALIGN.CENTER,
    )
    ui.add_notes(
        slide,
        goal="Show exactly how the planned SEA-AD universe reduces to 34 executable calls.",
        walkthrough="Of 1,548 planned slots, 786 have a non-estimable source contrast, 708 have no FDR-only mitochondrial query, and 20 retain only one or two mapped genes. Fifteen calls have 3–9 genes and 19 have at least 10.",
        boundary="A skipped slot has no KDA result and is not a completed null. Network-specific mapping changes which DEG queries meet the minimum size.",
        transition="Inspect the returned non-mitochondrial candidates from the completed calls.",
    )


def append_top5(prs: Presentation, facts: dict[str, Any]) -> None:
    slide = ui.new_slide(prs)
    ui.add_title_block(
        slide,
        f"Top-five display: {facts['top5_count']} non-MT entries across {facts['category_count']} categories",
    )
    ui.add_picture_contain(
        slide, SEA_TOP5, 0.62, 0.98, 12.10, 6.10,
        alt="VH12 SEA-AD-network top-five non-mitochondrial key drivers by sex/APOE and broad-cell category",
    )
    ui.add_text(
        slide,
        "Blue hatching = ACAT across ≥2 returned calls; orange = one-call q passthrough.",
        0.88, 6.72, 11.60, 0.25, size=9.8, color=ui.GRAY, italic=True, align=PP_ALIGN.CENTER,
    )
    ui.add_notes(
        slide,
        goal="Present the leading VH12 non-mitochondrial candidates without hiding category sparsity.",
        walkthrough="The display contains 20 entries across five populated categories: F_e33 excitatory, F_e4 astrocytes, and M_e33 excitatory, inhibitory, and oligodendrocytes.",
        boundary="Top five is a display cap, not a significance threshold. Colors and hatching identify the aggregation route, not an evidence tier.",
        transition="Next, summarize recurrence across the five populated categories.",
    )


def append_recurrence(prs: Presentation, facts: dict[str, Any]) -> None:
    recurrence = facts["recurrence"]
    slide = ui.new_slide(prs)
    ui.add_title_block(
        slide,
        "PAPOLA and PJVK each recur across 2 SEA-AD categories",
        "Returned-only recurrence is descriptive across sex/APOE × broad-cell categories.",
    )
    ui.add_rect(slide, 0.47, 1.27, 7.16, 5.40, color=ui.WHITE, outline=ui.LIGHT)
    ui.add_picture_contain(
        slide, SEA_RECURRENCE, 0.58, 1.37, 6.94, 5.18,
        alt="VH12 SEA-AD-network bar chart of the twenty most recurrent non-mitochondrial key drivers",
    )
    ui.add_rect(slide, 7.91, 1.42, 4.75, 4.98, color=ui.PALE_GREEN, outline=ui.TEAL)
    ui.add_panel_title(slide, "How to read it", 8.22, 1.74, 4.13, accent=ui.TEAL)
    ui.add_bullets(
        slide,
        [
            "Bar length = categories containing the gene.",
            "Fill = category occurrences combining ≥2 returned calls.",
            "PAPOLA: excitatory neurons + oligodendrocytes.",
            "PJVK: inhibitory neurons + oligodendrocytes.",
            f"The remaining {len(recurrence) - 2} displayed genes occur in one category.",
        ],
        8.23, 2.28, 4.02, size=11.3, line_h=0.61, accent=ui.TEAL,
    )
    ui.add_rect(slide, 8.22, 5.55, 4.10, 0.56, color=ui.WHITE, outline=ui.LIGHT)
    ui.add_text(slide, "Only 5 categories return non-MT genes.", 8.41, 5.66, 3.72, 0.36, size=10.1, color=ui.VERMILION_TEXT, bold=True, align=PP_ALIGN.CENTER)
    ui.add_notes(
        slide,
        goal="Interpret recurrence under the SEA-AD-specific networks.",
        walkthrough="PAPOLA and PJVK each appear in two M_e33 broad-cell categories. Every other gene in the top-20 recurrence display appears in one category.",
        boundary="Five return-bearing categories impose a low recurrence ceiling. Category recurrence is neither independent replication nor call count.",
        transition="Summarize the row, unit, gene, and category counts.",
    )


def append_part2_summary(prs: Presentation, facts: dict[str, Any]) -> None:
    slide = ui.new_slide(prs)
    ui.add_title_block(
        slide,
        f"SEA-AD output: {facts['sea_units']} non-MT category units represent {facts['sea_genes']} genes",
    )
    metrics = [
        (facts["non_mt_rows"], "non-MT returned call rows", ui.TEAL),
        (facts["sea_units"], "gene × category units", ui.TEAL),
        (facts["sea_genes"], "distinct gene symbols", ui.TEAL),
        (facts["category_count"], "categories with returns", ui.BLUE),
        (facts["top5_count"], "top-five displayed rows", ui.GOLD),
    ]
    for i, (value, label, accent) in enumerate(metrics):
        ui.add_metric(slide, str(value), label, 0.72 + i * 2.34, 1.52, 2.12, accent=accent)
    ui.add_rect(slide, 0.72, 3.14, 5.76, 3.10, color=ui.PALE_GREEN, outline=ui.TEAL)
    ui.add_panel_title(slide, "Where the 96 units occur", 1.02, 3.46, 5.16, accent=ui.TEAL)
    ui.add_text(slide, "By sex/APOE group", 1.03, 4.03, 2.05, 0.24, size=10.2, color=ui.TEAL_TEXT, bold=True)
    ui.add_text(slide, "M_e33 91  •  F_e4 3  •  F_e33 2  •  other groups 0", 1.03, 4.40, 4.97, 0.60, size=11.2, color=ui.NAVY, bold=True)
    ui.add_text(slide, "By broad network", 1.03, 5.16, 2.05, 0.24, size=10.2, color=ui.TEAL_TEXT, bold=True)
    ui.add_text(slide, "Excitatory 43 • Inhibitory 27 • Oligo 23 • Astrocytes 3\nMicroglia, OPCs, Vasculature 0", 1.03, 5.48, 4.97, 0.62, size=10.1, color=ui.NAVY, bold=True)
    ui.add_rect(slide, 6.76, 3.14, 5.48, 3.10, color=ui.PALE_SKY, outline=ui.SKY)
    ui.add_panel_title(slide, "Execution and filtering", 7.06, 3.46, 4.88, accent=ui.BLUE)
    ui.add_bullets(
        slide,
        [
            f"{facts['significant_calls']} calls with returns; {facts['empty_calls']} completed empty.",
            f"{facts['stock_rows']} significant return rows before mitochondrial filtering.",
            f"{facts['mt_excluded']} core-MitoCarta rows excluded.",
            f"{facts['non_mt_rows']} non-MT rows retained for aggregation.",
        ],
        7.08, 4.02, 4.78, size=11.0, line_h=0.48, accent=ui.BLUE,
    )
    ui.add_text(slide, "Returned-only ACAT scores are exploratory and not formally FDR-controlled across genes.", 0.88, 6.60, 11.60, 0.28, size=10.8, color=ui.VERMILION_TEXT, bold=True, align=PP_ALIGN.CENTER)
    ui.add_notes(
        slide,
        goal="Summarize the scale and distribution of the VH12 result set.",
        walkthrough="The 135 retained rows form 96 category units representing 94 genes across five categories. Most units and 32 of 34 active calls are in M_e33.",
        boundary="The result conditions on within-call significant returns and has no final across-gene correction. Vasculature remains encode-only exploratory and contributes no active call here.",
        transition="Compare the resulting driver set with the existing ROSMAP aggregation.",
    )


def append_part3_divider(prs: Presentation) -> None:
    slide = ui.new_slide(prs, bg=ui.NAVY)
    ui.add_text(slide, "PART 3", 0.78, 0.67, 2.4, 0.28, size=10.5, color=ui.PURPLE, bold=True)
    ui.add_rect(slide, 0.78, 1.26, 0.10, 2.38, color=ui.PURPLE, outline=None, radius=False)
    ui.add_text(slide, "Do the same drivers appear in both cohorts?", 1.18, 1.36, 8.65, 1.34, size=32, color=ui.WHITE, bold=True, font=ui.FONT_HEAD, valign=MSO_ANCHOR.MIDDLE)
    ui.add_text(slide, "ROSMAP and VH12 SEA-AD returned-only results are compared at gene and matched-context levels.", 1.20, 3.00, 8.30, 0.84, size=15.0, color=ui.RGBColor(204, 219, 234))
    ui.add_rect(slide, 9.78, 1.28, 2.70, 4.70, color=ui.NAVY_2, outline=None)
    ui.add_text(slide, "IN THIS PART", 10.08, 1.67, 2.10, 0.24, size=9.4, color=ui.PURPLE, bold=True)
    for index, topic in enumerate(["Comparison design", "Gene overlap", "Context overlap", "Interpretation"], start=1):
        y = 2.20 + (index - 1) * 0.82
        ui.add_circle(slide, 10.06, y + 0.03, 0.34, ui.PURPLE)
        ui.add_text(slide, str(index), 10.06, y + 0.08, 0.34, 0.17, size=8.2, color=ui.WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        ui.add_text(slide, topic, 10.53, y, 1.58, 0.52, size=11.1, color=ui.WHITE, bold=True)
    ui.add_notes(
        slide,
        goal="Frame the cross-cohort comparison using the VH12 SEA-AD network rerun.",
        walkthrough="Part 3 compares driver symbols and their sex/APOE-by-cell-type contexts, then contrasts the observed context agreement with descriptive baselines.",
        boundary="This is simultaneously a cross-cohort and cross-network comparison. It describes convergence but does not isolate a biological cohort effect.",
        transition="Define the two agreement levels before reading the overlap counts.",
    )


def append_design(prs: Presentation, facts: dict[str, Any]) -> None:
    slide = ui.new_slide(prs)
    ui.add_title_block(slide, "Two ways to ask whether a driver reappears", "Gene-level overlap ignores context; category-level overlap requires the same group and broad cell type.")
    ui.add_rect(slide, 0.64, 1.52, 5.86, 4.30, color=ui.PALE_SKY, outline=ui.BLUE)
    ui.add_text(slide, "GENE LEVEL", 0.95, 1.84, 3.2, 0.27, size=10, color=ui.BLUE, bold=True)
    ui.add_text(slide, "Does the same gene return anywhere?", 0.95, 2.18, 4.95, 0.64, size=20, color=ui.NAVY, bold=True, font=ui.FONT_HEAD)
    ui.add_metric(slide, str(facts["ros_genes"]), "ROSMAP driver genes", 0.95, 3.10, 2.30, accent=ui.BLUE)
    ui.add_metric(slide, str(facts["sea_genes"]), "SEA-AD driver genes", 3.50, 3.10, 2.30, accent=ui.BLUE)
    ui.add_text(slide, "Permissive overlap, but still affected by the network used to nominate drivers.", 0.96, 4.60, 5.20, 0.60, size=11.5, color=ui.GRAY)
    ui.add_rect(slide, 6.83, 1.52, 5.86, 4.30, color=ui.PALE_GOLD, outline=ui.GOLD)
    ui.add_text(slide, "CATEGORY LEVEL", 7.14, 1.84, 3.2, 0.27, size=10, color=ui.GOLD_TEXT, bold=True)
    ui.add_text(slide, "Same gene, same group, same cell type?", 7.14, 2.18, 4.95, 0.64, size=20, color=ui.NAVY, bold=True, font=ui.FONT_HEAD)
    ui.add_metric(slide, str(facts["ros_units"]), "ROSMAP gene × category units", 7.14, 3.10, 2.30, accent=ui.GOLD)
    ui.add_metric(slide, str(facts["sea_units"]), "SEA-AD gene × category units", 9.69, 3.10, 2.30, accent=ui.GOLD)
    ui.add_text(slide, "Strict context agreement; bounded by five SEA-AD return-bearing categories.", 7.15, 4.60, 5.20, 0.60, size=11.5, color=ui.GRAY)
    ui.add_text(slide, "Unlike the earlier deck, VH12 uses SEA-AD-specific network topology; network choice is part of the comparison.", 0.88, 6.35, 11.60, 0.28, size=11.0, color=ui.PURPLE, bold=True, align=PP_ALIGN.CENTER)
    ui.add_notes(
        slide,
        goal="Define gene-level and category-level agreement for the updated analysis.",
        walkthrough="Gene-level overlap compares symbols anywhere in each result. Category-level overlap additionally requires the same sex/APOE group and broad cell type.",
        boundary="ROSMAP and SEA-AD use cohort-specific network inputs in this comparison, so overlap depends on topology as well as biological replication.",
        transition="First examine the permissive gene-level result.",
    )


def append_gene_overlap(prs: Presentation, facts: dict[str, Any]) -> None:
    shared = facts["shared_genes"]
    pct = round(100 * len(shared) / facts["sea_genes"])
    slide = ui.new_slide(prs)
    ui.add_title_block(slide, f"{len(shared)} of {facts['sea_genes']} SEA-AD drivers also return in ROSMAP", "Gene-level overlap is limited after rerunning KDA with SEA-AD-specific networks.")
    nodes = [
        (str(facts["sea_genes"]), "SEA-AD driver genes"),
        (f"{len(shared)} ({pct}%)", "also ROSMAP drivers"),
        (str(len(facts["matched"])), "exact category match"),
    ]
    for i, (value, label) in enumerate(nodes):
        x = 0.90 + i * 3.95
        ui.add_rect(slide, x, 1.60, 3.60, 1.45, color=ui.WHITE, outline=ui.LIGHT)
        ui.add_text(slide, value, x + 0.25, 1.78, 3.10, 0.55, size=26, color=ui.BLUE, bold=True, font=ui.FONT_HEAD, align=PP_ALIGN.CENTER)
        ui.add_text(slide, label, x + 0.25, 2.42, 3.10, 0.40, size=11.5, color=ui.GRAY, bold=True, align=PP_ALIGN.CENTER)
    ui.add_rect(slide, 0.90, 3.55, 11.50, 2.30, color=ui.PALE_SKY, outline=ui.SKY)
    ui.add_panel_title(slide, "The 10 shared genes", 1.20, 3.85, 10.80, accent=ui.BLUE)
    ui.add_text(slide, "CHD3  •  CTNND2  •  CXXC1  •  DEF8  •  LAGE3\nMIPOL1  •  MKLN1  •  PAPOLA  •  SECISBP2  •  SYN1", 1.22, 4.42, 10.70, 0.82, size=15.0, color=ui.NAVY, bold=True, align=PP_ALIGN.CENTER)
    ui.add_text(slide, "LAGE3 and SYN1 also appear on both cohorts’ top-20 recurrence displays.", 1.22, 5.37, 10.70, 0.28, size=10.4, color=ui.GRAY, italic=True, align=PP_ALIGN.CENTER)
    ui.add_text(slide, "The earlier common-network SEA-AD result should not be mixed with this network-specific rerun.", 0.88, 6.35, 11.60, 0.28, size=11.0, color=ui.PURPLE, bold=True, align=PP_ALIGN.CENTER)
    ui.add_notes(
        slide,
        goal="State the updated gene-level overlap and name the shared set.",
        walkthrough="Ten of 94 SEA-AD non-MT drivers appear among the 433 ROSMAP drivers. Only MIPOL1 also keeps the full category context.",
        boundary="The reduction from the earlier deck coincides with switching the SEA-AD KDA network, so it is not evidence of poorer biological replication by itself.",
        transition="Inspect the sole exact category match.",
    )


def q_text(value: float) -> str:
    return f"{value:.4f}" if value >= 0.001 else f"{value:.1e}".replace("e-0", "e-")


def append_exact_match(prs: Presentation, facts: dict[str, Any]) -> None:
    row = facts["matched"].iloc[0]
    slide = ui.new_slide(prs)
    ui.add_title_block(slide, "One driver-category pair matches exactly", "MIPOL1 returns in male e33 excitatory neurons in both cohorts.")
    ui.add_table(
        slide,
        ["Sex/APOE · cell type", "Gene", "SEA-AD score (rank)", "ROSMAP score (rank)"],
        [[
            "M_e33 · Excitatory neurons", "MIPOL1",
            f"{q_text(float(row['score_sea']))}  (#{int(row['display_rank_sea'])})",
            f"{q_text(float(row['score_ros']))}  (#{int(row['display_rank_ros'])})",
        ]],
        0.90, 1.55, [4.05, 1.85, 2.80, 2.80], row_h=0.62, header_h=0.52, font_size=10.4,
    )
    ui.add_rect(slide, 0.90, 3.15, 5.50, 2.45, color=ui.PALE_GREEN, outline=ui.TEAL)
    ui.add_panel_title(slide, "SEA-AD", 1.20, 3.47, 4.90, accent=ui.TEAL)
    ui.add_bullets(slide, [
        f"Rank #{int(row['display_rank_sea'])} within M_e33 excitatory neurons.",
        f"Score {q_text(float(row['score_sea']))} from {int(row['calls_sea'])} returned calls.",
        "Nominated using the SEA-AD-specific excitatory network.",
    ], 1.20, 4.00, 4.88, size=11.5, line_h=0.48, accent=ui.TEAL)
    ui.add_rect(slide, 6.90, 3.15, 5.50, 2.45, color=ui.PALE_GOLD, outline=ui.GOLD)
    ui.add_panel_title(slide, "ROSMAP", 7.20, 3.47, 4.90, accent=ui.GOLD)
    ui.add_bullets(slide, [
        f"Rank #{int(row['display_rank_ros'])} within M_e33 excitatory neurons.",
        f"Score {q_text(float(row['score_ros']))} from {int(row['calls_ros'])} returned call.",
        "The only exact category overlap among 96 SEA-AD units.",
    ], 7.20, 4.00, 4.88, size=11.5, line_h=0.48, accent=ui.GOLD)
    ui.add_text(slide, "Exact agreement is descriptive convergence—not a calibrated replication P value.", 0.88, 6.25, 11.60, 0.28, size=11.0, color=ui.PURPLE, bold=True, align=PP_ALIGN.CENTER)
    ui.add_notes(
        slide,
        goal="Show the complete exact category-level overlap.",
        walkthrough="MIPOL1 is the only gene with the same sex/APOE group and broad cell type in both cohorts. It ranks seventh in SEA-AD and twenty-ninth in ROSMAP within M_e33 excitatory neurons.",
        boundary="Both scores are exploratory returned-only quantities; an exact match is not a formal replication test.",
        transition="Separate exact agreement from partial context retention.",
    )


def append_context_examples(prs: Presentation) -> None:
    slide = ui.new_slide(prs)
    ui.add_title_block(slide, "MIPOL1 is the sole exact context match", "A few other shared genes retain either cell type or stratum—but not both.")
    cards = [
        ("Exact context", "1 pair", "MIPOL1\nM_e33 · excitatory in both cohorts", ui.PALE_GREEN, ui.TEAL),
        ("Same cell type", "2 additional pairs", "CXXC1 · excitatory\nCTNND2 · inhibitory\nBoth switch sex/APOE group", ui.PALE_SKY, ui.BLUE),
        ("Same group", "3 additional pairs", "DEF8, LAGE3, SYN1\nAll remain M_e33 but switch\nbroad-cell network", ui.PALE_GOLD, ui.GOLD),
    ]
    for i, (heading, count, body, bg, accent) in enumerate(cards):
        x = 0.66 + i * 4.10
        ui.add_rect(slide, x, 1.60, 3.83, 4.55, color=bg, outline=accent)
        ui.add_text(slide, heading, x + 0.28, 1.92, 3.30, 0.46, size=21, color=ui.NAVY, bold=True, font=ui.FONT_HEAD)
        ui.add_text(slide, count, x + 0.28, 2.52, 3.30, 0.30, size=11.5, color=ui.readable_accent(accent), bold=True)
        ui.add_rect(slide, x + 0.26, 3.05, 3.31, 2.80, color=ui.WHITE, outline=ui.LIGHT)
        ui.add_text(slide, body, x + 0.44, 3.35, 2.96, 2.10, size=12.0, color=ui.GRAY, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    ui.add_notes(
        slide,
        goal="Distinguish exact context replication from partial context overlap.",
        walkthrough="MIPOL1 preserves both group and cell type. CXXC1 and CTNND2 preserve broad cell type while changing stratum. DEF8, LAGE3, and SYN1 preserve M_e33 while changing broad cell type.",
        boundary="The listed pairs are descriptive and non-independent because genes can occupy multiple ROSMAP categories.",
        transition="Quantify where the shared M_e33 genes appear across ROSMAP strata.",
    )


def append_group_matrix(prs: Presentation, facts: dict[str, Any]) -> None:
    slide = ui.new_slide(prs)
    ui.add_title_block(slide, "Shared genes are not concentrated in matching M_e33", "Where the 10 shared SEA-AD male-e33 drivers appear within ROSMAP.")
    ui.add_text(slide, f"{facts['sea_m_shared_count']} of {facts['sea_m_gene_count']} SEA-AD M_e33 drivers are also ROSMAP drivers. A gene may occur in more than one ROSMAP group.", 0.90, 1.50, 11.40, 0.34, size=12.5, color=ui.DARK)
    matrix = facts["group_matrix"]
    ordered = sorted(matrix.items(), key=lambda item: (-item[1], item[0]))
    rows = []
    for group, count in ordered:
        note = "highest overlap" if group == "F_e4" else "matching group" if group == "M_e33" else ""
        rows.append([group, str(count), note])
    ui.add_table(slide, ["ROSMAP group", "Shared genes (of 10)", "Note"], rows, 2.60, 2.05, [2.60, 2.60, 2.90], row_h=0.50, header_h=0.50, font_size=10.6)
    ui.add_rect(slide, 0.90, 5.75, 11.50, 0.62, color=ui.PALE_RED, outline=ui.VERMILION)
    ui.add_text(slide, "F_e4 contains seven of the shared genes; matching M_e33 contains four.", 1.14, 5.93, 11.05, 0.28, size=12.2, color=ui.VERMILION_TEXT, bold=True, align=PP_ALIGN.CENTER)
    ui.add_notes(
        slide,
        goal="Test descriptively whether shared SEA-AD M_e33 genes concentrate in the matching ROSMAP stratum.",
        walkthrough="Seven of the ten shared genes appear in ROSMAP F_e4 categories, while four appear in matching M_e33 categories. Counts can exceed ten in total because a ROSMAP gene may return in multiple groups.",
        boundary="Unequal category sizes and repeated genes make these descriptive counts, not a calibrated enrichment test.",
        transition="Compare cell-type and group retention across every matched-gene unit pairing.",
    )


def append_context_rates(prs: Presentation, facts: dict[str, Any]) -> None:
    net_pct = round(100 * facts["same_network_count"] / facts["pair_count"])
    group_pct = round(100 * facts["same_group_count"] / facts["pair_count"])
    chance_pct = round(100 * facts["chance_network"])
    slide = ui.new_slide(prs)
    ui.add_title_block(slide, "Neither cell type nor sex/APOE context is strongly preserved", f"All {facts['pair_count']} cross-cohort unit pairings of the 10 shared genes.")
    ui.add_rect(slide, 0.90, 1.62, 5.55, 3.05, color=ui.PALE_GREEN, outline=ui.TEAL)
    ui.add_text(slide, f"{net_pct}%", 1.20, 1.95, 4.95, 0.95, size=52, color=ui.TEAL_TEXT, bold=True, font=ui.FONT_HEAD, align=PP_ALIGN.CENTER)
    ui.add_text(slide, f"share broad cell type ({facts['same_network_count']} of {facts['pair_count']})", 1.20, 3.15, 4.95, 0.32, size=14.0, color=ui.NAVY, bold=True, align=PP_ALIGN.CENTER)
    ui.add_text(slide, f"descriptive random-pairing expectation ≈{chance_pct}%", 1.20, 3.60, 4.95, 0.30, size=11.0, color=ui.GRAY, align=PP_ALIGN.CENTER)
    ui.add_rect(slide, 6.85, 1.62, 5.55, 3.05, color=ui.PALE_RED, outline=ui.VERMILION)
    ui.add_text(slide, f"{group_pct}%", 7.15, 1.95, 4.95, 0.95, size=52, color=ui.VERMILION_TEXT, bold=True, font=ui.FONT_HEAD, align=PP_ALIGN.CENTER)
    ui.add_text(slide, f"share sex/APOE group ({facts['same_group_count']} of {facts['pair_count']})", 7.15, 3.15, 4.95, 0.32, size=14.0, color=ui.NAVY, bold=True, align=PP_ALIGN.CENTER)
    ui.add_text(slide, "four M_e33 pairings; only MIPOL1 also matches cell type", 7.15, 3.60, 4.95, 0.30, size=11.0, color=ui.GRAY, align=PP_ALIGN.CENTER)
    ui.add_rect(slide, 0.90, 5.00, 11.50, 1.30, color=ui.WHITE, outline=ui.LIGHT)
    ui.add_panel_title(slide, "Interpretation", 1.20, 5.22, 10.80, accent=ui.PURPLE)
    ui.add_text(slide, "The VH12 network-specific driver set shows little context transfer to ROSMAP. Network topology and cohort both changed, so the two causes cannot be separated here.", 1.22, 5.70, 10.80, 0.40, size=11.8, color=ui.DARK)
    ui.add_notes(
        slide,
        goal="Quantify the limited context agreement in the VH12 comparison.",
        walkthrough="Three of 22 cross-cohort unit pairings share broad cell type and four share sex/APOE group. The descriptive random-pairing expectation for cell type is about 32 percent.",
        boundary="Pairings are non-independent and the chance value is descriptive, not a formal null. Network and cohort effects are confounded.",
        transition="Summarize the sex distribution of the shared M_e33 genes in ROSMAP.",
    )


def append_sex_split(prs: Presentation, facts: dict[str, Any]) -> None:
    split = facts["sex_split"]
    slide = ui.new_slide(prs)
    ui.add_title_block(slide, "Shared SEA-AD male-e33 genes span both ROSMAP sexes", "ROSMAP sex context is mixed across the 10 shared M_e33 drivers.")
    ui.add_rect(slide, 0.66, 1.55, 5.95, 4.55, color=ui.PALE_SKY, outline=ui.BLUE)
    ui.add_panel_title(slide, "Gene membership", 0.96, 1.87, 5.35, accent=ui.BLUE)
    ui.add_text(slide, "Male-only (2)", 0.98, 2.52, 1.55, 0.28, size=11.2, color=ui.BLUE, bold=True)
    ui.add_text(slide, ", ".join(split["male_only"]), 2.58, 2.52, 3.55, 0.42, size=12.0, color=ui.NAVY, bold=True)
    ui.add_text(slide, "Female-only (4)", 0.98, 3.38, 1.55, 0.28, size=11.2, color=ui.VERMILION_TEXT, bold=True)
    ui.add_text(slide, ", ".join(split["female_only"]), 2.58, 3.38, 3.55, 0.62, size=11.5, color=ui.NAVY, bold=True)
    ui.add_text(slide, "Both sexes (4)", 0.98, 4.46, 1.55, 0.28, size=11.2, color=ui.GRAY, bold=True)
    ui.add_text(slide, ", ".join(split["both"]), 2.58, 4.46, 3.55, 0.62, size=11.5, color=ui.NAVY, bold=True)
    ui.add_rect(slide, 6.90, 1.55, 5.78, 4.55, color=ui.PALE_GRAY, outline=ui.LIGHT)
    ui.add_panel_title(slide, "ROSMAP sex distribution", 7.20, 1.87, 5.15, accent=ui.GRAY)
    bars = [
        ("Male-only", len(split["male_only"]), ui.BLUE),
        ("Female-only", len(split["female_only"]), ui.VERMILION),
        ("Both sexes", len(split["both"]), ui.GRAY),
    ]
    maximum = max(count for _, count, _ in bars)
    for i, (label, count, color) in enumerate(bars):
        y = 2.60 + i * 0.90
        ui.add_text(slide, label, 7.22, y, 2.60, 0.26, size=11.0, color=ui.NAVY, bold=True)
        ui.add_rect(slide, 7.22, y + 0.32, 3.90 * count / maximum, 0.24, color=color, outline=None, radius=False)
        ui.add_text(slide, str(count), 11.35, y + 0.28, 0.95, 0.28, size=13.0, color=color, bold=True)
    ui.add_text(slide, "Only CHD3 and DEF8 are restricted to male ROSMAP categories; four shared genes appear only in female ROSMAP categories.", 7.22, 5.35, 5.20, 0.62, size=10.4, color=ui.GRAY)
    ui.add_notes(
        slide,
        goal="Show whether shared SEA-AD M_e33 genes retain a male context in ROSMAP.",
        walkthrough="Two shared genes are male-only in ROSMAP, four are female-only, and four appear in both sexes. The full membership is shown rather than only counts.",
        boundary="Observed sex restriction reflects where returned-only units occur under unequal run availability; it is not a sex-interaction test.",
        transition="Close with the supported findings and the interpretation limits.",
    )


def append_takeaway(prs: Presentation, facts: dict[str, Any]) -> None:
    slide = ui.new_slide(prs)
    ui.add_title_block(slide, "What the SEA-AD-network rerun does—and does not—show")
    ui.add_rect(slide, 0.66, 1.30, 5.95, 5.05, color=ui.PALE_GREEN, outline=ui.TEAL)
    ui.add_panel_title(slide, "Supported", 0.97, 1.62, 5.35, accent=ui.TEAL)
    ui.add_bullets(slide, [
        f"{len(facts['shared_genes'])} of {facts['sea_genes']} SEA-AD drivers also occur in ROSMAP.",
        "MIPOL1 is the sole exact M_e33-excitatory context match.",
        "CXXC1 and CTNND2 retain broad cell type while switching strata.",
        "PAPOLA and PJVK each recur across two SEA-AD categories.",
    ], 0.99, 2.20, 5.40, size=11.5, line_h=0.78, accent=ui.TEAL)
    ui.add_rect(slide, 6.90, 1.30, 5.78, 5.05, color=ui.PALE_RED, outline=ui.VERMILION)
    ui.add_panel_title(slide, "Not supported / limits", 7.21, 1.62, 5.18, accent=ui.VERMILION)
    ui.add_bullets(slide, [
        "No broad preservation of cell-type or sex/APOE context is evident.",
        "Only M_e33 is well represented among SEA-AD active calls.",
        "Cohort and network topology changed together; their effects are inseparable here.",
        "Returned-only ACAT scores are exploratory, post-selected, and not final-FDR controlled.",
    ], 7.23, 2.20, 5.25, size=11.5, line_h=0.78, accent=ui.VERMILION)
    ui.add_text(slide, "Takeaway: SEA-AD-specific topology yields a distinct driver set; MIPOL1 is the clearest matched-context candidate.", 0.88, 6.65, 11.60, 0.30, size=11.7, color=ui.PURPLE, bold=True, align=PP_ALIGN.CENTER)
    ui.add_notes(
        slide,
        goal="State the updated conclusion without overclaiming validation.",
        walkthrough="The rerun finds ten shared genes, one exact context match, and weak overall context agreement. MIPOL1 is the clearest matched-context candidate, while PAPOLA and PJVK lead within-SEA-AD recurrence.",
        boundary="The comparison is descriptive, post-selected, unbalanced across strata, and confounds cohort with network topology. It does not establish causal regulation or formal replication.",
        transition="End of deck.",
    )


def build(source: Path, output: Path, facts: dict[str, Any]) -> None:
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {output}")
    prs = Presentation(str(source))
    if len(prs.slides) != 24:
        raise RuntimeError(f"Expected 24-slide source deck, found {len(prs.slides)}")
    ui.set_notes_body_template(prs.slides[0].notes_slide.notes_placeholder._element)
    update_overview_slides(prs, facts)
    drop_slides_after(prs, 10)
    append_part2_divider(prs, facts)
    append_run_universe(prs, facts)
    append_top5(prs, facts)
    append_recurrence(prs, facts)
    append_part2_summary(prs, facts)
    append_part3_divider(prs)
    append_design(prs, facts)
    append_gene_overlap(prs, facts)
    append_exact_match(prs, facts)
    append_context_examples(prs)
    append_group_matrix(prs, facts)
    append_context_rates(prs, facts)
    append_sex_split(prs, facts)
    append_takeaway(prs, facts)
    if len(prs.slides) != 24:
        raise RuntimeError(f"Expected 24 output slides, built {len(prs.slides)}")
    prs.core_properties.title = "ROSMAP and SEA-AD sex/APOE key-driver analysis — VH12 network rerun"
    prs.core_properties.subject = "VH12 SEA-AD-specific Bayesian-network KDA rerun"
    prs.core_properties.author = "Alzheimer project analysis team"
    prs.core_properties.comments = "Built from validated VH12 outputs on 2026-09-10 for the 2026-09-16 presentation."
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temp_name = tempfile.mkstemp(prefix=f".{output.name}.", suffix=".tmp", dir=output.parent)
    os.close(descriptor)
    temporary = Path(temp_name)
    try:
        prs.save(str(temporary))
        temporary.replace(output)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def validate(source: Path, output: Path, original_3_10: dict[int, tuple[str, int]]) -> tuple[list[dict[str, str]], list[dict[str, Any]]]:
    checks: list[dict[str, str]] = []

    def check(check_id: str, observed: Any, expected: Any, passed: bool) -> None:
        checks.append({
            "schema_version": "vh12_seaad_network_deck_checks_v1",
            "check_id": check_id,
            "observed": str(observed),
            "expected": str(expected),
            "passed": str(bool(passed)).upper(),
        })

    prs = Presentation(str(output))
    check("pptx_exists", output.is_file(), True, output.is_file())
    check("source_deck_unchanged", sha256(source), sha256(SOURCE_DECK), sha256(source) == sha256(SOURCE_DECK))
    check("slide_count", len(prs.slides), 24, len(prs.slides) == 24)
    check("widescreen_ratio", round(prs.slide_width / prs.slide_height, 3), 1.778, round(prs.slide_width / prs.slide_height, 3) == 1.778)
    title_hits = sum(title in slide_text(slide) for title, slide in zip(EXPECTED_TITLES, prs.slides, strict=True))
    check("title_sequence", title_hits, 24, title_hits == 24)
    unchanged = all(
        original_3_10[number] == (slide_text(prs.slides[number - 1]), len(prs.slides[number - 1].shapes))
        for number in range(3, 11)
    )
    check("slides_3_to_10_unchanged", unchanged, True, unchanged)
    notes_ok = 0
    bounded_new = 0
    picture_count = 0
    alt_count = 0
    inventory: list[dict[str, Any]] = []
    for number, slide in enumerate(prs.slides, start=1):
        notes_frame = slide.notes_slide.notes_text_frame
        notes = notes_frame.text if notes_frame is not None else ""
        if all(marker in notes for marker in ("Teaching goal:", "Walk through:", "Scientific boundary:", "Transition:")):
            notes_ok += 1
        in_bounds = True
        pictures = 0
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width + 10 or shape.top + shape.height > prs.slide_height + 10:
                in_bounds = False
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
                picture_count += 1
                c_nv_pr = shape._element.xpath(".//p:cNvPr")[0]
                if c_nv_pr.get("descr", "").strip():
                    alt_count += 1
        if number >= 11 and in_bounds:
            bounded_new += 1
        inventory.append({
            "schema_version": "vh12_seaad_network_deck_slide_inventory_v1",
            "slide_number": number,
            "title": EXPECTED_TITLES[number - 1],
            "shape_count": len(slide.shapes),
            "picture_count": pictures,
            "speaker_note_words": len(notes.split()),
        })
    check("structured_notes", notes_ok, 24, notes_ok == 24)
    # Slide 8 in the supplied deck intentionally extends its full-bleed image
    # 0.05 inches below the canvas. Preserve slides 3-10 byte-semantically and
    # apply the strict boundary gate to every newly built slide instead.
    check("all_new_shapes_within_slide", bounded_new, 14, bounded_new == 14)
    check("picture_count", picture_count, 4, picture_count == 4)
    check("picture_alt_text_count", alt_count, 4, alt_count == 4)
    all_text = "\n".join(slide_text(slide) for slide in prs.slides)
    stale = [phrase for phrase in ("42 active KDA calls", "35 of 91", "Eight driver-category", "HGSNAT recurs") if phrase in all_text]
    check("stale_phase11_claims_absent", stale, [], not stale)
    with zipfile.ZipFile(output) as archive:
        bad_member = archive.testzip()
        media_hashes = {
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        }
    check("pptx_zip_integrity", bad_member, None, bad_member is None)
    embedded = sum(sha256(path) in media_hashes for path in (SEA_TOP5, SEA_RECURRENCE))
    check("vh12_pngs_embedded_byte_exact", embedded, 2, embedded == 2)
    check("core_title", prs.core_properties.title, "ROSMAP and SEA-AD sex/APOE key-driver analysis — VH12 network rerun", prs.core_properties.title == "ROSMAP and SEA-AD sex/APOE key-driver analysis — VH12 network rerun")
    return checks, inventory


def write_audits(source: Path, output: Path, audit_dir: Path, checks: list[dict[str, str]], inventory: list[dict[str, Any]]) -> None:
    inputs = [
        ("source_deck", source),
        ("vh12_status", SEA_STATUS),
        ("vh12_manifest", SEA_MANIFEST),
        ("vh12_categories", SEA_CATEGORIES),
        ("vh12_top5_png", SEA_TOP5),
        ("vh12_recurrence_png", SEA_RECURRENCE),
        ("rosmap_categories", ROS_CATEGORIES),
        ("builder", Path(__file__).resolve()),
        ("output_deck", output),
    ]
    manifest = [{
        "schema_version": "vh12_seaad_network_deck_manifest_v1",
        "role": role,
        "path": str(path),
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
    } for role, path in inputs]
    write_tsv(manifest, audit_dir / "input_manifest.tsv", ["schema_version", "role", "path", "bytes", "sha256"])
    write_tsv(checks, audit_dir / "checks.tsv", ["schema_version", "check_id", "observed", "expected", "passed"])
    write_tsv(inventory, audit_dir / "slide_inventory.tsv", ["schema_version", "slide_number", "title", "shape_count", "picture_count", "speaker_note_words"])
    failed = sum(row["passed"] != "TRUE" for row in checks)
    write_tsv([{
        "schema_version": "vh12_seaad_network_deck_status_v1",
        "output_path": str(output),
        "slide_count": len(inventory),
        "validation_status": "validated_complete" if failed == 0 else "failed",
        "failed_checks": failed,
        "sha256": sha256(output),
    }], audit_dir / "status.tsv", ["schema_version", "output_path", "slide_count", "validation_status", "failed_checks", "sha256"])


def main() -> int:
    args = parse_args()
    source = args.source.resolve()
    output = args.output.resolve()
    audit_dir = args.audit_dir.resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    source_hash = sha256(source)
    source_prs = Presentation(str(source))
    original_3_10 = {
        number: (slide_text(source_prs.slides[number - 1]), len(source_prs.slides[number - 1].shapes))
        for number in range(3, 11)
    }
    facts = load_facts()
    build(source, output, facts)
    if sha256(source) != source_hash:
        raise RuntimeError("Source deck changed during build")
    checks, inventory = validate(source, output, original_3_10)
    write_audits(source, output, audit_dir, checks, inventory)
    failed = [row["check_id"] for row in checks if row["passed"] != "TRUE"]
    if failed:
        raise RuntimeError("Deck validation failed: " + ", ".join(failed))
    print(f"output={output}")
    print(f"slides={len(inventory)}")
    print(f"sha256={sha256(output)}")
    print(f"checks={len(checks)} passed")
    print(f"audit={audit_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
