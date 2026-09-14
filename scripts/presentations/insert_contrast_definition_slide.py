#!/usr/bin/env python3
"""Insert the contrast slide and place the four-step slide immediately after it."""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
import update_phase11_seaad_simple_aggr_part2 as ui  # noqa: E402


ROOT = Path(__file__).resolve().parents[2]
DECK = (
    ROOT
    / "docs"
    / "presentations"
    / "09162026"
    / "09162026_sex_apoe_kda_fine_broad.pptx"
)
AUDIT_ROOT = (
    ROOT
    / "results"
    / "presentations"
    / "09162026_contrast_definition_slide"
)

EXPECTED_INPUT_SLIDES = 27
INSERT_AFTER = 2
NEW_TITLE = "What is a contrast"
FOUR_STEPS_TITLE = "Four steps: one KDA query per contrast"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DECK)
    parser.add_argument("--output", type=Path, default=DECK)
    parser.add_argument("--audit-root", type=Path, default=AUDIT_ROOT)
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def slide_text(slide) -> str:
    return "\n".join(
        shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
    )


def semantic_slide_fingerprint(slide) -> tuple[Any, ...]:
    shapes = []
    for shape in slide.shapes:
        text = shape.text if getattr(shape, "has_text_frame", False) else ""
        shapes.append(
            (
                int(shape.shape_type),
                shape.name,
                int(shape.left),
                int(shape.top),
                int(shape.width),
                int(shape.height),
                text,
            )
        )
    notes = slide.notes_slide.notes_text_frame
    return tuple(shapes), notes.text if notes is not None else ""


def add_formula_card(
    slide,
    x: float,
    width: float,
    value: str,
    label: str,
    background,
    accent,
) -> None:
    ui.add_rect(
        slide,
        x,
        1.38,
        width,
        0.92,
        color=background,
        outline=accent,
    )
    ui.add_text(
        slide,
        value,
        x + 0.12,
        1.55,
        0.76,
        0.42,
        size=23,
        color=ui.readable_accent(accent),
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        label,
        x + 0.94,
        1.60,
        width - 1.06,
        0.30,
        size=10.6,
        color=ui.NAVY,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_operator(slide, text: str, x: float) -> None:
    ui.add_text(
        slide,
        text,
        x,
        1.63,
        0.40,
        0.34,
        size=19,
        color=ui.GRAY,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        font=ui.FONT_HEAD,
    )


def add_numbered_row(
    slide,
    number: str,
    title: str,
    detail: str,
    x: float,
    y: float,
    accent,
) -> None:
    ui.add_rect(
        slide,
        x,
        y,
        0.42,
        0.42,
        color=accent,
        outline=None,
    )
    ui.add_text(
        slide,
        number,
        x + 0.04,
        y + 0.08,
        0.34,
        0.20,
        size=10.2,
        color=ui.WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    ui.add_text(
        slide,
        title,
        x + 0.58,
        y - 0.01,
        4.72,
        0.27,
        size=11.2,
        color=ui.NAVY,
        bold=True,
    )
    ui.add_text(
        slide,
        detail,
        x + 0.58,
        y + 0.28,
        4.72,
        0.42,
        size=9.7,
        color=ui.GRAY,
    )


def build_contrast_slide(prs: Presentation):
    slide = ui.new_slide(prs)
    ui.add_title_block(
        slide,
        NEW_TITLE,
        "Fine cell type × sex × APOE defines the contrast; broad-cell aggregation happens after KDA.",
    )

    add_formula_card(
        slide,
        0.66,
        2.36,
        "N",
        "fine cell types",
        ui.PALE_SKY,
        ui.BLUE,
    )
    add_operator(slide, "×", 3.12)
    add_formula_card(
        slide,
        3.62,
        2.16,
        "2",
        "sexes",
        ui.PALE_GREEN,
        ui.TEAL,
    )
    add_operator(slide, "×", 5.88)
    add_formula_card(
        slide,
        6.38,
        2.36,
        "3",
        "APOE groups",
        ui.PALE_GOLD,
        ui.GOLD,
    )
    add_operator(slide, "=", 8.84)
    add_formula_card(
        slide,
        9.34,
        3.33,
        "6N",
        "planned contrasts",
        ui.PALE_RED,
        ui.VERMILION,
    )
    ui.add_text(
        slide,
        "ROSMAP example: 54 fine cell types × 2 sexes × 3 APOE groups = 324 planned contrasts",
        0.68,
        2.42,
        11.98,
        0.28,
        size=10.4,
        color=ui.PURPLE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    ui.add_rect(
        slide,
        0.66,
        2.88,
        5.93,
        3.52,
        color=ui.PALE_SKY,
        outline=None,
    )
    ui.add_panel_title(
        slide,
        "One contrast → at most one KDA call",
        0.96,
        3.16,
        5.30,
        accent=ui.BLUE,
    )
    add_numbered_row(
        slide,
        "1",
        "Define one AD-versus-NCI comparison",
        "Example: Ast GRM3 × Female ε2 carrier.",
        0.98,
        3.68,
        ui.BLUE,
    )
    add_numbered_row(
        slide,
        "2",
        "Remove invalid contrasts",
        "Remove source-invalid contrasts; among valid contrasts, only mapped mitochondrial queries with ≥3 genes proceed.",
        0.98,
        4.48,
        ui.BLUE,
    )
    add_numbered_row(
        slide,
        "3",
        "Run call_key_drivers()",
        "The call returns a list of significant key-driver genes for that contrast.",
        0.98,
        5.28,
        ui.BLUE,
    )
    ui.add_rect(
        slide,
        0.98,
        5.98,
        5.28,
        0.28,
        color=ui.WHITE,
        outline=ui.BLUE,
    )
    ui.add_text(
        slide,
        "Each gene can appear at most once in one contrast's returned list.",
        1.10,
        6.04,
        5.04,
        0.16,
        size=9.4,
        color=ui.NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    ui.add_rect(
        slide,
        6.75,
        2.88,
        5.92,
        3.52,
        color=ui.PALE_GOLD,
        outline=None,
    )
    ui.add_panel_title(
        slide,
        "Multiple fine-cell contrasts → one category",
        7.05,
        3.16,
        5.30,
        accent=ui.GOLD,
    )
    example_rows = [
        ("Fine type A", "Gene X returned · adjusted P₁"),
        ("Fine type B", "Gene X returned · adjusted P₂"),
        ("Fine type C", "Gene X not returned"),
    ]
    for index, (fine_type, result) in enumerate(example_rows):
        y = 3.72 + index * 0.52
        ui.add_rect(
            slide,
            7.06,
            y,
            1.22,
            0.38,
            color=ui.WHITE,
            outline=ui.LIGHT,
        )
        ui.add_text(
            slide,
            fine_type,
            7.16,
            y + 0.09,
            1.02,
            0.18,
            size=9.1,
            color=ui.GRAY,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        ui.add_text(
            slide,
            result,
            8.50,
            y + 0.08,
            3.64,
            0.20,
            size=9.7,
            color=ui.DARK,
            valign=MSO_ANCHOR.MIDDLE,
        )
    ui.add_text(
        slide,
        "↓",
        9.36,
        5.22,
        0.70,
        0.34,
        size=19,
        color=ui.GOLD_TEXT,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    ui.add_rect(
        slide,
        7.08,
        5.56,
        5.26,
        0.58,
        color=ui.WHITE,
        outline=ui.GOLD,
    )
    ui.add_text(
        slide,
        "Female ε2 × Astrocytes:  Gene X score = ACAT(P₁, P₂)",
        7.28,
        5.72,
        4.86,
        0.24,
        size=10.4,
        color=ui.GOLD_TEXT,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    ui.add_rect(
        slide,
        0.66,
        6.62,
        12.01,
        0.48,
        color=ui.PALE_RED,
        outline=ui.VERMILION,
    )
    ui.add_text(
        slide,
        "Why contrasts outnumber categories: many fine cell types map to one broad cell type; recurring returned values are combined by equal-weight, returned-only ACAT.",
        0.90,
        6.75,
        11.52,
        0.22,
        size=10.1,
        color=ui.VERMILION_TEXT,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    ui.add_notes(
        slide,
        goal="Distinguish the contrast, KDA call, and broad-cell category units.",
        walkthrough="A contrast is one AD-versus-NCI comparison within one fine cell type and one sex/APOE group. Therefore N fine cell types create six N planned contrasts; in ROSMAP, 54 fine types create 324. Remove source-invalid contrasts; among valid contrasts, call_key_drivers runs only when the mapped mitochondrial query contains at least three genes. One call returns each significant key-driver gene at most once. The same gene can return from several fine-cell contrasts belonging to the same broad-cell category, so its returned within-call adjusted P values are combined by equal-weight ACAT into one gene-by-category score.",
        boundary="The ACAT is returned-only: a contrast in which the gene was not returned does not contribute P equals one. A skipped contrast is unavailable, not a completed null KDA test.",
        transition="With contrast and category units separated, compare the two cohort funnels.",
    )
    return slide


def write_audit(path: Path, checks: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["check_id", "observed", "expected", "passed"],
            delimiter="\t",
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(checks)


def main() -> int:
    args = parse_args()
    input_path = args.input.resolve()
    output_path = args.output.resolve()
    audit_root = args.audit_root.resolve()
    if not input_path.is_file():
        raise FileNotFoundError(input_path)

    prs = Presentation(str(input_path))
    if len(prs.slides) != EXPECTED_INPUT_SLIDES:
        raise RuntimeError(
            f"Expected {EXPECTED_INPUT_SLIDES} slides, found {len(prs.slides)}"
        )
    all_text = "\n".join(slide_text(slide) for slide in prs.slides)
    if NEW_TITLE in all_text:
        raise RuntimeError("Contrast slide already present; refusing to insert twice")
    if "What is a category" not in slide_text(prs.slides[1]):
        raise RuntimeError("Slide 2 is not the expected category-definition slide")
    if FOUR_STEPS_TITLE not in slide_text(prs.slides[4]):
        raise RuntimeError("Slide 5 is not the expected four-step slide")

    before = [semantic_slide_fingerprint(slide) for slide in prs.slides]
    input_hash = sha256_file(input_path)
    backup_dir = audit_root / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / f"input_deck_{input_hash[:12]}.pptx"
    if not backup_path.exists():
        shutil.copy2(input_path, backup_path)

    ui.set_notes_body_template(
        prs.slides[0].notes_slide.notes_placeholder._element
    )
    build_contrast_slide(prs)
    slide_id_list = prs.slides._sldIdLst
    slide_ids = list(slide_id_list)
    slide_id_list.remove(slide_ids[-1])
    slide_id_list.insert(INSERT_AFTER, slide_ids[-1])
    four_steps_index = next(
        index
        for index, slide in enumerate(prs.slides)
        if FOUR_STEPS_TITLE in slide_text(slide)
    )
    four_steps_id = slide_id_list[four_steps_index]
    slide_id_list.remove(four_steps_id)
    slide_id_list.insert(INSERT_AFTER + 1, four_steps_id)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.name}.",
        suffix=".tmp",
        dir=output_path.parent,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        prs.save(str(temporary))
        reloaded = Presentation(str(temporary))
        after = [
            semantic_slide_fingerprint(slide) for slide in reloaded.slides
        ]
        other_slides_reordered_without_changes = (
            after[:INSERT_AFTER] == before[:INSERT_AFTER]
            and after[INSERT_AFTER + 1] == before[4]
            and after[INSERT_AFTER + 2 : INSERT_AFTER + 4] == before[2:4]
            and after[INSERT_AFTER + 4 :] == before[5:]
        )
        inserted_text = slide_text(reloaded.slides[INSERT_AFTER])
        four_steps_text = slide_text(reloaded.slides[INSERT_AFTER + 1])
        notes = reloaded.slides[INSERT_AFTER].notes_slide.notes_text_frame
        checks: list[dict[str, Any]] = [
            {
                "check_id": "output_slide_count",
                "observed": len(reloaded.slides),
                "expected": EXPECTED_INPUT_SLIDES + 1,
                "passed": len(reloaded.slides) == EXPECTED_INPUT_SLIDES + 1,
            },
            {
                "check_id": "contrast_slide_is_slide_3",
                "observed": NEW_TITLE in inserted_text,
                "expected": True,
                "passed": NEW_TITLE in inserted_text,
            },
            {
                "check_id": "contrast_formula_present",
                "observed": "54 fine cell types × 2 sexes × 3 APOE groups = 324"
                in inserted_text,
                "expected": True,
                "passed": "54 fine cell types × 2 sexes × 3 APOE groups = 324"
                in inserted_text,
            },
            {
                "check_id": "acat_aggregation_present",
                "observed": "ACAT(P₁, P₂)" in inserted_text,
                "expected": True,
                "passed": "ACAT(P₁, P₂)" in inserted_text,
            },
            {
                "check_id": "four_steps_slide_is_slide_4",
                "observed": FOUR_STEPS_TITLE in four_steps_text,
                "expected": True,
                "passed": FOUR_STEPS_TITLE in four_steps_text,
            },
            {
                "check_id": "all_other_slides_reordered_without_changes",
                "observed": other_slides_reordered_without_changes,
                "expected": True,
                "passed": other_slides_reordered_without_changes,
            },
            {
                "check_id": "new_slide_has_notes",
                "observed": bool(notes and notes.text.strip()),
                "expected": True,
                "passed": bool(notes and notes.text.strip()),
            },
        ]
        failed = [row["check_id"] for row in checks if not row["passed"]]
        if failed:
            raise RuntimeError("Slide insertion checks failed: " + ", ".join(failed))
        os.replace(temporary, output_path)
    finally:
        temporary.unlink(missing_ok=True)

    write_audit(audit_root / "contrast_slide_insert_checks.tsv", checks)
    print(f"Updated deck: {output_path}")
    print(f"Inserted slide: {INSERT_AFTER + 1}")
    print(f"Backup: {backup_path}")
    print(f"Audit: {audit_root}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
