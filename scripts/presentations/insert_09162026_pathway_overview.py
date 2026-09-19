#!/usr/bin/env python3
"""Insert two focused pathway-analysis slides after the reviewed slide 5.

Slide 6 lists the four frozen mitochondrial pathways used throughout the
presentation. Slide 7 shows, for each sex/APOE group, how many unique genes in
each pathway appeared as a significant DEG in at least one ROSMAP fine-cell
comparison. Existing slides, their order, and their notes are preserved.
"""

from __future__ import annotations

import argparse
import hashlib
import os
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Any

os.environ.setdefault(
    "MPLCONFIGDIR", str(Path(tempfile.gettempdir()) / "pathway_overlap_mpl")
)
os.environ.setdefault("MPLBACKEND", "Agg")

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Inches  # noqa: E402


ROOT = Path(__file__).resolve().parents[2]
PRESENTATION_SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(PRESENTATION_SCRIPTS))
import update_phase11_seaad_simple_aggr_part2 as ui  # noqa: E402

SKILL_SCRIPTS = ROOT / ".agents/skills/scientific-visualization/scripts"
sys.path.insert(0, str(SKILL_SCRIPTS))
from figure_export import save_publication_figure  # noqa: E402
from style_presets import apply_publication_style  # noqa: E402


DEFAULT_DECK = (
    ROOT
    / "docs"
    / "presentations"
    / "09162026"
    / "09162026_sex_apoe_kda_fine_broad.pptx"
)
AUDIT_ROOT = ROOT / "results" / "presentations" / "09162026_pathway_overview"
MODULE_PATH = ROOT / "config" / "phase13_respiratory_modules.tsv"
PROGRAM_MANIFEST = (
    ROOT
    / "results"
    / "minerva_production"
    / "11_pathway_deg_fine"
    / "fine_deg_pathway_program_manifest.tsv"
)
QUERY_GENES = (
    ROOT
    / "results"
    / "minerva_production"
    / "11_pathway_deg_fine"
    / "fine_deg_pathway_query_genes.tsv.gz"
)

INPUT_SLIDES = 50
OUTPUT_SLIDES = 52
UPDATED_OUTPUT_SLIDES = {52, 157, 158, 159}
INSERT_INDEX = 5
PREVIOUS_TITLE = "Pathway analysis"
NEXT_TITLE = "OXPHOS gene directions vary across sex/APOE groups"
LIST_TITLE = "Four mitochondrial pathways highlighted"
OVERLAP_TITLE = "Which pathway genes appeared among the DEGs?"
GROUP_ORDER = ["F_e2", "F_e33", "F_e4", "M_e2", "M_e33", "M_e4"]
GROUP_LABELS = [
    "Female\nε2",
    "Female\nε3/ε3",
    "Female\nε4",
    "Male\nε2",
    "Male\nε3/ε3",
    "Male\nε4",
]
PROGRAMS = (
    (
        "mtdna_oxphos_13",
        "mtDNA-encoded OXPHOS",
        13,
        "Core respiratory-complex subunits encoded by mitochondrial DNA.",
        ui.PALE_SKY,
        ui.BLUE,
    ),
    (
        "nuclear_oxphos_structural_86",
        "Nuclear structural OXPHOS",
        86,
        "Structural respiratory-complex subunits encoded by nuclear DNA.",
        ui.PALE_RED,
        ui.VERMILION,
    ),
    (
        "mitochondrial_translation_155",
        "Mitochondrial translation",
        155,
        "Ribosomal proteins and factors that make proteins inside mitochondria.",
        ui.PALE_GREEN,
        ui.TEAL_TEXT,
    ),
    (
        "mib_micos_inner_membrane_19",
        "MIB/MICOS inner membrane",
        19,
        "Proteins that shape cristae—the folds of the inner mitochondrial membrane.",
        ui.PALE_GOLD,
        ui.PURPLE,
    ),
)
EXPECTED_OVERLAPS = {
    "mtdna_oxphos_13": [12, 12, 12, 12, 12, 12],
    "nuclear_oxphos_structural_86": [56, 43, 59, 67, 48, 52],
    "mitochondrial_translation_155": [56, 41, 79, 130, 63, 58],
    "mib_micos_inner_membrane_19": [5, 2, 9, 17, 5, 7],
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_DECK)
    parser.add_argument("--output", type=Path, default=DEFAULT_DECK)
    parser.add_argument("--audit-root", type=Path, default=AUDIT_ROOT)
    parser.add_argument(
        "--figure-only",
        action="store_true",
        help="Render and validate overlap artifacts without modifying the deck.",
    )
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def slide_title(slide) -> str:
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text.strip():
            return " ".join(shape.text.split())
    return ""


def slide_fingerprint(slide) -> tuple[Any, ...]:
    notes = slide.notes_slide.notes_text_frame
    return (
        tuple(
            (
                int(shape.shape_type),
                shape.name,
                int(shape.left),
                int(shape.top),
                int(shape.width),
                int(shape.height),
                shape.text if getattr(shape, "has_text_frame", False) else "",
            )
            for shape in slide.shapes
        ),
        notes.text if notes is not None else "",
    )


def prepare_overlap_data() -> pd.DataFrame:
    for path in (MODULE_PATH, PROGRAM_MANIFEST, QUERY_GENES):
        if not path.is_file():
            raise FileNotFoundError(path)

    modules = pd.read_csv(MODULE_PATH, sep="\t", low_memory=False)
    manifest = pd.read_csv(PROGRAM_MANIFEST, sep="\t", low_memory=False)
    queries = pd.read_csv(
        QUERY_GENES,
        sep="\t",
        usecols=["query_mode", "signature_group", "symbol_hgnc_current"],
        low_memory=False,
    )
    queries = queries.loc[queries["query_mode"].eq("AD_any_mito")].copy()
    program_manifest = manifest.loc[manifest["pathway_collection"].eq("legacy_four")]

    rows: list[dict[str, Any]] = []
    for program_order, (module_id, pathway_name, expected_size, *_rest) in enumerate(
        PROGRAMS, start=1
    ):
        module_genes = set(
            modules.loc[
                modules["module_id"].eq(module_id), "current_approved_symbol"
            ].dropna()
        )
        if len(module_genes) != expected_size:
            raise RuntimeError(
                f"{module_id} contains {len(module_genes)} genes, expected {expected_size}"
            )
        manifest_rows = program_manifest.loc[
            program_manifest["pathway_id"].eq(module_id)
        ]
        if len(manifest_rows) != 1:
            raise RuntimeError(f"Missing unique Phase 11 manifest row for {module_id}")
        manifest_row = manifest_rows.iloc[0]
        if (
            manifest_row["pathway_name"] != pathway_name
            or int(manifest_row["source_pathway_size"]) != expected_size
        ):
            raise RuntimeError(f"Phase 11 manifest changed for {module_id}")

        for group_order, group in enumerate(GROUP_ORDER, start=1):
            deg_genes = set(
                queries.loc[
                    queries["signature_group"].eq(group), "symbol_hgnc_current"
                ].dropna()
            )
            overlap_genes = sorted(module_genes & deg_genes)
            rows.append(
                {
                    "program_order": program_order,
                    "group_order": group_order,
                    "module_id": module_id,
                    "pathway_name": pathway_name,
                    "signature_group": group,
                    "pathway_size": expected_size,
                    "overlap_gene_count": len(overlap_genes),
                    "overlap_percent": 100.0 * len(overlap_genes) / expected_size,
                    "overlap_genes": ",".join(overlap_genes),
                }
            )

    frame = pd.DataFrame(rows).sort_values(["program_order", "group_order"])
    observed = {
        module_id: frame.loc[
            frame["module_id"].eq(module_id), "overlap_gene_count"
        ].astype(int).tolist()
        for module_id, *_rest in PROGRAMS
    }
    if observed != EXPECTED_OVERLAPS:
        raise RuntimeError(f"Pathway-DEG overlap counts changed: {observed!r}")
    return frame


def render_overlap_figure(frame: pd.DataFrame, output_dir: Path) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    apply_publication_style("presentation")
    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"],
            "figure.facecolor": "white",
            "axes.facecolor": "white",
        }
    )
    matrix = frame.pivot(
        index="program_order", columns="group_order", values="overlap_percent"
    ).reindex(index=range(1, 5), columns=range(1, 7))
    counts = frame.pivot(
        index="program_order", columns="group_order", values="overlap_gene_count"
    ).reindex(index=range(1, 5), columns=range(1, 7))
    sizes = [program[2] for program in PROGRAMS]
    row_labels = [
        "mtDNA-encoded OXPHOS\n(13 genes)",
        "Nuclear structural OXPHOS\n(86 genes)",
        "Mitochondrial translation\n(155 genes)",
        "Inner-membrane organization\n(19 genes)",
    ]

    fig, ax = plt.subplots(figsize=(12.0, 4.20), constrained_layout=False)
    fig.subplots_adjust(left=0.25, right=0.90, top=0.94, bottom=0.16)
    image = ax.imshow(matrix.to_numpy(), cmap="cividis", vmin=0, vmax=100, aspect="auto")
    ax.set_xticks(np.arange(6))
    ax.set_xticklabels(GROUP_LABELS, fontsize=12.0, fontweight="bold")
    for index, tick in enumerate(ax.get_xticklabels()):
        if index in {1, 4}:
            tick.set_color("#5B3A87")
    ax.set_yticks(np.arange(4))
    ax.set_yticklabels(row_labels, fontsize=11.5, fontweight="bold")
    ax.tick_params(which="both", length=0, pad=8)
    ax.set_xticks(np.arange(-0.5, 6, 1), minor=True)
    ax.set_yticks(np.arange(-0.5, 4, 1), minor=True)
    ax.grid(which="minor", color="white", linewidth=2.0)
    for spine in ax.spines.values():
        spine.set_visible(False)

    cmap = plt.get_cmap("cividis")
    for row in range(4):
        for column in range(6):
            percent = float(matrix.iloc[row, column])
            count = int(counts.iloc[row, column])
            red, green, blue, _alpha = cmap(percent / 100.0)
            luminance = 0.2126 * red + 0.7152 * green + 0.0722 * blue
            color = "#111827" if luminance > 0.58 else "white"
            ax.text(
                column,
                row - 0.08,
                f"{percent:.0f}%",
                ha="center",
                va="center",
                fontsize=14.0,
                fontweight="bold",
                color=color,
            )
            ax.text(
                column,
                row + 0.24,
                f"{count}/{sizes[row]} genes",
                ha="center",
                va="center",
                fontsize=9.2,
                color=color,
            )

    colorbar = fig.colorbar(image, ax=ax, fraction=0.035, pad=0.035)
    colorbar.set_ticks([0, 25, 50, 75, 100])
    colorbar.ax.tick_params(labelsize=9.5, length=0)
    colorbar.set_label("Pathway genes observed as DEGs", fontsize=11.0)
    colorbar.outline.set_visible(False)

    base = output_dir / "rosmap_pathway_deg_overlap_by_sex_apoe"
    saved = save_publication_figure(
        fig,
        base,
        formats=["png"],
        dpi=320,
        bbox_inches="tight",
        pad_inches=0.08,
    )
    plt.close(fig)
    artifacts = {path.suffix.lstrip("."): path for path in saved}
    if set(artifacts) != {"png"}:
        raise RuntimeError(f"Missing figure outputs: {artifacts}")
    frame.to_csv(base.with_name(base.name + "_statistics.tsv"), sep="\t", index=False)
    return artifacts


def clear_slide(slide) -> None:
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = ui.OFF_WHITE


def add_program_card(
    slide,
    *,
    x: float,
    y: float,
    title: str,
    size: int,
    description: str,
    background,
    accent,
) -> None:
    ui.add_rect(slide, x, y, 5.74, 1.70, color=background, outline=None)
    ui.add_text(
        slide,
        title,
        x + 0.30,
        y + 0.27,
        4.28,
        0.42,
        size=16.0,
        color=ui.NAVY,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    ui.add_rect(slide, x + 4.67, y + 0.24, 0.80, 0.52, color=ui.WHITE, outline=None)
    ui.add_text(
        slide,
        str(size),
        x + 4.67,
        y + 0.31,
        0.80,
        0.25,
        size=18.0,
        color=ui.readable_accent(accent),
        bold=True,
        font=ui.FONT_HEAD,
        align=PP_ALIGN.CENTER,
    )
    ui.add_text(
        slide,
        "genes",
        x + 4.67,
        y + 0.57,
        0.80,
        0.16,
        size=8.5,
        color=ui.GRAY,
        align=PP_ALIGN.CENTER,
    )
    ui.add_text(
        slide,
        description,
        x + 0.30,
        y + 0.88,
        5.12,
        0.50,
        size=12.0,
        color=ui.DARK,
        valign=MSO_ANCHOR.MIDDLE,
    )


def build_program_list_slide(slide) -> None:
    ui.add_title_block(slide, LIST_TITLE)
    slide.shapes[-1].height = Inches(0.47)
    positions = [(0.72, 1.42), (6.87, 1.42), (0.72, 3.48), (6.87, 3.48)]
    for position, program in zip(positions, PROGRAMS, strict=True):
        _module_id, title, size, description, background, accent = program
        add_program_card(
            slide,
            x=position[0],
            y=position[1],
            title=title,
            size=size,
            description=description,
            background=background,
            accent=accent,
        )
    ui.add_rect(slide, 0.72, 5.62, 11.90, 1.04, color=ui.PALE_GRAY, outline=None)
    ui.add_text(
        slide,
        "Two related questions",
        1.02,
        5.82,
        1.72,
        0.26,
        size=12.0,
        color=ui.PURPLE,
        bold=True,
    )
    ui.add_text(
        slide,
        "Overlap: which pathway genes are DEGs?   →   Enrichment: is that overlap larger than expected among detectable genes?",
        2.90,
        5.78,
        9.20,
        0.48,
        size=13.2,
        color=ui.DARK,
        valign=MSO_ANCHOR.MIDDLE,
    )
    ui.add_notes(
        slide,
        goal="Introduce the four mitochondrial pathways highlighted in the focused results.",
        walkthrough=(
            "The focused analysis highlights four fixed pathway definitions, using "
            "the same gene membership in every contrast. The first contains the 13 "
            "OXPHOS subunits encoded by mitochondrial DNA. The second contains 86 "
            "structural OXPHOS subunits encoded by nuclear DNA. The third contains "
            "155 genes involved in translating proteins inside mitochondria. The "
            "fourth contains 19 MIB/MICOS genes that organize the inner mitochondrial "
            "membrane and its folds, called cristae."
        ),
        boundary=(
            "These are focused benchmark pathways, not the full pathway search. Phase "
            "11 also tested 46 Level 1 and Level 2 MitoCarta pathway categories and "
            "all 149 MitoCarta pathways. Listing a pathway, or observing some of its "
            "genes among the DEGs, does not establish enrichment or altered function."
        ),
        transition="First inspect which genes in these programs appeared among the DEGs.",
    )


def build_overlap_slide(slide, figure_path: Path) -> None:
    ui.add_title_block(
        slide,
        OVERLAP_TITLE,
        "Each cell shows unique DEG genes ÷ all genes in that pathway for one sex/APOE group.",
    )
    ui.add_picture_contain(
        slide,
        figure_path,
        0.72,
        1.25,
        11.90,
        4.84,
        alt=(
            "Heatmap showing the percentage and count of genes from four mitochondrial "
            "pathways that appeared as significant DEGs in at least one fine-cell "
            "comparison for each of six sex/APOE groups."
        ),
    )
    ui.add_rect(slide, 0.72, 6.14, 11.90, 0.72, color=ui.PALE_GRAY, outline=None)
    ui.add_text(
        slide,
        "Interpret carefully",
        1.02,
        6.35,
        1.52,
        0.25,
        size=12.0,
        color=ui.PURPLE,
        bold=True,
    )
    ui.add_text(
        slide,
        "Male ε2 has the broadest overlap in three pathways, consistent with its larger overall DEG count; overlap alone is not enrichment.",
        2.70,
        6.29,
        9.44,
        0.36,
        size=13.0,
        color=ui.DARK,
        valign=MSO_ANCHOR.MIDDLE,
    )
    ui.add_notes(
        slide,
        goal="Show the breadth of pathway-gene overlap with the fine-cell DEG results.",
        walkthrough=(
            "For each sex/APOE group, a pathway gene is counted once if it was a "
            "significant DEG in at least one fine-cell comparison. The denominator is "
            "the full frozen pathway size shown on the previous slide. Twelve of the "
            "13 mitochondrial-DNA OXPHOS genes appeared in every group. Male epsilon-2 "
            "had the broadest overlap in the other three pathways: 67 of 86 nuclear "
            "structural OXPHOS genes, 130 of 155 mitochondrial-translation genes, and "
            "17 of 19 inner-membrane genes. Female epsilon-3 homozygous had narrower "
            "overlap in mitochondrial translation and inner-membrane organization."
        ),
        boundary=(
            "This heatmap collapses across fine cell types and records presence, not "
            "direction, recurrence, effect size, or statistical enrichment. Groups "
            "with more overall DEGs will tend to overlap more pathway genes. Formal "
            "over-representation analysis instead uses the detectable-gene background "
            "within each comparison and controls the false-discovery rate."
        ),
        transition=(
            "Now move from overlap breadth to direction within the two OXPHOS gene sets."
        ),
    )


def validate_deck(prs: Presentation, preserved: tuple[Any, ...]) -> None:
    expected_slides = len(preserved) + 2
    if len(prs.slides) != expected_slides:
        raise RuntimeError(
            f"Expected {expected_slides} slides, found {len(prs.slides)}"
        )
    expected = (PREVIOUS_TITLE, LIST_TITLE, OVERLAP_TITLE, NEXT_TITLE)
    observed = tuple(
        slide_title(prs.slides[index])
        for index in range(INSERT_INDEX - 1, INSERT_INDEX + 3)
    )
    if observed != expected:
        raise RuntimeError(f"Unexpected pathway-section order: {observed!r}")
    retained = tuple(
        slide_fingerprint(slide)
        for index, slide in enumerate(prs.slides)
        if index not in {INSERT_INDEX, INSERT_INDEX + 1}
    )
    if retained != preserved:
        raise RuntimeError("A retained slide changed during pathway-slide insertion")
    for index, required in (
        (INSERT_INDEX, ["13", "86", "155", "19", "Two related questions"]),
        (INSERT_INDEX + 1, ["unique DEG genes", "overlap alone is not enrichment"]),
    ):
        slide = prs.slides[index]
        visible = "\n".join(
            shape.text
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
        )
        for item in required:
            if item not in visible:
                raise RuntimeError(f"Missing {item!r} from slide {index + 1}")
        notes = slide.notes_slide.notes_text_frame
        if notes is None or not notes.text.strip():
            raise RuntimeError(f"Missing speaker notes on slide {index + 1}")
        for shape in slide.shapes:
            if (
                shape.left < 0
                or shape.top < 0
                or shape.left + shape.width > prs.slide_width + 10
                or shape.top + shape.height > prs.slide_height + 10
            ):
                raise RuntimeError(
                    f"Out-of-bounds shape on slide {index + 1}: {shape.name}"
                )


def main() -> int:
    args = parse_args()
    input_path = args.input.resolve()
    output_path = args.output.resolve()
    audit_root = args.audit_root.resolve()
    overlap = prepare_overlap_data()
    artifacts = render_overlap_figure(overlap, audit_root / "figures")
    if args.figure_only:
        print(f"Rendered figure: {artifacts['png']}")
        return 0
    if not input_path.is_file():
        raise FileNotFoundError(input_path)

    input_hash = sha256(input_path)
    prs = Presentation(input_path)
    if (
        len(prs.slides) == INPUT_SLIDES
        and slide_title(prs.slides[INSERT_INDEX - 1]) == PREVIOUS_TITLE
        and slide_title(prs.slides[INSERT_INDEX]) == NEXT_TITLE
    ):
        preserved = tuple(slide_fingerprint(slide) for slide in prs.slides)
        ui.set_notes_body_template(
            prs.slides[0].notes_slide.notes_placeholder._element
        )
        list_slide = ui.new_slide(prs)
        overlap_slide = ui.new_slide(prs)
        build_program_list_slide(list_slide)
        build_overlap_slide(overlap_slide, artifacts["png"])
        slide_ids = prs.slides._sldIdLst
        list_id, overlap_id = slide_ids[-2], slide_ids[-1]
        slide_ids.remove(list_id)
        slide_ids.remove(overlap_id)
        slide_ids.insert(INSERT_INDEX, list_id)
        slide_ids.insert(INSERT_INDEX + 1, overlap_id)
        action = "Inserted"
    elif (
        len(prs.slides) in UPDATED_OUTPUT_SLIDES
        and slide_title(prs.slides[INSERT_INDEX - 1]) == PREVIOUS_TITLE
        and slide_title(prs.slides[INSERT_INDEX]) == LIST_TITLE
        and slide_title(prs.slides[INSERT_INDEX + 1]) == OVERLAP_TITLE
        and slide_title(prs.slides[INSERT_INDEX + 2]) == NEXT_TITLE
    ):
        preserved = tuple(
            slide_fingerprint(slide)
            for index, slide in enumerate(prs.slides)
            if index not in {INSERT_INDEX, INSERT_INDEX + 1}
        )
        clear_slide(prs.slides[INSERT_INDEX])
        clear_slide(prs.slides[INSERT_INDEX + 1])
        build_program_list_slide(prs.slides[INSERT_INDEX])
        build_overlap_slide(prs.slides[INSERT_INDEX + 1], artifacts["png"])
        action = "Updated"
    else:
        raise RuntimeError(
            f"Unexpected deck state: {len(prs.slides)} slides; slides 5-8 are "
            f"{[slide_title(prs.slides[i]) for i in range(4, min(8, len(prs.slides)))]!r}"
        )
    validate_deck(prs, preserved)

    audit_root.mkdir(parents=True, exist_ok=True)
    backup_dir = audit_root / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / f"input_deck_{input_hash[:12]}.pptx"
    if not backup_path.exists():
        shutil.copy2(input_path, backup_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    handle, temp_name = tempfile.mkstemp(
        prefix=output_path.stem + ".", suffix=".pptx", dir=output_path.parent
    )
    os.close(handle)
    temp_path = Path(temp_name)
    try:
        prs.save(temp_path)
        reopened = Presentation(temp_path)
        validate_deck(reopened, preserved)
        os.replace(temp_path, output_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()

    print(f"{action} pathway slides 6-7: {output_path}")
    print(f"Figure: {artifacts['png']}")
    print(f"Backup: {backup_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
