#!/usr/bin/env python3
"""Insert a category-definition slide after slide 1 of the sex/APOE KDA deck."""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
import update_phase11_seaad_simple_aggr_part2 as ui  # noqa: E402


ROOT = Path(__file__).resolve().parents[2]
DECK = (
    ROOT
    / "docs"
    / "presentations"
    / "09162026"
    / "09162026_sex_apoe_kda_fine_broad.pptx"
)
AUDIT_ROOT = (
    ROOT
    / "results"
    / "presentations"
    / "09162026_sex_apoe_categories_slide"
)

EXPECTED_INPUT_SLIDES = 26
INSERT_AFTER = 1
NEW_TITLE = "What is a category"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DECK)
    parser.add_argument("--output", type=Path, default=DECK)
    parser.add_argument("--audit-root", type=Path, default=AUDIT_ROOT)
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def slide_text(slide) -> str:
    return "\n".join(
        shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
    )


def semantic_slide_fingerprint(slide) -> tuple[Any, ...]:
    shapes = []
    for shape in slide.shapes:
        text = shape.text if getattr(shape, "has_text_frame", False) else ""
        shapes.append(
            (
                int(shape.shape_type),
                shape.name,
                int(shape.left),
                int(shape.top),
                int(shape.width),
                int(shape.height),
                text,
            )
        )
    notes = slide.notes_slide.notes_text_frame
    return tuple(shapes), notes.text if notes is not None else ""


def add_factor_card(
    slide,
    x: float,
    width: float,
    number: str,
    label: str,
    detail: str,
    background,
    accent,
) -> None:
    ui.add_rect(
        slide,
        x,
        1.42,
        width,
        1.14,
        color=background,
        outline=accent,
    )
    ui.add_text(
        slide,
        number,
        x + 0.18,
        1.62,
        0.62,
        0.54,
        size=28,
        color=ui.readable_accent(accent),
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        font=ui.FONT_HEAD,
    )
    ui.add_text(
        slide,
        label,
        x + 0.88,
        1.61,
        width - 1.05,
        0.28,
        size=11.8,
        color=ui.NAVY,
        bold=True,
    )
    ui.add_text(
        slide,
        detail,
        x + 0.88,
        1.96,
        width - 1.05,
        0.34,
        size=9.4,
        color=ui.GRAY,
    )


def add_operator(slide, value: str, x: float) -> None:
    ui.add_text(
        slide,
        value,
        x,
        1.73,
        0.42,
        0.42,
        size=22,
        color=ui.GRAY,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        font=ui.FONT_HEAD,
    )


def add_group_chip(slide, text: str, x: float, y: float, accent) -> None:
    ui.add_rect(
        slide,
        x,
        y,
        1.34,
        0.46,
        color=ui.WHITE,
        outline=accent,
    )
    ui.add_text(
        slide,
        text,
        x + 0.08,
        y + 0.10,
        1.18,
        0.22,
        size=11.2,
        color=ui.NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_cell_chip(slide, text: str, x: float, y: float, width: float) -> None:
    ui.add_rect(
        slide,
        x,
        y,
        width,
        0.48,
        color=ui.WHITE,
        outline=ui.LIGHT,
    )
    ui.add_text(
        slide,
        text,
        x + 0.12,
        y + 0.11,
        width - 0.24,
        0.23,
        size=10.5,
        color=ui.DARK,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )


def build_category_slide(prs: Presentation):
    slide = ui.new_slide(prs)
    ui.add_title_block(
        slide,
        NEW_TITLE,
        "One category selects one sex, one APOE group, and one broad cell type.",
    )

    add_factor_card(
        slide,
        0.66,
        2.22,
        "2",
        "SEXES",
        "Female • Male",
        ui.PALE_SKY,
        ui.BLUE,
    )
    add_operator(slide, "×", 2.98)
    add_factor_card(
        slide,
        3.50,
        2.78,
        "3",
        "APOE GROUPS",
        "ε2 carrier • ε3/ε3 • ε4 carrier",
        ui.PALE_GREEN,
        ui.TEAL,
    )
    add_operator(slide, "×", 6.38)
    add_factor_card(
        slide,
        6.90,
        2.32,
        "7",
        "BROAD TYPES",
        "Frozen cell classes",
        ui.PALE_GOLD,
        ui.GOLD,
    )
    add_operator(slide, "=", 9.32)
    add_factor_card(
        slide,
        9.84,
        2.83,
        "42",
        "CATEGORIES",
        "Possible analysis contexts",
        ui.PALE_RED,
        ui.VERMILION,
    )
    ui.add_text(
        slide,
        "Key-driver genes are ranked separately within each category.",
        0.66,
        2.62,
        12.01,
        0.26,
        size=11.6,
        color=ui.VERMILION_TEXT,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    ui.add_rect(
        slide,
        0.66,
        2.90,
        5.93,
        3.16,
        color=ui.PALE_SKY,
        outline=None,
    )
    ui.add_panel_title(
        slide,
        "Six sex/APOE groups",
        0.96,
        3.18,
        5.30,
        accent=ui.BLUE,
    )
    ui.add_text(
        slide,
        "FEMALE",
        0.98,
        3.88,
        0.86,
        0.24,
        size=9.6,
        color=ui.GRAY,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    ui.add_text(
        slide,
        "MALE",
        0.98,
        4.56,
        0.86,
        0.24,
        size=9.6,
        color=ui.GRAY,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    for index, suffix in enumerate(("e2", "e33", "e4")):
        x = 1.86 + index * 1.50
        add_group_chip(slide, f"F_{suffix}", x, 3.76, ui.BLUE)
        add_group_chip(slide, f"M_{suffix}", x, 4.44, ui.BLUE)
    ui.add_text(
        slide,
        "e2 = ε2 carrier   •   e33 = ε3/ε3   •   e4 = ε4 carrier",
        0.98,
        5.28,
        5.26,
        0.30,
        size=9.8,
        color=ui.GRAY,
        align=PP_ALIGN.CENTER,
    )

    ui.add_rect(
        slide,
        6.75,
        2.90,
        5.92,
        3.16,
        color=ui.PALE_GOLD,
        outline=None,
    )
    ui.add_panel_title(
        slide,
        "Seven broad cell types",
        7.05,
        3.18,
        5.30,
        accent=ui.GOLD,
    )
    left_types = ["Astrocytes", "Excitatory neurons", "Inhibitory neurons", "Microglia"]
    right_types = ["OPCs", "Oligodendrocytes", "Vasculature"]
    for index, label in enumerate(left_types):
        add_cell_chip(slide, label, 7.05, 3.72 + index * 0.55, 2.55)
    for index, label in enumerate(right_types):
        add_cell_chip(slide, label, 9.82, 3.72 + index * 0.55, 2.55)

    ui.add_rect(
        slide,
        0.66,
        6.30,
        12.01,
        0.66,
        color=ui.WHITE,
        outline=ui.LIGHT,
    )
    ui.add_text(
        slide,
        "Example category",
        0.94,
        6.48,
        1.42,
        0.26,
        size=10.2,
        color=ui.GRAY,
        bold=True,
    )
    ui.add_text(
        slide,
        "M_e33 × Excitatory neurons",
        2.34,
        6.43,
        3.36,
        0.32,
        size=13.2,
        color=ui.PURPLE,
        bold=True,
    )
    ui.add_text(
        slide,
        "42 is the design space; later slides show which categories have usable calls and returned drivers.",
        5.64,
        6.44,
        6.72,
        0.34,
        size=9.8,
        color=ui.GRAY,
        align=PP_ALIGN.RIGHT,
    )
    ui.add_notes(
        slide,
        goal="Define the category unit before presenting cohort-specific counts.",
        walkthrough="Start with two sexes and three APOE groups: epsilon-2 carriers, epsilon-3 homozygotes, and epsilon-4 carriers. Their cross creates six sex/APOE groups. Crossing each group with the seven frozen broad cell types creates 42 possible categories. For example, M_e33 by excitatory neurons is one category. Key-driver genes are ranked separately within each category.",
        boundary="Forty-two is the structural design space. A possible category may lack an eligible KDA call or a returned non-mitochondrial driver, so it should not be interpreted as an observed negative result.",
        transition="With the category unit defined, compare the two cohorts and their shared aggregation rule.",
    )
    return slide


def write_audit(
    audit_path: Path,
    checks: list[dict[str, Any]],
) -> None:
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    with audit_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["check_id", "observed", "expected", "passed"],
            delimiter="\t",
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(checks)


def main() -> int:
    args = parse_args()
    input_path = args.input.resolve()
    output_path = args.output.resolve()
    audit_root = args.audit_root.resolve()
    if not input_path.is_file():
        raise FileNotFoundError(input_path)

    prs = Presentation(str(input_path))
    if len(prs.slides) != EXPECTED_INPUT_SLIDES:
        raise RuntimeError(
            f"Expected {EXPECTED_INPUT_SLIDES} slides, found {len(prs.slides)}"
        )
    if NEW_TITLE in "\n".join(slide_text(slide) for slide in prs.slides):
        raise RuntimeError("Category slide already present; refusing to insert twice")

    before_fingerprints = [
        semantic_slide_fingerprint(slide) for slide in prs.slides
    ]
    input_hash = sha256_file(input_path)
    backup_dir = audit_root / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / f"input_deck_{input_hash[:12]}.pptx"
    if not backup_path.exists():
        shutil.copy2(input_path, backup_path)

    ui.set_notes_body_template(
        prs.slides[0].notes_slide.notes_placeholder._element
    )
    build_category_slide(prs)
    slide_id_list = prs.slides._sldIdLst
    slide_ids = list(slide_id_list)
    slide_id_list.remove(slide_ids[-1])
    slide_id_list.insert(INSERT_AFTER, slide_ids[-1])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.name}.",
        suffix=".tmp",
        dir=output_path.parent,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        prs.save(str(temporary))
        reloaded = Presentation(str(temporary))
        after_fingerprints = [
            semantic_slide_fingerprint(slide) for slide in reloaded.slides
        ]
        other_slides_unchanged = (
            after_fingerprints[:INSERT_AFTER]
            == before_fingerprints[:INSERT_AFTER]
            and after_fingerprints[INSERT_AFTER + 1 :]
            == before_fingerprints[INSERT_AFTER:]
        )
        new_slide_text = slide_text(reloaded.slides[INSERT_AFTER])
        notes = reloaded.slides[INSERT_AFTER].notes_slide.notes_text_frame
        checks: list[dict[str, Any]] = [
            {
                "check_id": "output_slide_count",
                "observed": len(reloaded.slides),
                "expected": EXPECTED_INPUT_SLIDES + 1,
                "passed": len(reloaded.slides) == EXPECTED_INPUT_SLIDES + 1,
            },
            {
                "check_id": "category_slide_is_slide_2",
                "observed": NEW_TITLE in new_slide_text,
                "expected": True,
                "passed": NEW_TITLE in new_slide_text,
            },
            {
                "check_id": "category_math_present",
                "observed": all(
                    value in new_slide_text
                    for value in ("2", "3", "7", "42", "F_e2", "M_e4")
                ),
                "expected": True,
                "passed": all(
                    value in new_slide_text
                    for value in ("2", "3", "7", "42", "F_e2", "M_e4")
                ),
            },
            {
                "check_id": "all_other_slides_semantically_unchanged",
                "observed": other_slides_unchanged,
                "expected": True,
                "passed": other_slides_unchanged,
            },
            {
                "check_id": "new_slide_has_notes",
                "observed": bool(notes and notes.text.strip()),
                "expected": True,
                "passed": bool(notes and notes.text.strip()),
            },
        ]
        failed = [row["check_id"] for row in checks if not row["passed"]]
        if failed:
            raise RuntimeError("Slide insertion checks failed: " + ", ".join(failed))
        os.replace(temporary, output_path)
    finally:
        temporary.unlink(missing_ok=True)

    write_audit(audit_root / "category_slide_insert_checks.tsv", checks)
    print(f"Updated deck: {output_path}")
    print(f"Inserted slide: {INSERT_AFTER + 1}")
    print(f"Backup: {backup_path}")
    print(f"Audit: {audit_root}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
